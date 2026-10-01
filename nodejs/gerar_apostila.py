# -*- coding: utf-8 -*-
"""
Gera a apostila impressa da Aula 1 — Preparando o ambiente Node.js no Ubuntu Server.

Requisitos:
    pip install python-docx

Uso:
    python gerar_apostila.py
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ARQUIVO_SAIDA = "Apostila_Ambiente_NodeJS.docx"

# ------------------------------------------------------------
# Cores e fontes
# ------------------------------------------------------------
COR_PRIMARIA = RGBColor(0x0B, 0x5C, 0x5C)   # verde-petróleo
COR_DESTAQUE = RGBColor(0xE8, 0x8B, 0x1A)   # laranja
COR_ALERTA = RGBColor(0xC5, 0x36, 0x2B)     # vermelho
COR_CINZA = RGBColor(0x55, 0x55, 0x55)
COR_TEXTO = RGBColor(0x1A, 0x1A, 0x1A)

FUNDO_CODIGO = "EDEDED"      # cinza claro
FUNDO_CHECK = "E6F4EF"       # verde suave
FUNDO_ALERTA = "FFF2E0"      # laranja suave
FUNDO_DICA = "EAF0FB"        # azul suave
FUNDO_CABECALHO = "0B5C5C"   # cabeçalho de tabela

FONTE_TEXTO = "Calibri"
FONTE_CODIGO = "Consolas"


# ------------------------------------------------------------
# Helpers de formatação
# ------------------------------------------------------------
def sombrear_paragrafo(paragrafo, cor_hex):
    """Aplica cor de fundo ao parágrafo inteiro."""
    p_pr = paragrafo._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    p_pr.append(shd)


def borda_esquerda(paragrafo, cor_hex, espessura=24):
    """Coloca uma barra colorida à esquerda do parágrafo."""
    p_pr = paragrafo._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(espessura))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), cor_hex)
    bdr.append(left)
    p_pr.append(bdr)


def sombrear_celula(celula, cor_hex):
    tc_pr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    tc_pr.append(shd)


def repetir_cabecalho(linha):
    """Repete a linha de cabeçalho se a tabela quebrar de página."""
    tr_pr = linha._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def fonte_run(run, nome):
    run.font.name = nome
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(atributo), nome)


# ------------------------------------------------------------
# Blocos de conteúdo
# ------------------------------------------------------------
def titulo(doc, texto, nivel=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if nivel == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(17 if nivel == 1 else 13)
    run.font.color.rgb = COR_PRIMARIA if nivel == 1 else COR_DESTAQUE
    if nivel == 1:
        p_pr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "0B5C5C")
        bdr.append(bottom)
        p_pr.append(bdr)
    return p


def paragrafo(doc, texto="", negrito_inicio=None, italico=False, cor=None, tamanho=11):
    """Parágrafo simples. Se negrito_inicio for dado, ele vem em negrito antes do texto."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if negrito_inicio:
        r = p.add_run(negrito_inicio)
        r.bold = True
        r.font.size = Pt(tamanho)
    r = p.add_run(texto)
    r.italic = italico
    r.font.size = Pt(tamanho)
    if cor:
        r.font.color.rgb = cor
    return p


def marcador(doc, texto, negrito_inicio=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if negrito_inicio:
        p.add_run(negrito_inicio).bold = True
    p.add_run(texto)
    return p


def numerado(doc, texto):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(texto)
    return p


def bloco_codigo(doc, linhas, legenda=None):
    """Bloco de código: fundo cinza claro, fonte monoespaçada."""
    if legenda:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(legenda)
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = COR_CINZA
    if isinstance(linhas, str):
        linhas = [linhas]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    sombrear_paragrafo(p, FUNDO_CODIGO)
    borda_esquerda(p, "0B5C5C", 24)
    for i, linha in enumerate(linhas):
        run = p.add_run(linha)
        fonte_run(run, FONTE_CODIGO)
        run.font.size = Pt(10.5)
        run.font.color.rgb = COR_TEXTO
        if i < len(linhas) - 1:
            run.add_break(WD_BREAK.LINE)
    return p


def caixa(doc, rotulo, texto, fundo, cor_borda, cor_rotulo):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    sombrear_paragrafo(p, fundo)
    borda_esquerda(p, cor_borda, 24)
    r = p.add_run(rotulo + " ")
    r.bold = True
    r.font.color.rgb = cor_rotulo
    p.add_run(texto)
    return p


def checkpoint(doc, texto):
    return caixa(doc, "✅ Checkpoint:", texto, FUNDO_CHECK, "0B5C5C", COR_PRIMARIA)


def alerta(doc, texto):
    return caixa(doc, "⚠️ Atenção:", texto, FUNDO_ALERTA, "E88B1A", COR_ALERTA)


def dica(doc, texto):
    return caixa(doc, "💡 Dica:", texto, FUNDO_DICA, "4A6FB5", RGBColor(0x2E, 0x4F, 0x8F))


def tabela(doc, cabecalho, linhas, larguras_cm=None, centralizar_ultima=False):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    linha_cab = t.rows[0]
    repetir_cabecalho(linha_cab)
    for i, texto in enumerate(cabecalho):
        cel = linha_cab.cells[i]
        sombrear_celula(cel, FUNDO_CABECALHO)
        cel.text = ""
        r = cel.paragraphs[0].add_run(texto)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for n, dados in enumerate(linhas):
        cels = t.add_row().cells
        for i, texto in enumerate(dados):
            cels[i].text = ""
            par = cels[i].paragraphs[0]
            r = par.add_run(texto)
            r.font.size = Pt(10.5)
            if i == 0:
                r.bold = True
            if centralizar_ultima and i == len(dados) - 1:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if n % 2 == 1:
                sombrear_celula(cels[i], "F4F7F7")
    if larguras_cm:
        for row in t.rows:
            for i, w in enumerate(larguras_cm):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def linhas_em_branco(doc, quantidade):
    for _ in range(quantidade):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(24)
        r = p.add_run("_" * 86)
        r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
        r.font.size = Pt(11)


def etapa(doc, numero, nome):
    titulo(doc, f"Etapa {numero} — {nome}", nivel=2)


# ------------------------------------------------------------
# Montagem do documento
# ------------------------------------------------------------
def configurar_documento():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = FONTE_TEXTO
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE_TEXTO)
    return doc


def rodape(doc):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Servidor Web com Node.js — Aula 1: Preparando o ambiente   |   Página ")
    r.font.size = Pt(9)
    r.font.color.rgb = COR_CINZA
    # Campo de número de página
    run = p.add_run()
    run.font.size = Pt(9)
    run.font.color.rgb = COR_CINZA
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tipo)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = texto
            run._r.append(instr)
    # Capa sem rodapé
    sec.different_first_page_header_footer = True


def capa(doc):
    for _ in range(5):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("REDES DE COMPUTADORES · AULA PRÁTICA 1 DE 2")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = COR_DESTAQUE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("Seu primeiro servidor web\ncom Node.js")
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = COR_PRIMARIA

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("Parte 1 — Preparando o ambiente no Ubuntu Server")
    r.font.size = Pt(15)
    r.font.color.rgb = COR_CINZA

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("Ubuntu Server · Node.js 22 · npm · PM2 · Firewall ufw")
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = COR_CINZA

    for _ in range(6):
        doc.add_paragraph()

    # Campos de identificação
    t = doc.add_table(rows=4, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    campos = ["Nome do aluno:", "Turma:", "Data:", "Professor:"]
    for i, campo in enumerate(campos):
        c0, c1 = t.rows[i].cells
        c0.width = Cm(4.0)
        c1.width = Cm(12.0)
        sombrear_celula(c0, "F4F7F7")
        c0.text = ""
        r = c0.paragraphs[0].add_run(campo)
        r.bold = True
        r.font.color.rgb = COR_PRIMARIA
        c0.paragraphs[0].paragraph_format.space_before = Pt(8)
        c0.paragraphs[0].paragraph_format.space_after = Pt(8)
        c1.paragraphs[0].paragraph_format.space_before = Pt(8)
        c1.paragraphs[0].paragraph_format.space_after = Pt(8)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def conteudo(doc):
    # ---------------- Objetivo ----------------
    titulo(doc, "Objetivo da aula")
    paragrafo(doc, "Nesta aula você vai preparar um computador para funcionar como servidor web. "
                   "Ao final, a máquina terá o Ubuntu Server instalado e atualizado, o Node.js e o PM2 "
                   "instalados, a pasta do projeto criada, a porta 3000 liberada no firewall e a conexão "
                   "com outros computadores testada.")
    paragrafo(doc, "Hoje não vamos escrever código. Na próxima aula você vai criar os arquivos do site "
                   "e colocar o servidor no ar.", italico=True, cor=COR_CINZA)

    # ---------------- Pré-requisitos ----------------
    titulo(doc, "Pré-requisitos")
    marcador(doc, "Um computador que possa ser formatado (todo o conteúdo do disco será apagado).")
    marcador(doc, "Um pendrive bootável com o Ubuntu Server (preparado pelo professor ou pela turma).")
    marcador(doc, "Cabo de rede ligado ao switch/roteador do laboratório, com acesso à internet.")
    marcador(doc, "Um segundo computador na mesma rede, para o teste final.")
    marcador(doc, "Caderno ou esta apostila para anotar usuário, senha e IP.")
    dica(doc, "A imagem ISO do Ubuntu Server pode ser baixada em ubuntu.com/download/server "
              "(opcional — normalmente o professor já traz o pendrive pronto). "
              "Para gravar o pendrive no Windows, use programas como Rufus ou balenaEtcher.")

    # ---------------- Conceito ----------------
    titulo(doc, "Conceito rápido: o que é um servidor web?")
    paragrafo(doc, "Pense num restaurante:")
    tabela(doc,
           ["No restaurante", "Na rede", "Papel"],
           [
               ["Cliente", "Navegador (Chrome, Firefox…)", "Faz o pedido"],
               ["Garçom", "Servidor web", "Recebe o pedido e entrega o prato"],
               ["Cozinha", "Arquivos do site", "O que é servido"],
               ["Número da mesa", "Porta (a nossa será a 3000)", "Onde o garçom atende"],
           ],
           larguras_cm=[4.0, 6.0, 6.6])
    paragrafo(doc, "Ou seja: um servidor web é um programa que fica esperando pedidos e entrega páginas "
                   "para quem pedir. O computador onde ele roda também é chamado de servidor.")

    # ---------------- Etapas ----------------
    titulo(doc, "Passo a passo")
    paragrafo(doc, "Siga as etapas na ordem. Só avance quando o ", negrito_inicio=None)
    doc.paragraphs[-1].add_run("✅ Checkpoint").bold = True
    doc.paragraphs[-1].add_run(" da etapa der certo. Os comandos aparecem nos blocos cinza: digite "
                               "exatamente como está e aperte Enter. O símbolo $ indica o início do "
                               "comando — não digite o $.")

    # Etapa 1
    etapa(doc, 1, "Instalar o Ubuntu Server")
    numerado(doc, "Conecte o pendrive e ligue o computador.")
    numerado(doc, "Abra o menu de boot (geralmente F12, F8, F11 ou Esc) e escolha o pendrive.")
    numerado(doc, "Escolha \"Try or Install Ubuntu Server\".")
    numerado(doc, "Siga o instalador conforme a tabela abaixo.")
    tabela(doc,
           ["Tela do instalador", "O que escolher"],
           [
               ["Idioma", "Português (ou English, se preferir)"],
               ["Teclado", "Portuguese (Brazil)"],
               ["Rede", "Deixar automático (DHCP) — o IP é recebido sozinho"],
               ["Disco", "Usar o disco inteiro (Use an entire disk)"],
               ["Perfil", "Seu nome, nome do servidor, usuário e senha"],
               ["SSH", "Marcar \"Install OpenSSH server\""],
               ["Pacotes extras (snaps)", "Não marcar nada"],
               ["Fim", "Reboot Now — retire o pendrive quando pedir"],
           ],
           larguras_cm=[5.0, 11.6])
    alerta(doc, "\"Usar o disco inteiro\" apaga tudo que está no HD. Confirme com o professor antes de avançar.")
    paragrafo(doc, "Anote aqui:  Usuário: ____________________   Nome do servidor: ____________________")
    dica(doc, "Ao digitar a senha no Linux, nada aparece na tela (nem asteriscos). É normal — digite e aperte Enter.")
    checkpoint(doc, "aparece a mensagem \"nome-do-servidor login:\" e você consegue entrar com seu usuário e senha.")

    # Etapa 2
    etapa(doc, 2, "Atualizar o sistema")
    paragrafo(doc, "Primeiro atualizamos a lista de programas disponíveis; depois instalamos as atualizações.")
    bloco_codigo(doc, "$ sudo apt update", "Atualiza a lista de programas:")
    bloco_codigo(doc, "$ sudo apt upgrade -y", "Instala as atualizações:")
    marcador(doc, " significa \"executar como administrador\". Vai pedir a sua senha.", negrito_inicio="sudo")
    marcador(doc, " responde \"sim\" automaticamente às perguntas.", negrito_inicio="-y")
    dica(doc, "Pode demorar alguns minutos. Se aparecer uma tela perguntando quais serviços reiniciar, apenas aperte Enter.")
    checkpoint(doc, "os comandos terminam sem linhas começando com \"E:\" (que indicam erro) e o prompt $ volta a aparecer.")

    # Etapa 3
    etapa(doc, 3, "Descobrir o IP da máquina")
    bloco_codigo(doc, "$ ip a")
    paragrafo(doc, "Procure a placa de rede (nome começando com enp, ens ou eth) e a linha que começa com inet. Exemplo:")
    bloco_codigo(doc, [
        "2: enp0s3: <BROADCAST,MULTICAST,UP,LOWER_UP> ...",
        "    inet 192.168.1.50/24 brd 192.168.1.255 scope global enp0s3",
    ], "Exemplo de saída (trecho):")
    paragrafo(doc, "Neste exemplo, o IP é 192.168.1.50 (o /24 não faz parte do IP). Ignore a interface lo "
                   "(127.0.0.1) — ela representa a própria máquina.")
    bloco_codigo(doc, "$ hostname -I", "Atalho que mostra só o IP:")
    paragrafo(doc, "Meu IP:  ______ . ______ . ______ . ______", negrito_inicio="")
    checkpoint(doc, "o IP do seu servidor está anotado acima.")

    # Etapa 4
    etapa(doc, 4, "Instalar o Node.js (versão 22, via NodeSource)")
    paragrafo(doc, "O Ubuntu já tem um Node.js no repositório dele, mas é uma versão antiga. "
                   "Por isso usamos a NodeSource, que fornece a versão 22.")
    bloco_codigo(doc, "$ curl -fsSL https://deb.nodesource.com/setup_22.x -o nodesource_setup.sh",
                 "1. Baixar o script de configuração:")
    bloco_codigo(doc, "$ sudo -E bash nodesource_setup.sh",
                 "2. Rodar o script (adiciona o repositório da NodeSource):")
    bloco_codigo(doc, "$ sudo apt install -y nodejs",
                 "3. Instalar o Node.js:")
    alerta(doc, "digite o endereço com cuidado. Um caractere errado faz o download falhar.")
    checkpoint(doc, "o último comando termina sem erros e o prompt $ volta a aparecer.")

    # Etapa 5
    etapa(doc, 5, "Verificar o Node.js e o npm")
    bloco_codigo(doc, "$ node -v")
    bloco_codigo(doc, "v22.x.x", "Resposta esperada (os x variam):")
    bloco_codigo(doc, "$ npm -v")
    bloco_codigo(doc, "10.x.x", "Resposta esperada (os x variam):")
    paragrafo(doc, "O -v pede a versão. Se o programa responde a versão, é porque está instalado. "
                   "O npm vem junto com o Node.js — não precisa instalar separado.")
    paragrafo(doc, "Versão do Node: ______________   Versão do npm: ______________")
    checkpoint(doc, "node -v começa com v22 e npm -v mostra um número de versão.")

    # Etapa 6
    etapa(doc, 6, "Criar a pasta do projeto")
    bloco_codigo(doc, "$ mkdir ~/meu-servidor-web", "Criar a pasta:")
    bloco_codigo(doc, "$ cd ~/meu-servidor-web", "Entrar na pasta:")
    bloco_codigo(doc, "$ pwd", "Confirmar onde você está:")
    bloco_codigo(doc, "/home/aluno/meu-servidor-web", "Resposta esperada (no lugar de \"aluno\" aparece o seu usuário):")
    tabela(doc,
           ["Símbolo / comando", "Significado"],
           [
               ["~", "Sua pasta pessoal (/home/seu-usuario)"],
               ["mkdir", "Cria uma pasta (make directory)"],
               ["cd", "Entra em uma pasta (change directory)"],
               ["pwd", "Mostra em qual pasta você está"],
           ],
           larguras_cm=[5.0, 11.6])
    checkpoint(doc, "o comando pwd mostra um caminho terminando em /meu-servidor-web.")

    # Etapa 7
    etapa(doc, 7, "Instalar o PM2")
    paragrafo(doc, "O PM2 funciona como um gerente de turno: mantém o servidor web funcionando e o "
                   "religa se ele cair. Hoje só vamos instalá-lo; ele será usado na próxima aula.")
    bloco_codigo(doc, "$ sudo npm install -g pm2", "Instalar globalmente (o -g vale para o sistema todo):")
    bloco_codigo(doc, "$ pm2 -v", "Conferir a versão:")
    dica(doc, "na primeira vez, o PM2 pode mostrar uma mensagem de inicialização antes da versão. É normal.")
    checkpoint(doc, "pm2 -v mostra um número de versão.")

    # Etapa 8
    etapa(doc, 8, "Liberar a porta 3000 no firewall (ufw)")
    paragrafo(doc, "O firewall é o porteiro da máquina: decide quais portas ficam abertas. "
                   "Vamos liberar o acesso remoto (SSH) e a porta 3000, onde o servidor web vai atender.")
    alerta(doc, "libere o SSH ANTES de ativar o firewall. Se ativar primeiro, quem estiver acessando "
                "a máquina remotamente perde a conexão.")
    bloco_codigo(doc, "$ sudo ufw allow OpenSSH", "1. Liberar o SSH:")
    bloco_codigo(doc, "$ sudo ufw allow 3000/tcp", "2. Liberar a porta 3000:")
    bloco_codigo(doc, "$ sudo ufw enable", "3. Ativar o firewall (responda y e Enter):")
    bloco_codigo(doc, "$ sudo ufw status", "4. Conferir:")
    bloco_codigo(doc, [
        "Status: active",
        "",
        "To                         Action      From",
        "--                         ------      ----",
        "OpenSSH                    ALLOW       Anywhere",
        "3000/tcp                   ALLOW       Anywhere",
    ], "Resposta esperada (podem aparecer também linhas com (v6)):")
    checkpoint(doc, "aparecem \"Status: active\" e a linha \"3000/tcp ALLOW\".")

    # Etapa 9
    etapa(doc, 9, "Testar a conectividade de outro computador")
    paragrafo(doc, "Vá até o computador de um colega (ou outro PC do laboratório) e faça um ping para o IP "
                   "que você anotou na Etapa 3. Troque 192.168.1.50 pelo seu IP.")
    bloco_codigo(doc, "> ping 192.168.1.50", "No Windows (abra o Prompt de Comando: tecla Windows, digite cmd):")
    bloco_codigo(doc, "$ ping -c 4 192.168.1.50", "No Linux:")
    bloco_codigo(doc, [
        "Resposta de 192.168.1.50: bytes=32 tempo<1ms TTL=64",
        "Resposta de 192.168.1.50: bytes=32 tempo<1ms TTL=64",
        "...",
        "Pacotes: Enviados = 4, Recebidos = 4, Perdidos = 0 (0% de perda)",
    ], "Resposta esperada (Windows):")
    dica(doc, "no Linux, se esquecer o -c 4, o ping não para sozinho. Aperte Ctrl + C para encerrar.")
    paragrafo(doc, "IP do computador de onde testei: ______ . ______ . ______ . ______")
    checkpoint(doc, "chegam respostas e a perda é de 0%. As duas máquinas se enxergam na rede.")

    # ---------------- Resumo ----------------
    titulo(doc, "Resumo do ambiente pronto")
    tabela(doc,
           ["Item", "Como conferir", "Status"],
           [
               ["Ubuntu Server instalado", "O login funciona", "✅"],
               ["Sistema atualizado", "sudo apt update termina sem erros", "✅"],
               ["IP anotado", "ip a ou hostname -I", "✅"],
               ["Node.js 22", "node -v mostra v22.x.x", "✅"],
               ["npm", "npm -v mostra a versão", "✅"],
               ["Pasta do projeto", "ls ~ mostra meu-servidor-web", "✅"],
               ["PM2", "pm2 -v mostra a versão", "✅"],
               ["Porta 3000 liberada", "sudo ufw status mostra 3000/tcp ALLOW", "✅"],
               ["Rede funcionando", "ping de outro computador responde", "✅"],
           ],
           larguras_cm=[5.2, 9.4, 2.0], centralizar_ultima=True)

    # ---------------- Problemas comuns ----------------
    titulo(doc, "Problemas comuns")
    tabela(doc,
           ["Problema", "Causa provável", "Solução"],
           [
               ["O computador não dá boot pelo pendrive",
                "Ordem de boot errada ou pendrive mal gravado",
                "Usar o menu de boot (F12/F8/Esc); regravar o pendrive"],
               ["Letras com acento ou ç saem erradas",
                "Layout de teclado errado na instalação",
                "Rodar: sudo dpkg-reconfigure keyboard-configuration"],
               ["\"Temporary failure resolving\" no apt update",
                "Máquina sem internet",
                "Conferir o cabo de rede e se ip a mostra um IP"],
               ["\"Could not get lock\" no apt",
                "Outra atualização rodando em segundo plano",
                "Esperar alguns minutos e tentar de novo"],
               ["ip a não mostra linha inet na placa de rede",
                "Cabo desconectado ou rede sem DHCP",
                "Conferir o cabo; reiniciar com sudo reboot"],
               ["\"node: command not found\"",
                "A instalação da Etapa 4 não terminou",
                "Repetir a Etapa 4 com atenção aos erros"],
               ["node -v mostra versão diferente de v22",
                "Instalou o Node.js antigo do Ubuntu",
                "Rodar o script da NodeSource e reinstalar o nodejs"],
               ["\"EACCES: permission denied\" no npm",
                "Faltou o sudo",
                "sudo npm install -g pm2"],
               ["Ping: \"Esgotado o tempo limite\"",
                "IP errado ou máquinas em redes diferentes",
                "Conferir o IP com ip a; confirmar que estão no mesmo switch/roteador"],
               ["Comando \"não encontrado\" em geral",
                "Erro de digitação",
                "Conferir letra por letra; o Linux diferencia maiúsculas de minúsculas"],
           ],
           larguras_cm=[5.2, 5.2, 6.2])

    # ---------------- Glossário ----------------
    titulo(doc, "Glossário rápido")
    tabela(doc,
           ["Termo", "Significado"],
           [
               ["HTTP", "A \"língua\" que o navegador e o servidor web usam para conversar"],
               ["Porta", "Número que identifica qual programa atende dentro da máquina (ex.: 3000)"],
               ["IP", "Endereço do computador na rede (ex.: 192.168.1.50)"],
               ["Node.js", "Programa que permite rodar JavaScript no servidor, fora do navegador"],
               ["npm", "Instalador de pacotes do Node.js; vem junto com ele"],
               ["PM2", "Gerenciador que mantém aplicações Node.js funcionando como serviço"],
               ["ufw", "Firewall simples do Ubuntu; abre e fecha portas"],
               ["sudo", "Executa um comando como administrador"],
               ["apt", "Gerenciador de programas do Ubuntu (instala e atualiza)"],
               ["ISO", "Arquivo com a imagem do instalador do sistema"],
               ["ping", "Comando que testa se outro computador responde na rede"],
           ],
           larguras_cm=[3.5, 13.1])

    # ---------------- Anotações ----------------
    titulo(doc, "Anotações")
    linhas_em_branco(doc, 14)

    # ---------------- Avaliação ----------------
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    titulo(doc, "Avaliação rápida")
    perguntas = [
        "Qual a diferença entre os comandos sudo apt update e sudo apt upgrade?",
        "Qual comando você usou para descobrir o IP da máquina? Qual era o IP do seu servidor?",
        "Como você confirmou que o Node.js e o npm foram instalados corretamente?",
        "Por que liberamos a porta 3000 no firewall? E por que liberamos o SSH antes de ativá-lo?",
        "O que significa quando o ping de outro computador recebe resposta do seu servidor?",
    ]
    for i, pergunta in enumerate(perguntas, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{i}. ")
        r.bold = True
        r.font.color.rgb = COR_PRIMARIA
        p.add_run(pergunta)
        linhas_em_branco(doc, 3)


def main():
    doc = configurar_documento()
    rodape(doc)
    capa(doc)
    conteudo(doc)
    doc.save(ARQUIVO_SAIDA)
    print(f"✅ Arquivo gerado: {os.path.abspath(ARQUIVO_SAIDA)}")


if __name__ == "__main__":
    main()
