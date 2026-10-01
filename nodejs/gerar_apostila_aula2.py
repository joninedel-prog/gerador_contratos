# -*- coding: utf-8 -*-
"""
Gera a apostila impressa da Aula 2 — Criando o site e colocando o servidor Node.js no ar.

Requisitos:
    pip install python-docx

Uso:
    python gerar_apostila_aula2.py
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ARQUIVO_SAIDA = "Apostila_Servidor_NodeJS.docx"

# ------------------------------------------------------------
# Cores e fontes
# ------------------------------------------------------------
COR_PRIMARIA = RGBColor(0x0B, 0x5C, 0x5C)   # verde-petróleo
COR_DESTAQUE = RGBColor(0xE8, 0x8B, 0x1A)   # laranja
COR_ALERTA = RGBColor(0xC5, 0x36, 0x2B)     # vermelho
COR_CINZA = RGBColor(0x55, 0x55, 0x55)
COR_TEXTO = RGBColor(0x1A, 0x1A, 0x1A)

FUNDO_CODIGO = "EDEDED"      # cinza claro
FUNDO_CHECK = "E6F4EF"       # verde suave
FUNDO_ALERTA = "FFF2E0"      # laranja suave
FUNDO_DICA = "EAF0FB"        # azul suave
FUNDO_CABECALHO = "0B5C5C"   # cabeçalho de tabela

FONTE_TEXTO = "Calibri"
FONTE_CODIGO = "Consolas"


# ------------------------------------------------------------
# Helpers de formatação
# ------------------------------------------------------------
def sombrear_paragrafo(paragrafo, cor_hex):
    """Aplica cor de fundo ao parágrafo inteiro."""
    p_pr = paragrafo._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    p_pr.append(shd)


def borda_esquerda(paragrafo, cor_hex, espessura=24):
    """Coloca uma barra colorida à esquerda do parágrafo."""
    p_pr = paragrafo._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(espessura))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), cor_hex)
    bdr.append(left)
    p_pr.append(bdr)


def sombrear_celula(celula, cor_hex):
    tc_pr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    tc_pr.append(shd)


def repetir_cabecalho(linha):
    """Repete a linha de cabeçalho se a tabela quebrar de página."""
    tr_pr = linha._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def fonte_run(run, nome):
    run.font.name = nome
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(atributo), nome)


# ------------------------------------------------------------
# Blocos de conteúdo
# ------------------------------------------------------------
def titulo(doc, texto, nivel=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if nivel == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(17 if nivel == 1 else 13)
    run.font.color.rgb = COR_PRIMARIA if nivel == 1 else COR_DESTAQUE
    if nivel == 1:
        p_pr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "0B5C5C")
        bdr.append(bottom)
        p_pr.append(bdr)
    return p


def paragrafo(doc, texto="", negrito_inicio=None, italico=False, cor=None, tamanho=11):
    """Parágrafo simples. Se negrito_inicio for dado, ele vem em negrito antes do texto."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if negrito_inicio:
        r = p.add_run(negrito_inicio)
        r.bold = True
        r.font.size = Pt(tamanho)
    r = p.add_run(texto)
    r.italic = italico
    r.font.size = Pt(tamanho)
    if cor:
        r.font.color.rgb = cor
    return p


def marcador(doc, texto, negrito_inicio=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if negrito_inicio:
        p.add_run(negrito_inicio).bold = True
    p.add_run(texto)
    return p


def numerado(doc, texto):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(texto)
    return p


def bloco_codigo(doc, linhas, legenda=None, tamanho=10.5):
    """Bloco de código: fundo cinza claro, fonte monoespaçada."""
    if legenda:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(legenda)
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = COR_CINZA
    if isinstance(linhas, str):
        linhas = [linhas]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    sombrear_paragrafo(p, FUNDO_CODIGO)
    borda_esquerda(p, "0B5C5C", 24)
    for i, linha in enumerate(linhas):
        run = p.add_run(linha)
        fonte_run(run, FONTE_CODIGO)
        run.font.size = Pt(tamanho)
        run.font.color.rgb = COR_TEXTO
        if i < len(linhas) - 1:
            run.add_break(WD_BREAK.LINE)
    return p


def caixa(doc, rotulo, texto, fundo, cor_borda, cor_rotulo):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    sombrear_paragrafo(p, fundo)
    borda_esquerda(p, cor_borda, 24)
    r = p.add_run(rotulo + " ")
    r.bold = True
    r.font.color.rgb = cor_rotulo
    p.add_run(texto)
    return p


def checkpoint(doc, texto):
    return caixa(doc, "✅ Checkpoint:", texto, FUNDO_CHECK, "0B5C5C", COR_PRIMARIA)


def alerta(doc, texto):
    return caixa(doc, "⚠️ Atenção:", texto, FUNDO_ALERTA, "E88B1A", COR_ALERTA)


def dica(doc, texto):
    return caixa(doc, "💡 Dica:", texto, FUNDO_DICA, "4A6FB5", RGBColor(0x2E, 0x4F, 0x8F))


def tabela(doc, cabecalho, linhas, larguras_cm=None, centralizar_ultima=False):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    linha_cab = t.rows[0]
    repetir_cabecalho(linha_cab)
    for i, texto in enumerate(cabecalho):
        cel = linha_cab.cells[i]
        sombrear_celula(cel, FUNDO_CABECALHO)
        cel.text = ""
        r = cel.paragraphs[0].add_run(texto)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for n, dados in enumerate(linhas):
        cels = t.add_row().cells
        for i, texto in enumerate(dados):
            cels[i].text = ""
            par = cels[i].paragraphs[0]
            r = par.add_run(texto)
            r.font.size = Pt(10.5)
            if i == 0:
                r.bold = True
            if centralizar_ultima and i == len(dados) - 1:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if n % 2 == 1:
                sombrear_celula(cels[i], "F4F7F7")
    if larguras_cm:
        for row in t.rows:
            for i, w in enumerate(larguras_cm):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def linhas_em_branco(doc, quantidade):
    for _ in range(quantidade):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(24)
        r = p.add_run("_" * 86)
        r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
        r.font.size = Pt(11)


def etapa(doc, numero, nome):
    titulo(doc, f"Etapa {numero} — {nome}", nivel=2)


# ------------------------------------------------------------
# Montagem do documento
# ------------------------------------------------------------
def configurar_documento():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = FONTE_TEXTO
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE_TEXTO)
    return doc


def rodape(doc):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Servidor Web com Node.js — Aula 2: Colocando o servidor no ar   |   Página ")
    r.font.size = Pt(9)
    r.font.color.rgb = COR_CINZA
    # Campo de número de página
    run = p.add_run()
    run.font.size = Pt(9)
    run.font.color.rgb = COR_CINZA
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tipo)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = texto
            run._r.append(instr)
    # Capa sem rodapé
    sec.different_first_page_header_footer = True



# ------------------------------------------------------------
# Códigos que os alunos vão digitar
# ------------------------------------------------------------
CODIGO_INDEX = [
    '<!DOCTYPE html>',
    '<html lang="pt-BR">',
    '<head>',
    '  <meta charset="UTF-8">',
    '  <title>Meu Servidor Web</title>',
    '  <style>',
    '    body { font-family: Arial, sans-serif; background: #0b5c5c; color: white;',
    '           text-align: center; padding-top: 80px; }',
    '    .caixa { background: rgba(0, 0, 0, 0.25); display: inline-block;',
    '             padding: 30px 50px; border-radius: 12px; }',
    '  </style>',
    '</head>',
    '<body>',
    '  <div class="caixa">',
    '    <h1>Meu primeiro servidor web</h1>',
    '    <p>Feito por: <strong>SEU NOME</strong></p>',
    '    <p>Curso de Redes de Computadores</p>',
    '    <p>Esta página está vindo do meu servidor Ubuntu, na porta 3000.</p>',
    '  </div>',
    '</body>',
    '</html>',
]

CODIGO_SERVER = [
    "// Módulos que já vêm com o Node.js",
    "const http = require('http');",
    "const fs = require('fs');",
    "const path = require('path');",
    "",
    "// Porta onde o servidor vai atender",
    "const PORTA = 3000;",
    "",
    "// O \"garçom\": roda a cada pedido que chega",
    "const servidor = http.createServer((pedido, resposta) => {",
    "  const hora = new Date().toLocaleTimeString('pt-BR');",
    "  console.log(`${hora} - ${pedido.socket.remoteAddress} pediu ${pedido.url}`);",
    "",
    "  const arquivo = path.join(__dirname, 'index.html');",
    "  fs.readFile(arquivo, (erro, conteudo) => {",
    "    if (erro) {",
    "      resposta.writeHead(500, { 'Content-Type': 'text/plain; charset=utf-8' });",
    "      resposta.end('Erro: não consegui ler o index.html');",
    "      return;",
    "    }",
    "    resposta.writeHead(200, { 'Content-Type': 'text/html; charset=utf-8' });",
    "    resposta.end(conteudo);",
    "  });",
    "});",
    "",
    "// Liga o servidor e espera pedidos de qualquer placa de rede",
    "servidor.listen(PORTA, '0.0.0.0', () => {",
    "  console.log(`Servidor rodando na porta ${PORTA}`);",
    "});",
]


def capa(doc):
    for _ in range(5):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("REDES DE COMPUTADORES · AULA PRÁTICA 2 DE 2")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = COR_DESTAQUE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("Seu primeiro servidor web\ncom Node.js")
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = COR_PRIMARIA

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("Parte 2 — Criando o site e colocando o servidor no ar")
    r.font.size = Pt(15)
    r.font.color.rgb = COR_CINZA

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("SSH · nano · index.html · server.js · PM2")
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = COR_CINZA

    for _ in range(6):
        doc.add_paragraph()

    # Campos de identificação
    t = doc.add_table(rows=4, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    campos = ["Nome do aluno:", "Turma:", "Data:", "Professor:"]
    for i, campo in enumerate(campos):
        c0, c1 = t.rows[i].cells
        c0.width = Cm(4.0)
        c1.width = Cm(12.0)
        sombrear_celula(c0, "F4F7F7")
        c0.text = ""
        r = c0.paragraphs[0].add_run(campo)
        r.bold = True
        r.font.color.rgb = COR_PRIMARIA
        c0.paragraphs[0].paragraph_format.space_before = Pt(8)
        c0.paragraphs[0].paragraph_format.space_after = Pt(8)
        c1.paragraphs[0].paragraph_format.space_before = Pt(8)
        c1.paragraphs[0].paragraph_format.space_after = Pt(8)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def conteudo(doc):
    # ---------------- Objetivo ----------------
    titulo(doc, "Objetivo da aula")
    paragrafo(doc, "Na Aula 1 você preparou o servidor. Hoje você vai criar a página do site (index.html) "
                   "e o programa que entrega essa página pela rede (server.js). No final, o site vai "
                   "abrir no navegador de outro computador e continuar no ar mesmo depois de a máquina "
                   "ser reiniciada.")

    # ---------------- Pré-requisitos ----------------
    titulo(doc, "Pré-requisitos")
    paragrafo(doc, "Tudo isto foi feito na Aula 1. Confira na Etapa 1.")
    marcador(doc, "Ubuntu Server instalado, com usuário e senha anotados.")
    marcador(doc, "Node.js 22, npm e PM2 instalados.")
    marcador(doc, "Pasta ~/meu-servidor-web criada.")
    marcador(doc, "Porta 3000 e SSH liberados no firewall (ufw).")
    marcador(doc, "Um segundo computador na mesma rede, com navegador e Prompt de Comando.")

    # ---------------- Conceito ----------------
    titulo(doc, "Conceito rápido: os dois arquivos de hoje")
    paragrafo(doc, "Continuando a analogia do restaurante da Aula 1:")
    tabela(doc,
           ["No restaurante", "No nosso projeto", "Escrito em"],
           [
               ["Cliente", "Navegador do outro computador", "—"],
               ["Garçom", "server.js — recebe o pedido e entrega a página", "JavaScript"],
               ["Prato", "index.html — a página que aparece na tela", "HTML"],
               ["Gerente de turno", "PM2 — mantém o garçom trabalhando", "—"],
           ],
           larguras_cm=[4.0, 9.0, 3.6])

    # ---------------- Etapas ----------------
    titulo(doc, "Passo a passo")
    p = paragrafo(doc, "Siga as etapas na ordem. Só avance quando o ")
    p.add_run("✅ Checkpoint").bold = True
    p.add_run(" da etapa der certo. O símbolo $ indica o início de um comando no servidor e > um "
              "comando no Windows — não digite esses símbolos.")

    # Etapa 1
    etapa(doc, 1, "Conferir o ambiente")
    paragrafo(doc, "Faça login no servidor e rode:")
    bloco_codigo(doc, ["$ hostname -I", "$ node -v", "$ pm2 -v", "$ sudo ufw status"])
    alerta(doc, "o IP pode ter mudado desde a última aula. Use sempre o IP de hoje.")
    paragrafo(doc, "IP do servidor hoje:  ______ . ______ . ______ . ______")
    checkpoint(doc, "IP anotado, node -v começa com v22, pm2 -v mostra uma versão e o ufw mostra 3000/tcp ALLOW. "
                    "Se algo falhar, refaça a etapa correspondente da apostila da Aula 1.")

    # Etapa 2
    etapa(doc, 2, "Entrar no servidor por SSH")
    paragrafo(doc, "Com o SSH você controla o servidor a partir de outro computador, pela rede. "
                   "A grande vantagem: dá para colar o código em vez de digitar tudo.")
    bloco_codigo(doc, "> ssh aluno@192.168.1.50",
                 "No PC do laboratório, abra o Prompt de Comando (tecla Windows, digite cmd) e rode:")
    paragrafo(doc, "Troque aluno pelo seu usuário e 192.168.1.50 pelo IP do seu servidor. Na primeira "
                   "conexão aparece a pergunta abaixo; digite yes e depois a sua senha.")
    bloco_codigo(doc, "Are you sure you want to continue connecting (yes/no/[fingerprint])? yes")
    dica(doc, "no Prompt de Comando do Windows, colar é com o botão direito do mouse (ou Ctrl + V).")
    checkpoint(doc, "o prompt da janela do Windows muda para aluno@servidor:~$.")

    # Etapa 3
    etapa(doc, 3, "Criar o index.html")
    bloco_codigo(doc, ["$ cd ~/meu-servidor-web", "$ nano index.html"], "Entrar na pasta e abrir o editor nano:")
    tabela(doc,
           ["Atalho no nano", "Faz o quê"],
           [
               ["Ctrl + O, depois Enter", "Salvar o arquivo"],
               ["Ctrl + X", "Sair do nano"],
               ["Ctrl + K", "Apagar a linha inteira"],
               ["Botão direito (no cmd)", "Colar"],
           ],
           larguras_cm=[6.0, 10.6])
    paragrafo(doc, "Digite (ou cole) o código abaixo. Troque SEU NOME pelo seu nome:")
    bloco_codigo(doc, CODIGO_INDEX, "Arquivo index.html", tamanho=9.5)
    tabela(doc,
           ["Marcação", "Significado"],
           [
               ["<h1> ... </h1>", "Título principal"],
               ["<p> ... </p>", "Parágrafo"],
               ["<strong> ... </strong>", "Texto em negrito"],
               ["<style> ... </style>", "Cores e aparência da página"],
               ["<meta charset=\"UTF-8\">", "Faz os acentos aparecerem certo"],
           ],
           larguras_cm=[6.0, 10.6])
    bloco_codigo(doc, "$ cat index.html", "Salve, saia e confira:")
    checkpoint(doc, "o comando cat index.html mostra o código que você digitou.")

    # Etapa 4
    etapa(doc, 4, "Criar o server.js")
    bloco_codigo(doc, "$ nano server.js")
    bloco_codigo(doc, CODIGO_SERVER, "Arquivo server.js", tamanho=9.5)
    alerta(doc, "cada parêntese, chave, aspas e ponto e vírgula importa. Um caractere faltando e o "
                "servidor não liga. Atenção às crases (`) nas linhas com console.log: no teclado ABNT2, "
                "a crase fica na tecla ao lado do P (com Shift), seguida de espaço.")
    paragrafo(doc, "O que cada parte faz:", negrito_inicio=None)
    tabela(doc,
           ["Trecho", "Em português"],
           [
               ["require('http')", "Pega a ferramenta de servidor web que já vem no Node.js"],
               ["const PORTA = 3000", "A \"mesa\" onde o garçom atende"],
               ["http.createServer(...)", "Cria o garçom. O que está dentro roda a cada pedido"],
               ["console.log(...)", "Escreve no terminal a hora, o IP de quem pediu e o que pediu"],
               ["fs.readFile(...)", "Lê o arquivo index.html do disco"],
               ["resposta.end(conteudo)", "Entrega a página para o navegador"],
               ["listen(PORTA, '0.0.0.0')", "Liga o servidor e aceita pedidos vindos da rede"],
           ],
           larguras_cm=[5.6, 11.0])
    bloco_codigo(doc, "$ ls", "Salve, saia e confira:")
    checkpoint(doc, "o comando ls mostra os dois arquivos: index.html e server.js.")

    # Etapa 5
    etapa(doc, 5, "Ligar o servidor pela primeira vez")
    bloco_codigo(doc, "$ node server.js")
    bloco_codigo(doc, "Servidor rodando na porta 3000", "Resposta esperada:")
    paragrafo(doc, "O terminal vai ficar \"parado\". É assim mesmo: o servidor está esperando pedidos. "
                   "Deixe essa janela aberta.")
    dica(doc, "se aparecer SyntaxError, a mensagem mostra o nome do arquivo e o número da linha com erro "
              "(ex.: server.js:12). Abra o arquivo no nano, corrija e rode de novo.")
    checkpoint(doc, "apareceu a mensagem Servidor rodando na porta 3000.")

    # Etapa 6
    etapa(doc, 6, "Abrir o site de outro computador")
    paragrafo(doc, "No navegador do outro computador, digite na barra de endereço (com o IP do seu servidor):")
    bloco_codigo(doc, "http://192.168.1.50:3000")
    alerta(doc, "escreva o http:// e o :3000 no final. Sem a porta, o navegador procura na porta 80, "
                "onde não há nada rodando.")
    paragrafo(doc, "Repare no terminal do servidor: cada acesso aparece no log, com o IP de quem pediu.")
    bloco_codigo(doc, [
        "19:42:10 - 192.168.1.61 pediu /",
        "19:42:10 - 192.168.1.61 pediu /favicon.ico",
    ], "Exemplo de log:")
    paragrafo(doc, "O pedido /favicon.ico é o navegador procurando o ícone da aba. É normal.",
              italico=True, cor=COR_CINZA)
    paragrafo(doc, "no terminal do servidor aperte Ctrl + C e atualize o navegador (F5).", negrito_inicio="Teste: ")
    paragrafo(doc, "O site sai do ar (\"recusou a conexão\"). Com node server.js o servidor só vive enquanto "
                   "o terminal está aberto. Na próxima etapa resolvemos isso.")
    checkpoint(doc, "a página abriu no navegador de outro PC e você viu o IP dele no log.")

    # Etapa 7
    etapa(doc, 7, "Colocar no ar com o PM2")
    bloco_codigo(doc, "$ pm2 start server.js --name meu-site", "Iniciar o site com um nome:")
    bloco_codigo(doc, "$ pm2 list", "Ver o que está rodando:")
    bloco_codigo(doc, "$ pm2 logs meu-site", "Ver os acessos ao vivo (Ctrl + C para sair dos logs):")
    paragrafo(doc, "Agora o terminal fica livre e o site continua no ar, mesmo se você fechar o SSH.")
    dica(doc, "sair dos logs com Ctrl + C não desliga o site. Só fecha a visualização dos logs.")
    checkpoint(doc, "pm2 list mostra meu-site com status online e o site abre no outro PC.")

    # Etapa 8
    etapa(doc, 8, "Alterar o site")
    paragrafo(doc, "", negrito_inicio="Mudou o index.html? ")
    doc.paragraphs[-1].add_run("Salve e aperte F5 no navegador. O server.js lê o arquivo a cada pedido, "
                               "então a mudança aparece na hora.")
    paragrafo(doc, "", negrito_inicio="Mudou o server.js? ")
    doc.paragraphs[-1].add_run("O código do servidor só é lido quando ele liga. Por isso precisa reiniciar:")
    bloco_codigo(doc, "$ pm2 restart meu-site")
    paragrafo(doc, "Pratique: abra o index.html no nano, troque a cor depois de background: "
                   "(ex.: #8e2de2, #c0392b, #1f6feb, darkgreen) e recarregue a página.")
    paragrafo(doc, "Cor que eu escolhi: ______________________")
    checkpoint(doc, "a nova cor apareceu no navegador do outro PC.")

    # Etapa 9
    etapa(doc, 9, "Ligar junto com a máquina")
    bloco_codigo(doc, "$ pm2 startup", "1. Gerar o comando de inicialização:")
    paragrafo(doc, "O PM2 responde com um comando começando com sudo, parecido com este:")
    bloco_codigo(doc, "sudo env PATH=$PATH:/usr/bin /usr/lib/node_modules/pm2/bin/pm2 startup systemd -u aluno --hp /home/aluno",
                 tamanho=9.5)
    paragrafo(doc, "2. Copie o comando que apareceu na SUA tela (selecione com o mouse), cole e aperte Enter.")
    bloco_codigo(doc, "$ pm2 save", "3. Salvar a lista de sites rodando:")
    bloco_codigo(doc, "$ sudo reboot", "4. Testar de verdade:")
    dica(doc, "depois do reboot a conexão SSH cai. Espere 1 a 2 minutos e recarregue o site no navegador.")
    checkpoint(doc, "depois de reiniciar, sem ninguém fazer login, o site abre no navegador do outro PC.")

    # ---------------- Comandos do PM2 ----------------
    titulo(doc, "Comandos do PM2")
    tabela(doc,
           ["Comando", "Para que serve"],
           [
               ["pm2 start server.js --name meu-site", "Liga o site com um nome"],
               ["pm2 list", "Mostra o que está rodando e o status"],
               ["pm2 logs meu-site", "Mostra os acessos e os erros ao vivo"],
               ["pm2 restart meu-site", "Reinicia (depois de mudar o server.js)"],
               ["pm2 stop meu-site", "Desliga o site"],
               ["pm2 startup  +  pm2 save", "Faz o site ligar junto com a máquina"],
           ],
           larguras_cm=[7.4, 9.2])

    # ---------------- Resumo ----------------
    titulo(doc, "Resumo: projeto no ar")
    tabela(doc,
           ["Item", "Como conferir", "Status"],
           [
               ["Acesso por SSH", "ssh usuario@IP funciona do PC do laboratório", "✅"],
               ["index.html criado", "cat index.html mostra o código", "✅"],
               ["server.js criado", "node server.js mostra \"Servidor rodando\"", "✅"],
               ["Site acessível pela rede", "http://IP:3000 abre em outro PC", "✅"],
               ["Site rodando no PM2", "pm2 list mostra meu-site online", "✅"],
               ["Alteração publicada", "nova cor aparece após F5", "✅"],
               ["Liga sozinho", "site volta depois de sudo reboot", "✅"],
           ],
           larguras_cm=[5.2, 9.4, 2.0], centralizar_ultima=True)

    # ---------------- Problemas comuns ----------------
    titulo(doc, "Problemas comuns")
    tabela(doc,
           ["Problema", "Causa provável", "Solução"],
           [
               ["ssh: \"Connection refused\" ou \"timed out\"",
                "IP errado, OpenSSH não instalado ou SSH bloqueado no ufw",
                "Conferir o IP; sudo apt install openssh-server; sudo ufw allow OpenSSH"],
               ["SyntaxError ao rodar node server.js",
                "Erro de digitação no código",
                "Ver o número da linha na mensagem e corrigir no nano"],
               ["\"Cannot find module\"",
                "Você não está na pasta do projeto ou o nome do arquivo está errado",
                "cd ~/meu-servidor-web e conferir com ls"],
               ["\"EADDRINUSE\" (porta em uso)",
                "Já existe um servidor rodando na porta 3000",
                "Fechar o node server.js (Ctrl + C) ou pm2 stop meu-site"],
               ["Navegador: \"não é possível acessar\"",
                "Servidor parado, IP errado ou faltou o :3000",
                "pm2 list; conferir IP; usar http://IP:3000"],
               ["Navegador demora e dá tempo esgotado",
                "Porta 3000 bloqueada no firewall",
                "sudo ufw allow 3000/tcp"],
               ["Acentos aparecem como sÃ­mbolos",
                "Faltou o charset (UTF-8)",
                "Conferir o <meta charset=\"UTF-8\"> e o charset=utf-8 no server.js"],
               ["Página mostra \"Erro: não consegui ler o index.html\"",
                "O index.html não está na mesma pasta ou está com outro nome",
                "ls na pasta; o nome deve ser index.html, tudo minúsculo"],
               ["pm2 list mostra status errored",
                "O server.js tem erro",
                "pm2 logs meu-site para ver o erro; corrigir; pm2 restart meu-site"],
               ["Após o reboot o site não volta",
                "Faltou rodar o comando do pm2 startup ou o pm2 save",
                "Refazer a Etapa 9 com o site online"],
           ],
           larguras_cm=[5.2, 5.2, 6.2])

    # ---------------- Glossário ----------------
    titulo(doc, "Glossário rápido")
    tabela(doc,
           ["Termo", "Significado"],
           [
               ["SSH", "Acesso remoto seguro ao terminal de outra máquina pela rede"],
               ["nano", "Editor de texto simples que funciona dentro do terminal"],
               ["HTML", "Linguagem que descreve o conteúdo de uma página (títulos, textos, cores)"],
               ["JavaScript", "Linguagem de programação; o Node.js roda JavaScript no servidor"],
               ["Módulo", "Pacote de funções pronto que o código usa (ex.: http, fs)"],
               ["Log", "Registro do que aconteceu, linha por linha (ex.: quem acessou e quando)"],
               ["0.0.0.0", "\"Todas as placas de rede\": aceita pedidos vindos de qualquer lugar da rede"],
               ["URL", "Endereço digitado no navegador (ex.: http://192.168.1.50:3000)"],
               ["PM2", "Gerenciador que mantém o servidor ligado e o religa se cair"],
           ],
           larguras_cm=[3.5, 13.1])

    # ---------------- Desafios ----------------
    titulo(doc, "Desafios extras (para quem terminar antes)")
    paragrafo(doc, "", negrito_inicio="1. Volta na turma. ")
    doc.paragraphs[-1].add_run("Abra o site de três colegas e preencha a tabela.")
    tabela(doc,
           ["Nome do colega", "IP do servidor", "Abriu? (sim/não)", "Cor do site"],
           [["", "", "", ""] for _ in range(3)],
           larguras_cm=[5.6, 4.2, 3.4, 3.4])
    paragrafo(doc, "", negrito_inicio="2. Detetive de IP. ")
    doc.paragraphs[-1].add_run("Rode pm2 logs meu-site e descubra quem acessou o seu site. "
                               "IPs que apareceram: __________________________________________")
    paragrafo(doc, "", negrito_inicio="3. Site com a sua cara. ")
    doc.paragraphs[-1].add_run("Adicione um parágrafo sobre você e uma lista com 3 comandos que aprendeu "
                               "(use <ul> para a lista e <li> para cada item).")
    paragrafo(doc, "", negrito_inicio="4. Quebre de propósito. ")
    doc.paragraphs[-1].add_run("Rode pm2 stop meu-site, veja a mensagem no navegador e anote: "
                               "__________________________. Depois ligue de novo com pm2 start meu-site.")

    # ---------------- Anotações ----------------
    titulo(doc, "Anotações")
    linhas_em_branco(doc, 10)

    # ---------------- Avaliação ----------------
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    titulo(doc, "Avaliação rápida")
    perguntas = [
        "Qual a função do index.html e qual a função do server.js no projeto?",
        "Qual endereço você digitou no navegador para abrir o seu site? Por que é preciso colocar o :3000?",
        "O que acontece com o site quando você aperta Ctrl + C no terminal onde rodou node server.js? Por quê?",
        "Para que serve o PM2? Cite dois comandos dele e o que cada um faz.",
        "Você mudou o index.html e o server.js. Em qual dos dois foi preciso rodar pm2 restart? Explique.",
    ]
    for i, pergunta in enumerate(perguntas, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{i}. ")
        r.bold = True
        r.font.color.rgb = COR_PRIMARIA
        p.add_run(pergunta)
        linhas_em_branco(doc, 3)


def main():
    doc = configurar_documento()
    rodape(doc)
    capa(doc)
    conteudo(doc)
    doc.save(ARQUIVO_SAIDA)
    print(f"✅ Arquivo gerado: {os.path.abspath(ARQUIVO_SAIDA)}")


if __name__ == "__main__":
    main()
