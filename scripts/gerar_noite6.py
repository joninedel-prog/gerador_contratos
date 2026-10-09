"""Gera os 4 arquivos da Noite 6 do curso de Redes.

Noite 6 — Endereço na rede: IP, MAC e como o aparelho é encontrado
Padrão prática-pesada. Parte da tarefa da Noite 5 (IP e MAC anotados do celular),
explica MAC e IP com a analogia do condomínio, e os alunos entram no painel do
roteador da equipe pra ver quem está na rede, caçam um intruso, usam ipconfig e ping,
e reconhecem o "sem endereço" (169.254) como diagnóstico de campo.
Não precisa de internet: a rede local da bancada funciona sozinha (liga com a Noite 1).
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from common import (
    COR_PRIMARIA, COR_ALERTA, COR_DESTAQUE, COR_CINZA,
    add_heading, add_para, add_bullet, add_table, add_hr,
    setup_page, cabecalho_documento,
    PPalette, new_presentation, add_slide, add_text, add_title_block,
    add_slide_cover, add_footer, add_table_slide, add_bar,
)
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PInches

OUTDIR = "/home/nedel/curso"
NOITE = "Noite 6"
FOOTER = f"Redes de Computadores · {NOITE}"


# A analogia do condomínio — usada no roteiro, nos slides, na folha e no resumo
CONDOMINIO = [
    ("Condomínio", "A rede da casa (tudo que está ligado no mesmo roteador)"),
    ("Portaria", "O roteador — por onde tudo entra e sai. É o \"gateway\""),
    ("Porteiro que distribui os apartamentos", "O DHCP — a função do roteador que dá endereço pra quem chega"),
    ("Número do apartamento", "O endereço IP — muda se o morador sai e volta depois de muito tempo"),
    ("CPF do morador", "O endereço MAC — vem de fábrica, é da placa de rede"),
    ("Endereço do condomínio na rua", "O IP de fora (público) — é um só pra todos os apartamentos"),
]


# Três telas de ipconfig pra leitura (folha p.3 e gabarito do roteiro)
IPCONFIG = [
    ("Tela A",
     ["Adaptador Ethernet Ethernet:",
      "   Endereço IPv4. . . . . . . . . . . : 192.168.0.104",
      "   Máscara de Sub-rede . . . . . . . . : 255.255.255.0",
      "   Gateway Padrão. . . . . . . . . . . : 192.168.0.1"],
     "Normal. Recebeu endereço do roteador (192.168.0.x) e sabe quem é a portaria (192.168.0.1)."),
    ("Tela B",
     ["Adaptador Ethernet Ethernet:",
      "   Endereço IPv4 de Configuração Automática: 169.254.37.112",
      "   Máscara de Sub-rede . . . . . . . . : 255.255.0.0",
      "   Gateway Padrão. . . . . . . . . . . :"],
     "Sem endereço. O cabo conectou, mas ninguém deu IP: roteador desligado/travado, DHCP desligado ou cabo ligado no lugar errado. O computador inventou um 169.254 sozinho."),
    ("Tela C",
     ["Adaptador Ethernet Ethernet:",
      "   Estado da mídia. . . . . . . . . . : mídia desconectada"],
     "Nem conectou. Problema físico: cabo solto, cabo ruim, porta apagada. Volta pro método da Noite 5."),
]


# ============================================================
# ROTEIRO
# ============================================================
def build_roteiro():
    doc = Document()
    setup_page(doc)

    cabecalho_documento(
        doc,
        "ENDEREÇO NA REDE: IP, MAC E COMO O APARELHO É ENCONTRADO",
        f"Roteiro do professor · Redes de Computadores · {NOITE} · Curso FIC · CEJA Itapiranga – SC",
    )

    # A IDEIA
    add_heading(doc, "A IDEIA DESTA NOITE")
    add_para(doc, "Até agora a turma mexeu no que dá pra pegar na mão: equipamento e cabo. Hoje entra a primeira coisa invisível — o endereço. Sem ele, o roteador não sabe pra quem entregar o vídeo que chegou da internet. E é no endereço que aparecem defeitos que nenhum testador de cabo pega: o \"conectado, sem internet\", a impressora que sumiu, o aparelho estranho na rede do cliente.")
    add_para(doc, "A frase da noite: \"Todo aparelho na rede tem dois documentos: o CPF (MAC) e o número do apartamento (IP).\"", bold=True, color=COR_PRIMARIA)

    add_heading(doc, "O que se resolve nesta noite", level=2)
    add_para(doc, "Cada aluno sai sabendo achar o IP e o MAC de um celular e de um computador, entrar no painel do roteador e identificar quem está na rede, usar o ping pra saber se um aparelho responde, e reconhecer o 169.254 como sinal de \"o aparelho não recebeu endereço\".")

    add_heading(doc, "Esta é a sexta de quinze noites", level=2)
    add_para(doc, "Abre o bloco do \"invisível\" (Noites 6 e 7). O ping já apareceu na Noite 5 como ferramenta. Hoje ele ganha sentido. A Noite 7 completa com nomes (DNS): \"o IP é o número; o DNS é a agenda que transforma nome em número\". A Noite 9 volta ao painel do roteador pra configurar de verdade — hoje é só OLHAR, nada de mudar configuração.")
    add_para(doc, "Não entre em máscara de sub-rede, classes de IP nem conta binária. Pra técnico de campo, basta: \"os três primeiros números iguais = mesma rede, na maioria das casas\".", italic=True, color=COR_ALERTA)

    add_hr(doc)

    # ANTES DA AULA
    add_heading(doc, "ANTES DA AULA")

    add_heading(doc, "Material da mesa", level=2)
    add_bullet(doc, "Um roteador por equipe de 3 ou 4 (os da bancada), cada um na tomada. NÃO precisa de internet")
    add_bullet(doc, "Um roteador extra, \"sabotado\", para a Estação B do Bloco 7 (ver abaixo)")
    add_bullet(doc, "1 ou 2 computadores/notebooks com porta de rede + cabos testados da Noite 5")
    add_bullet(doc, "Celular do professor (vai ser o \"intruso\")")
    add_bullet(doc, "Fita crepe e caneta pra etiquetar os roteadores: EQUIPE 1, EQUIPE 2…")
    add_bullet(doc, "Régua de tomadas")

    add_heading(doc, "Prepare os roteadores das equipes (em casa, antes)", level=2)
    add_bullet(doc, "Restaure cada um para o padrão de fábrica (botão Reset, ~10 segundos com o aparelho ligado)")
    add_bullet(doc, "Entre no painel (http://192.168.0.1 ou tplinkwifi.net) e crie a senha de administrador. Use a mesma em todos e anote num papel que fica com você")
    add_bullet(doc, "Troque o nome do Wi-Fi pra EQUIPE-1, EQUIPE-2… e a senha do Wi-Fi pra algo simples (ex.: redes2026). Escreva o nome e a senha do Wi-Fi na fita crepe colada no roteador")
    add_bullet(doc, "Descubra, em cada modelo, onde fica a lista de aparelhos conectados. O nome muda: \"Clientes DHCP\", \"Lista de clientes DHCP\", \"Dispositivos conectados\", \"Mapa da rede → Clientes\" ou \"Estatísticas wireless\". Anote o caminho na fita do roteador")
    add_para(doc, "Os roteadores vão ficar sem internet e o celular vai reclamar (\"Wi-Fi sem internet\"). É de propósito. Diga à turma: \"Lembram da Noite 1? Rede não é internet. A rede funciona sem ela.\"", italic=True, color=COR_PRIMARIA)

    add_heading(doc, "Prepare o roteador sabotado (Estação B)", level=2)
    add_bullet(doc, "Num roteador extra, no painel, DESATIVE o servidor DHCP (fica em Rede/LAN → DHCP, ou Avançado → Rede → Servidor DHCP)")
    add_bullet(doc, "Etiquete só como \"ESTAÇÃO B\" — não escreva que está sabotado")
    add_bullet(doc, "Teste antes: ligue o computador nele por cabo e rode ipconfig. Tem que aparecer 169.254.x.x. Se aparecer, está pronto")
    add_para(doc, "Se não conseguir desativar o DHCP: ligue o computador por cabo direto noutro computador (sem roteador). Os dois vão pegar 169.254 — o efeito é o mesmo.", italic=True, color=COR_CINZA)

    add_heading(doc, "Deixe o computador pronto", level=2)
    add_bullet(doc, "Wi-Fi do computador DESLIGADO (como na Noite 5) — senão ele mostra dois endereços e confunde")
    add_bullet(doc, "Prompt de Comando fixado na barra de tarefas")
    add_bullet(doc, "Teste: ipconfig e ipconfig /all. Ache as linhas \"Endereço IPv4\", \"Gateway Padrão\" e \"Endereço Físico\" (é o MAC)")
    add_bullet(doc, "O Windows pode bloquear o ping vindo de fora (firewall). Não importa: hoje o computador é quem PINGA, não quem é pingado")

    add_heading(doc, "Imprima", level=2)
    add_bullet(doc, "Folha do aluno (3 páginas), uma por pessoa")
    add_bullet(doc, "Resumo do aluno, uma por pessoa (pra levar pra casa)")

    add_hr(doc)

    # QUADRO DE TEMPO
    add_heading(doc, "QUADRO DE TEMPO — 210 MINUTOS")
    add_table(doc, [
        ["Bloco", "Tempo", "O que acontece"],
        ["1. Retomada + censo dos endereços", "15 min", "Cada um escreve no quadro o IP da tarefa. A turma acha o padrão"],
        ["2. Dois documentos: MAC e IP", "20 min", "Analogia do condomínio"],
        ["3. PRÁTICA — Ache seus endereços", "20 min", "Celular na rede da equipe + ipconfig no computador"],
        ["4. Quem dá o endereço + IP de dentro e de fora", "15 min", "DHCP, gateway, IP público"],
        ["INTERVALO", "15 min", ""],
        ["5. PRÁTICA — Quem está na rede?", "35 min", "Equipes entram no painel do roteador e identificam cada aparelho"],
        ["6. PRÁTICA — Caça ao intruso", "15 min", "O professor entra escondido numa rede. A equipe tem que achar"],
        ["7. PRÁTICA — Ping e o sem-endereço", "25 min", "Estação A: ping. Estação B: o 169.254. Fila: ler telas de ipconfig"],
        ["8. Chamados em role-play", "20 min", "Conectado sem internet, impressora sumida, vizinho no Wi-Fi"],
        ["9. Duelo relâmpago", "15 min", "Time A × Time B"],
        ["10. Fechamento + tarefa", "15 min", "Quiz de saída e anúncio da Noite 7"],
    ], col_widths=[6.5, 2, 8])
    add_para(doc, "Blocos 3, 5, 6, 7, 8 e 9: 130 min de atividade do aluno (62% do total), mais o censo do Bloco 1. A exposição está em dois blocos curtos (35 min somados), separados por prática.", italic=True, color=COR_CINZA)

    add_hr(doc)

    # BLOCO 1
    add_heading(doc, "BLOCO 1 — RETOMADA + CENSO DOS ENDEREÇOS · 15 MIN")
    add_heading(doc, "Recap relâmpago — 5 min", level=2)
    add_bullet(doc, "Quais os três jeitos de testar um cabo? (testador, luz, velocidade)")
    add_bullet(doc, "Luz acesa garante cabo bom? (não)")
    add_bullet(doc, "Quais os 4 passos do método? (olhar, testar, isolar, resolver)")
    add_bullet(doc, "No teste do balanço a gente pingou um número. Alguém lembra qual? (192.168.0.1)")

    add_heading(doc, "O censo — 10 min", level=2)
    add_para(doc, "Divida o quadro em duas colunas: IP e MAC. Cada aluno que fez a tarefa vai lá e escreve os números que anotou do celular. Quem não fez, pega o celular agora (Configurações → Wi-Fi → toque na rede → detalhes).")
    add_para(doc, "Depois pergunte à turma, sem explicar nada ainda:")
    add_bullet(doc, "\"O que esses IPs têm em comum?\" (quase todos começam com 192.168…)")
    add_bullet(doc, "\"Tem dois iguais?\" (pode ter! — 192.168.0.100 em duas casas diferentes. Guarde isso pro Bloco 2)")
    add_bullet(doc, "\"E os MACs, tem algum igual?\" (não pode ter)")
    add_bullet(doc, "\"Quem anotou o IP no 4G, fora do Wi-Fi?\" (vai ser bem diferente — às vezes começa com 10. ou 100.)")
    add_para(doc, "\"Dois colegas com o mesmo IP, e nenhum com o mesmo MAC. Guardem essa esquisitice. Em 20 minutos ela vai fazer sentido.\"", italic=True, color=COR_PRIMARIA)

    add_hr(doc)

    # BLOCO 2
    add_heading(doc, "BLOCO 2 — DOIS DOCUMENTOS: MAC E IP · 20 MIN")
    add_heading(doc, "A pergunta — 3 min", level=2)
    add_para(doc, "\"Tem vinte celulares no mesmo Wi-Fi. Chega um vídeo da internet. Como o roteador sabe que é pro celular do Pedro e não pro da Ana?\" Deixe a turma chutar. Aceite tudo. A resposta é: endereço.", italic=True)

    add_heading(doc, "O condomínio — 12 min", level=2)
    add_para(doc, "Desenhe no quadro um prédio com portaria e apartamentos. Construa a tabela com a turma:")
    add_table(doc, [["No condomínio…", "Na rede…"]] + [list(r) for r in CONDOMINIO], col_widths=[6, 10.5])
    add_para(doc, "Os pontos que a turma precisa sair falando:")
    add_bullet(doc, "MAC é o CPF: vem gravado de fábrica na placa de rede. Cada placa tem um. Lembra da etiqueta do roteador na Noite 2? Tinha \"MAC: 60-32-B1-…\". Mostre uma etiqueta de novo")
    add_bullet(doc, "IP é o número do apartamento: quem dá é o porteiro (o roteador), na hora que o aparelho chega. Sai e volta depois de dias? Pode ganhar outro número")
    add_bullet(doc, "Por isso dois colegas tinham o mesmo IP: são condomínios diferentes. Apartamento 101 existe em todo prédio. CPF repetido não existe em lugar nenhum")
    add_bullet(doc, "Como se lê um IP: são 4 números de 0 a 255 com ponto. Em casa, os três primeiros dizem a rua (a rede) e o último diz a casa (o aparelho). 192.168.0.104 e 192.168.0.105 são vizinhos. 192.168.1.104 é outra rua")
    add_bullet(doc, "Como se lê um MAC: 6 pares de letras e números (de 0 a 9 e de A a F), separados por traço ou dois pontos. Ex.: 3C-52-82-1A-7F-09")

    add_heading(doc, "O aviso de campo — 5 min", level=2)
    add_para(doc, "\"Agora uma pegadinha da vida real: celular novo NÃO mostra o CPF verdadeiro. Ele inventa um MAC diferente pra cada Wi-Fi, por privacidade. Chama MAC aleatório (Android) ou endereço privado (iPhone). Então o MAC que vocês anotaram em casa pode não ser o mesmo que o roteador da escola vai ver hoje. Vamos conferir daqui a pouco.\"", italic=True, color=COR_ALERTA)
    add_para(doc, "Onde dilatar: pergunte \"por que o celular faria isso?\" e deixe a turma descobrir (pra loja, shopping e aeroporto não ficarem rastreando por onde você passa).", italic=True, color=COR_DESTAQUE)

    add_hr(doc)

    # BLOCO 3
    add_heading(doc, "BLOCO 3 — PRÁTICA · ACHE SEUS ENDEREÇOS · 20 MIN")
    add_para(doc, "Turma em equipes de 3 ou 4. Cada equipe recebe um roteador (EQUIPE-1, EQUIPE-2…). Página 1 da folha.")
    add_heading(doc, "No celular — 12 min", level=2)
    add_bullet(doc, "Desligue os dados móveis (senão o celular foge pro 4G quando vê \"sem internet\")")
    add_bullet(doc, "Conecte no Wi-Fi da sua equipe. Se aparecer \"Sem internet — conectar mesmo assim?\", diga que sim")
    add_bullet(doc, "Android: Configurações → Wi-Fi → engrenagem ao lado da rede → Endereço IP. O MAC fica em \"Avançado\" ou \"Privacidade\"")
    add_bullet(doc, "iPhone: Ajustes → Wi-Fi → (i) ao lado da rede → Endereço IP e Endereço Wi-Fi")
    add_bullet(doc, "Anote IP, MAC e se o MAC é aleatório/privado. Compare com o que anotou em casa")
    add_heading(doc, "No computador — 8 min (estação, a turma roda)", level=2)
    add_bullet(doc, "Ligue o cabo no roteador de uma equipe. Prompt de Comando → ipconfig")
    add_bullet(doc, "Ache o Endereço IPv4 e o Gateway Padrão")
    add_bullet(doc, "Depois ipconfig /all → Endereço Físico (é o MAC da placa do computador)")
    add_para(doc, "Pergunta pra fechar: \"O IP do computador e o do seu celular começam igual? Por quê?\" (mesmo roteador = mesmo condomínio = mesma rua)", italic=True, color=COR_PRIMARIA)
    add_para(doc, "Onde dilatar: se o MAC de alguém NÃO mudou (aleatório desligado), compare com quem mudou. Mostre no celular de um aluno onde se liga e desliga isso.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)

    # BLOCO 4
    add_heading(doc, "BLOCO 4 — QUEM DÁ O ENDEREÇO + IP DE DENTRO E DE FORA · 15 MIN")
    add_heading(doc, "DHCP: o porteiro — 5 min", level=2)
    add_bullet(doc, "Quando o celular chega, ele grita: \"alguém me dá um endereço?\" O roteador responde com: um IP, quem é a portaria (gateway) e por quanto tempo vale")
    add_bullet(doc, "Essa função do roteador se chama DHCP. Ninguém precisa digitar nada. Por isso conectar no Wi-Fi \"só funciona\"")
    add_bullet(doc, "Se o porteiro não responde, o aparelho fica sem endereço. O Windows, aí, inventa um número que começa com 169.254. \"Quando vocês virem 169.254, já sabem: ninguém deu endereço pra esse aparelho\"")
    add_heading(doc, "IP de dentro e IP de fora — 10 min", level=2)
    add_bullet(doc, "O 192.168… só vale dentro de casa. Lá fora, na internet, a casa inteira aparece com um endereço só: o do roteador do lado de fora, que o provedor dá")
    add_bullet(doc, "É o endereço do condomínio na rua. O carteiro não sabe o número do apartamento — ele entrega na portaria, e a portaria distribui")
    add_bullet(doc, "Demonstração (se a escola tiver Wi-Fi com internet): dois alunos no Wi-Fi da escola pesquisam \"qual meu ip\" no Google. Aparece o MESMO número pros dois. Agora um passa pro 4G: número totalmente diferente")
    add_para(doc, "Ligue com a Noite 2: \"O provedor que dá esse IP de fora. Às vezes ele dá um IP compartilhado com outros clientes — isso vai importar quando a gente falar de câmera, na Noite 13.\"", italic=True, color=COR_CINZA)

    add_hr(doc)
    add_heading(doc, "INTERVALO · 15 MIN", color=COR_CINZA)
    add_hr(doc)

    # BLOCO 5
    add_heading(doc, "BLOCO 5 — PRÁTICA · QUEM ESTÁ NA REDE? · 35 MIN")
    add_para(doc, "A atividade principal da noite. Cada equipe entra no painel do SEU roteador e descobre quem é cada aparelho conectado. Página 2 da folha.")
    add_heading(doc, "Passo a passo (escreva no quadro)", level=2)
    add_bullet(doc, "Celular conectado no Wi-Fi da equipe. Abra o navegador e digite: 192.168.0.1 (é o gateway que vocês anotaram!)")
    add_bullet(doc, "Senha de administrador: você fala em voz alta ou passa de equipe em equipe")
    add_bullet(doc, "Ache a lista de aparelhos (o caminho está na fita do roteador)")
    add_bullet(doc, "Preencha a tabela da folha: nome que aparece, IP, MAC — e quem é (de qual colega é cada aparelho)")
    add_bullet(doc, "Cada um confere: o MAC que o roteador mostra bate com o que você anotou no Bloco 3? E com o da tarefa de casa?")
    add_heading(doc, "O que vai aparecer — e o que dizer", level=2)
    add_bullet(doc, "Aparelho com nome esquisito (\"Galaxy-A15\", \"android-3f2a…\", \"iPhone\", ou só \"Desconhecido\"). O nome é o aparelho que diz, não é confiável. O MAC e o IP é que identificam")
    add_bullet(doc, "Dois celulares iguais com o mesmo nome: só o MAC separa")
    add_bullet(doc, "O computador aparece com cabo, os celulares pelo Wi-Fi — alguns painéis mostram separado")
    add_bullet(doc, "Um MAC na lista que bate com a etiqueta do roteador? Não, o roteador não aparece na própria lista. Mas a etiqueta embaixo dele mostra o MAC dele. Peça pra conferirem")
    add_para(doc, "\"Isto aqui é o que vocês vão fazer no cliente que diz 'acho que o vizinho tá usando meu Wi-Fi'. Entra no painel, conta os aparelhos da casa, compara com a lista. Sobrou aparelho? Tem alguém de fora.\"", italic=True, color=COR_PRIMARIA)
    add_para(doc, "REGRA: hoje não se muda NADA no painel. Quem clicar em \"Bloquear\" ou \"Salvar\" sem querer, avisa. A Noite 9 é a de configurar.", bold=True, color=COR_ALERTA)
    add_para(doc, "Onde dilatar: equipe que terminar rápido desliga e religa o Wi-Fi do celular e confere se o IP mudou (geralmente não — o roteador lembra pelo MAC por um tempo). Depois \"esquece a rede\" e conecta de novo: com MAC aleatório, pode ganhar outro IP.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)

    # BLOCO 6
    add_heading(doc, "BLOCO 6 — PRÁTICA · CAÇA AO INTRUSO · 15 MIN")
    add_para(doc, "Durante o Bloco 5, sem ninguém ver, conecte o seu celular (o \"intruso\") no Wi-Fi de uma equipe — a senha está na fita. Se der, conecte também um aparelho de outra equipe noutra rede.")
    add_bullet(doc, "Anuncie: \"Uma das redes foi invadida. Cada equipe confere a sua lista e diz se tem intruso.\"")
    add_bullet(doc, "A equipe tem que dizer: qual aparelho é o intruso, o IP e o MAC dele, e como chegou nessa conclusão")
    add_bullet(doc, "A conclusão certa é por eliminação: contaram os aparelhos da equipe, identificaram cada um, sobrou um")
    add_bullet(doc, "Mostre que o nome do intruso podia ser qualquer coisa. No seu celular, mude o nome do aparelho pra \"Impressora-HP\" antes — fica mais divertido")
    add_para(doc, "\"Como o intruso entrou? Sabia a senha — estava escrita na fita. Na casa do cliente é igual: senha na geladeira, senha que o filho passou pro amigo. A gente resolve isso na Noite 11.\"", italic=True, color=COR_PRIMARIA)

    add_hr(doc)

    # BLOCO 7
    add_heading(doc, "BLOCO 7 — PRÁTICA · PING E O SEM-ENDEREÇO · 25 MIN")
    add_para(doc, "Duas estações e uma atividade de papel. As equipes rodam: ~8 min por estação, e quem espera faz a leitura de telas (página 3 da folha).")

    add_heading(doc, "Estação A — Quem responde? (computador no roteador de uma equipe)", level=2)
    add_bullet(doc, "ping 192.168.0.1 — a portaria responde? (sim, tempo baixinho, tipo 1 ms)")
    add_bullet(doc, "ping no IP do celular de um colega da equipe — responde? (Android em geral sim; iPhone às vezes não. Não responder não é defeito — tem aparelho que ignora ping por segurança)")
    add_bullet(doc, "O colega desliga o Wi-Fi do celular e pinga de novo: \"Host de destino inacessível\" ou \"Esgotado o tempo limite\"")
    add_bullet(doc, "ping num IP que ninguém tem (ex.: 192.168.0.250) — mesma resposta")
    add_para(doc, "\"Ping é bater na porta. Respondeu: o aparelho está lá e o caminho até ele funciona. Não respondeu: ou ele não está, ou o caminho está quebrado, ou ele não abre a porta pra ninguém.\"", italic=True, color=COR_PRIMARIA)

    add_heading(doc, "Estação B — O roteador mudo (o sabotado)", level=2)
    add_bullet(doc, "Liga o computador por cabo no roteador da Estação B. Luz da porta acende. O Windows mostra \"Sem internet\" ou \"Rede não identificada\"")
    add_bullet(doc, "ipconfig → aparece 169.254.x.x e Gateway Padrão vazio")
    add_bullet(doc, "Tenta conectar o celular no Wi-Fi dele: fica em \"Obtendo endereço IP…\" e depois \"Falha ao obter endereço IP\"")
    add_bullet(doc, "A equipe escreve o diagnóstico. Depois você conta o segredo: o porteiro (DHCP) estava de folga")
    add_para(doc, "\"Cabo bom, luz acesa, e não funciona. Se vocês só soubessem testar cabo, iam trocar o cabo à toa. Agora vocês sabem olhar o endereço.\"", italic=True, color=COR_PRIMARIA)

    add_heading(doc, "Fila — Lendo telas de ipconfig (papel)", level=2)
    add_para(doc, "Três telas na folha. A equipe escreve o que está acontecendo em cada uma. Gabarito:")
    for nome, _, resp in IPCONFIG:
        add_para(doc, f"{nome}: {resp}", italic=True, color=COR_CINZA)

    add_hr(doc)

    # BLOCO 8
    add_heading(doc, "BLOCO 8 — CHAMADOS EM ROLE-PLAY · 20 MIN")
    add_para(doc, "Duplas: um cliente, um técnico. O técnico diz o que faria, passo a passo, usando o que viu hoje. Troca os papéis a cada chamado. Chamados na página 3 da folha:")
    chamados = [
        ("\"O computador diz 'conectado, sem internet'. O cabo é novo, a luz acende.\"",
         "ipconfig. Se 169.254: ninguém deu endereço — reinicia o roteador, confere se o cabo está numa porta LAN (não na WAN). Se IP normal: ping no gateway e segue o diagnóstico."),
        ("\"Faltou luz ontem e agora a impressora sumiu do computador.\"",
         "O roteador reiniciou e a impressora ganhou outro IP; o computador procura no endereço antigo. Achar o IP novo no painel (pelo MAC). Solução definitiva: reserva de IP (Noite 9)."),
        ("\"Acho que o vizinho está usando meu Wi-Fi. A internet anda lenta à noite.\"",
         "Painel do roteador → lista de aparelhos. Contar os aparelhos da casa com o cliente, identificar cada um. Sobrou? Trocar a senha do Wi-Fi (Noite 11)."),
        ("\"Meu filho disse que o IP dele 'vazou' no jogo. É perigoso?\"",
         "O que aparece no jogo é o IP de fora (do condomínio), não o da casa. Muda quando o roteador reinicia na maioria dos provedores. Calma, não é senha."),
    ]
    for ch, resp in chamados:
        add_para(doc, ch, bold=True)
        add_para(doc, "> " + resp, italic=True, color=COR_CINZA)

    add_hr(doc)

    # BLOCO 9
    add_heading(doc, "BLOCO 9 — DUELO RELÂMPAGO · 15 MIN")
    add_para(doc, "Turma em dois times. Resposta em 5 segundos. Alterna.")
    add_heading(doc, "14 perguntas de reserva", level=2)
    add_bullet(doc, "Qual dos dois vem de fábrica: IP ou MAC? — MAC")
    add_bullet(doc, "Quem dá o IP pro celular em casa? — O roteador (DHCP)")
    add_bullet(doc, "Na analogia, o MAC é o quê? — O CPF")
    add_bullet(doc, "E o IP? — O número do apartamento")
    add_bullet(doc, "E o gateway? — A portaria (o roteador)")
    add_bullet(doc, "192.168.0.10 e 192.168.0.200 estão na mesma rede? — Sim (mesmos três primeiros)")
    add_bullet(doc, "192.168.0.10 e 192.168.1.10? — Não, na maioria das casas")
    add_bullet(doc, "Computador com 169.254 quer dizer o quê? — Ninguém deu endereço pra ele")
    add_bullet(doc, "\"Mídia desconectada\" no ipconfig. Problema de quê? — Físico: cabo ou porta")
    add_bullet(doc, "Dois aparelhos na mesma casa podem ter o mesmo IP? — Não (dá conflito)")
    add_bullet(doc, "E em casas diferentes? — Pode")
    add_bullet(doc, "Pra que serve o ping? — Saber se um aparelho responde")
    add_bullet(doc, "Por que o celular usa MAC aleatório? — Privacidade, pra não ser rastreado")
    add_bullet(doc, "Onde o técnico vê quem está no Wi-Fi do cliente? — No painel do roteador")

    add_hr(doc)

    # BLOCO 10
    add_heading(doc, "BLOCO 10 — FECHAMENTO E TAREFA · 15 MIN")
    add_heading(doc, "Quiz de saída", level=2)
    add_bullet(doc, "Diferença entre IP e MAC, numa frase cada")
    add_bullet(doc, "O que é o gateway e qual o número dele na rede de hoje?")
    add_bullet(doc, "O que o 169.254 conta pro técnico?")
    add_bullet(doc, "Como você descobre se tem aparelho estranho no Wi-Fi do cliente?")

    add_heading(doc, "A tarefa", level=2)
    add_para(doc, "\"Hoje a gente digitou 192.168.0.1 no navegador. Mas ninguém digita 142.250.… pra entrar no Google — digita google.com.br. Quem transforma o nome em número? Isso é a Noite 7. A tarefa: no celular, abra três sites que você usa e anote o endereço completo que aparece na barra lá em cima — com o https:// e tudo. Repare se aparece um cadeado.\"", italic=True)

    add_heading(doc, "Anuncie a próxima", level=2)
    add_para(doc, "\"Na próxima noite: nomes na rede. Por que às vezes o WhatsApp funciona e o site não abre, o que é aquele cadeado do navegador e como reconhecer um site falso pelo endereço.\"", italic=True)

    add_hr(doc)

    # SE PERGUNTAREM
    add_heading(doc, "SE PERGUNTAREM")
    add_para(doc, "\"Dá pra descobrir onde a pessoa mora pelo IP?\"", bold=True)
    add_para(doc, "> \"Pelo IP de fora dá pra saber o provedor e, mais ou menos, a cidade. A rua e a casa, só o provedor sabe — e só entrega com ordem da Justiça.\"", italic=True)
    add_para(doc, "\"Posso bloquear o vizinho pelo MAC?\"", bold=True)
    add_para(doc, "> \"O roteador deixa, mas é fraco: MAC se troca fácil, e celular novo já troca sozinho. O que resolve é senha forte. Noite 11.\"", italic=True)
    add_para(doc, "\"E se dois aparelhos ficarem com o mesmo IP na mesma casa?\"", bold=True)
    add_para(doc, "> \"É o conflito de IP. O Windows avisa 'conflito de endereço IP' e um dos dois para de funcionar. Quase sempre acontece quando alguém digitou um IP fixo na mão. Com DHCP cuidando, não acontece.\"", italic=True)
    add_para(doc, "\"Ouvi falar de IPv6. O que é?\"", bold=True)
    add_para(doc, "> \"Os números do IP de hoje (IPv4) acabaram no mundo. O IPv6 é o endereço novo, bem mais comprido, com letras e números. Ele já está funcionando junto — vocês viram uma linha 'IPv6' no ipconfig. No dia a dia do técnico, quase sempre a gente trabalha com o IPv4.\"", italic=True)
    add_para(doc, "\"Por que o roteador de casa é 192.168.1.1 e o da escola 192.168.0.1?\"", bold=True)
    add_para(doc, "> \"Cada fabricante escolhe a rua de fábrica. TP-Link costuma ser 0, muitos modems de provedor são 1. Por isso o técnico sempre olha o 'Gateway Padrão' no ipconfig em vez de chutar.\"", italic=True)

    add_hr(doc)

    # DEPOIS DA AULA
    add_heading(doc, "DEPOIS DA AULA")
    add_bullet(doc, "Reative o DHCP do roteador da Estação B, ou etiquete ele como \"SEM DHCP — DEMONSTRAÇÃO\". Ele serve de novo na Noite 14")
    add_bullet(doc, "Mantenha os roteadores com os nomes EQUIPE-1, EQUIPE-2… — a Noite 9 parte deles")
    add_bullet(doc, "Anote quantas equipes acharam o intruso sozinhas. É o indicador da noite")
    add_bullet(doc, "Desconecte seu celular das redes das equipes (e esqueça as redes)")

    add_hr(doc)

    # DESCRITIVO
    add_heading(doc, "DESCRITIVO PARA O DIÁRIO")
    add_para(doc, "Desenvolvi atividade sobre endereçamento em redes locais. Apresentei o endereço físico (MAC) e o endereço lógico (IPv4) por meio de analogia com um condomínio, relacionando gateway padrão, atribuição automática de endereços pelo serviço DHCP do roteador e a diferença entre endereço privado (rede interna) e endereço público atribuído pelo provedor, incluindo o uso de MAC aleatório pelos dispositivos móveis. Os estudantes identificaram os endereços IP e MAC de celulares e computadores (ipconfig), acessaram a interface de administração dos roteadores da bancada em equipes e identificaram os dispositivos conectados pela lista de clientes, incluindo a localização de um dispositivo não autorizado. Realizaram testes de conectividade com o comando ping e diagnosticaram uma falha de atribuição de endereço (endereço de configuração automática 169.254) em roteador com DHCP desativado, aplicando os conceitos em simulações de atendimento ao cliente.")

    add_hr(doc)

    # CAIXA DE FERRAMENTAS DIDÁTICAS
    add_heading(doc, "CAIXA DE FERRAMENTAS DIDÁTICAS", color=COR_DESTAQUE)
    add_para(doc, "Se sobrar tempo, ou pra usar em outra noite:", italic=True, color=COR_CINZA)

    add_para(doc, "1. Quem fabricou esse MAC?", bold=True)
    add_para(doc, "Os 3 primeiros pares do MAC dizem o fabricante. No celular, pesquisar \"mac vendor lookup\" e digitar o MAC do roteador (da etiqueta) e do computador: aparece TP-Link, Intel, Realtek… Depois tente com um MAC aleatório de celular: não acha ninguém. Prova que é inventado. 10 min.")
    add_para(doc, "2. O porteiro vivo", bold=True)
    add_para(doc, "Encenação: um aluno é o roteador com uma folha de \"apartamentos\" (192.168.0.100 a .110). Os outros chegam um por um dizendo seu \"MAC\" (o nome). O roteador anota e entrega um número. Depois um sai e volta: recebe o mesmo número? Um chega sem o porteiro estar lá: fica com 169.254. 15 min.")
    add_para(doc, "3. Detetive da rede de casa", bold=True)
    add_para(doc, "Pra fazer em casa, com celular: entrar no painel do próprio roteador (o endereço é o gateway) e contar os aparelhos conectados. Quem não souber a senha do painel, só anota quantos aparelhos a família tem com Wi-Fi. Traz o número na próxima aula. 5 min pra explicar.")
    add_para(doc, "4. Cabo direto entre dois computadores", bold=True)
    add_para(doc, "Sem roteador no meio, os dois pegam 169.254. E mesmo assim conseguem se pingar! Mostra que 169.254 é \"rede improvisada\" — funciona entre os dois, mas não tem portaria pra sair. 10 min.")
    add_para(doc, "5. Mapa da turma", bold=True)
    add_para(doc, "Cada equipe desenha no papel o \"condomínio\" da sua rede: o roteador no meio com o gateway, e cada aparelho com IP, MAC e dono. Cola na parede. Na Noite 8 esse desenho vira o ponto de partida da planta de instalação. 15 min.")

    path = os.path.join(OUTDIR, "Roteiro - Redes 6 - Endereco na rede.docx")
    doc.save(path)
    return path


# ============================================================
# SLIDES
# ============================================================
def build_slides():
    prs = new_presentation()
    P, D, K, R = PPalette.P, PPalette.D, PPalette.K, PPalette.R

    def lista(kicker, titulo, linhas, subtitulo=None, size=20, nota=None, nota_cor=None):
        s = add_slide(prs)
        add_title_block(s, kicker, titulo, subtitulo)
        add_text(s, linhas, 0.7, 2.8, 11.9, 3.8, size=size, color=K)
        if nota:
            add_text(s, nota, 0.7, 6.4, 11.9, 0.5, size=16, italic=True, color=nota_cor or D, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    def frase(kicker, texto, cor=None):
        s = add_slide(prs)
        add_text(s, kicker, 0.7, 1.6, 11.9, 0.5, size=16, bold=True, color=D, align=PP_ALIGN.CENTER)
        add_text(s, texto.split("\n"), 0.9, 2.4, 11.5, 3.0, size=36, bold=True, color=cor or P, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    def tabela(kicker, titulo, rows, col_widths=None, subtitulo=None, size=16, nota=None, height=3.6):
        s = add_slide(prs)
        add_title_block(s, kicker, titulo, subtitulo)
        add_table_slide(s, rows, top=2.6, height=height, col_widths=col_widths, size=size)
        if nota:
            add_text(s, nota, 0.7, 6.5, 11.9, 0.5, size=15, italic=True, color=D, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    def tela(kicker, titulo, linhas_tela, nota=None, nota_cor=None):
        s = add_slide(prs)
        add_title_block(s, kicker, titulo)
        fundo = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, PInches(0.7), PInches(2.5), PInches(11.9), PInches(3.4))
        fundo.fill.solid()
        fundo.fill.fore_color.rgb = K
        fundo.line.fill.background()
        tb = add_text(s, linhas_tela, 1.0, 2.75, 11.3, 3.0, size=20, color=PPalette.W)
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = "Consolas"
        if nota:
            add_text(s, nota, 0.7, 6.15, 11.9, 0.7, size=17, italic=True, color=nota_cor or D, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    # 1: Capa
    add_slide_cover(prs, "NOITE 6", "ENDEREÇO", "NA REDE",
                    "IP, MAC e como o aparelho é encontrado",
                    "Redes de Computadores · Curso FIC · CEJA Itapiranga – SC")

    # 2: Recap
    lista("RETOMADA", "Da Noite 5 pra hoje", [
        "•  Quais os três jeitos de testar um cabo?",
        "•  Luz acesa garante cabo bom?",
        "•  Quais os 4 passos do método?",
        "•  No teste do balanço a gente pingou um número. Qual?",
    ])

    # 3: Censo
    lista("O CENSO", "Escreva no quadro o que você anotou", [
        "•  O IP e o MAC do seu celular (a tarefa)",
        "",
        "Agora olhe o quadro:",
        "•  O que os IPs têm em comum?",
        "•  Tem dois IPs iguais?",
        "•  E os MACs, tem algum igual?",
    ])

    # 4: Pergunta
    frase("A PERGUNTA DA NOITE", "Vinte celulares no mesmo Wi-Fi.\nChega um vídeo da internet.\nComo o roteador sabe de quem é?")

    # 5: Dois documentos
    frase("A RESPOSTA", "Todo aparelho tem dois documentos:\no CPF (MAC) e o número\ndo apartamento (IP).")

    # 6: Condomínio
    tabela("A REDE É UM CONDOMÍNIO", "Quem é quem", [["No condomínio…", "Na rede…"]] +
           [list(r) for r in CONDOMINIO],
           col_widths=[4.6, 7.3], size=15, height=4.0)

    # 7: MAC
    lista("O CPF", "Endereço MAC", [
        "•  Vem gravado de fábrica na placa de rede",
        "•  Cada placa tem o seu — não se repete no mundo",
        "•  6 pares de letras e números:   3C-52-82-1A-7F-09",
        "•  Está na etiqueta do roteador (lembra da Noite 2?)",
    ])

    # 8: IP
    lista("O NÚMERO DO APARTAMENTO", "Endereço IP", [
        "•  Quem dá é o roteador, na hora que o aparelho chega",
        "•  4 números de 0 a 255:   192.168.0.104",
        "•  Em casa: os três primeiros = a rua.  O último = a casa",
        "•  192.168.0.104 e 192.168.0.105 são vizinhos",
        "•  Em outra casa, o mesmo número pode existir",
    ], nota="Por isso dois colegas tinham o mesmo IP. Apartamento 101 tem em todo prédio.")

    # 9: MAC aleatório
    lista("PEGADINHA DA VIDA REAL", "O celular esconde o CPF", [
        "•  Celular novo inventa um MAC diferente pra cada Wi-Fi",
        "•  Android: \"MAC aleatório\"     iPhone: \"Endereço privado\"",
        "•  Motivo: privacidade — loja e shopping não te rastreiam",
        "",
        "O MAC que você anotou em casa pode não ser o de hoje.",
    ], nota_cor=R, nota="Vamos conferir na prática.")

    # 10: Prática - ache seus endereços
    lista("PRÁTICA · 20 MIN", "Ache seus endereços", [
        "1.  Desligue os dados móveis",
        "2.  Conecte no Wi-Fi da sua equipe (\"sem internet\"? conecte assim mesmo)",
        "3.  Android: Wi-Fi → engrenagem da rede    iPhone: Wi-Fi → (i)",
        "4.  Anote IP, MAC e se é aleatório",
        "5.  No computador: ipconfig  e  ipconfig /all",
    ], size=19, nota="Folha, página 1.")

    # 11: Rede sem internet
    frase("LEMBRA DA NOITE 1?", "Os roteadores estão sem internet.\nE a rede funciona assim mesmo.")

    # 12: DHCP
    lista("O PORTEIRO", "DHCP: quem distribui os endereços", [
        "O celular chega e pergunta: \"alguém me dá um endereço?\"",
        "",
        "O roteador responde com:",
        "•  um IP pra ele",
        "•  quem é a portaria (o gateway)",
        "•  por quanto tempo esse endereço vale",
    ])

    # 13: 169.254
    frase("QUANDO O PORTEIRO NÃO RESPONDE", "169.254.x.x\n=\nninguém deu endereço pra esse aparelho", cor=R)

    # 14: IP de fora
    lista("IP DE DENTRO × IP DE FORA", "O carteiro entrega na portaria", [
        "•  192.168… só vale dentro de casa",
        "•  Lá fora, a casa inteira tem UM endereço: o que o provedor dá",
        "•  O roteador recebe tudo e distribui pros apartamentos",
        "",
        "Teste: pesquise \"qual meu ip\" no Wi-Fi da escola e depois no 4G.",
    ])

    # 15: Intervalo
    frase("", "INTERVALO\n15 minutos", cor=K)

    # 16: Quem está na rede
    lista("PRÁTICA · 35 MIN", "Quem está na rede?", [
        "1.  No navegador do celular: 192.168.0.1",
        "2.  Senha de administrador: o professor passa",
        "3.  Ache a lista de aparelhos (caminho na fita do roteador)",
        "4.  Anote nome, IP e MAC de cada um — e de quem é",
        "5.  O MAC bate com o que você anotou?",
    ], nota="Hoje NÃO se muda nada no painel. Só olhar.", nota_cor=R)

    # 17: Nome não é documento
    frase("CUIDADO", "O nome que aparece na lista\nquem escolhe é o aparelho.\nQuem identifica é o MAC.")

    # 18: Intruso
    lista("PRÁTICA · 15 MIN", "Caça ao intruso", [
        "Uma das redes foi invadida.",
        "",
        "•  Confira a lista da sua equipe",
        "•  Qual aparelho é o intruso? IP? MAC?",
        "•  Como vocês chegaram nessa conclusão?",
    ], size=22, nota="No cliente: \"acho que o vizinho tá usando meu Wi-Fi\".")

    # 19: Ping
    lista("ESTAÇÃO A", "Ping: bater na porta", [
        "ping 192.168.0.1                  → a portaria responde?",
        "ping (IP do celular do colega)  → ele responde?",
        "Colega desliga o Wi-Fi e pinga de novo",
        "",
        "Respondeu: está lá e o caminho funciona.",
        "Não respondeu: não está, caminho quebrado — ou ele não abre a porta.",
    ], size=19)

    # 20: Estação B
    lista("ESTAÇÃO B", "O roteador mudo", [
        "•  Liga o computador por cabo. A luz acende",
        "•  ipconfig — o que aparece?",
        "•  Tenta o celular no Wi-Fi dele — o que acontece?",
        "•  Escreva o diagnóstico",
    ], size=22)

    # 21-23: Telas
    for nome, linhas_tela, _ in IPCONFIG:
        tela("LEIA A TELA", f"{nome}: o que está acontecendo?", linhas_tela)

    # 24: Gabarito telas
    tabela("GABARITO", "As três telas", [
        ["Tela", "O que quer dizer", "Por onde começa"],
        ["A · 192.168.0.104", "Normal", "Endereço ok — o problema está em outro lugar"],
        ["B · 169.254…", "Ninguém deu endereço", "Roteador, DHCP, porta LAN × WAN"],
        ["C · mídia desconectada", "Nem conectou", "Físico: cabo e porta (Noite 5)"],
    ], col_widths=[3.6, 3.4, 4.9], size=16)

    # 25: Chamados
    lista("ROLE-PLAY · 20 MIN", "Chamados", [
        "•  \"Conectado, sem internet. O cabo é novo, a luz acende.\"",
        "•  \"Faltou luz ontem e a impressora sumiu do computador.\"",
        "•  \"Acho que o vizinho está usando meu Wi-Fi.\"",
        "•  \"Meu filho disse que o IP dele vazou no jogo.\"",
    ], size=19, nota="Técnico: o que você olha primeiro? E depois?")

    # 26: Duelo
    lista("DUELO RELÂMPAGO", "Time A × Time B", [
        "•  5 segundos pra responder",
        "•  Errou, passa pro outro time",
        "•  Vale ponto quem explicar o porquê",
    ], size=24)

    # 27: O que fica
    lista("O QUE FICA", "Três frases pra guardar", [
        "\"MAC é o CPF. IP é o número do apartamento.\"",
        "",
        "\"169.254: ninguém deu endereço.\"",
        "",
        "\"Quem está na rede? Pergunte ao roteador.\"",
    ], size=24)

    # 28: Próxima
    lista("PARA A PRÓXIMA NOITE", "Nomes na rede: DNS e HTTPS", [
        "Uma tarefa simples:",
        "",
        "•  Abra três sites que você usa no celular",
        "•  Anote o endereço completo da barra — com o https://",
        "•  Repare: tem cadeado?",
    ])

    path = os.path.join(OUTDIR, "Aula - Redes 6 - Endereco na rede.pptx")
    prs.save(path)
    return path


# ============================================================
# FOLHA DO ALUNO
# ============================================================
def linhas(doc, n):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("_" * 100)
        r.font.size = Pt(11)
        r.font.color.rgb = COR_CINZA


def pergunta(doc, texto, size=11):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.font.size = Pt(size)
    r.bold = True


def bloco_tela(doc, linhas_tela):
    for ln in linhas_tela:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        r = p.add_run(ln)
        r.font.name = "Consolas"
        r.font.size = Pt(9)


def build_folha():
    doc = Document()
    setup_page(doc, margem=Cm(1.8))

    # PÁGINA 1 — Censo, condomínio, meus endereços
    cabecalho_documento(
        doc,
        "MEUS ENDEREÇOS NA REDE",
        f"Redes de Computadores · {NOITE} · Curso FIC · CEJA Itapiranga · 1 de 3",
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Nome: ")
    r.font.size = Pt(11)
    r.bold = True
    r = p.add_run("_________________________________________________     Equipe: ________")
    r.font.size = Pt(11)

    add_heading(doc, "1  A rede é um condomínio", level=2)
    add_para(doc, "Complete a coluna da direita durante a explicação.", italic=True, size=10, color=COR_CINZA)
    tbl = add_table(doc, [["No condomínio…", "Na rede…"]] + [[a, ""] for a, _ in CONDOMINIO],
                    col_widths=[6.5, 11])
    for row in tbl.rows[1:]:
        row.height = Cm(0.8)

    add_heading(doc, "2  Os endereços do meu celular", level=2)
    tbl = add_table(doc, [
        ["", "Em casa (tarefa)", "Hoje, no Wi-Fi da equipe"],
        ["Endereço IP", "", ""],
        ["Endereço MAC", "", ""],
        ["MAC aleatório / privado?", "", "☐ Sim   ☐ Não   ☐ Não achei"],
    ], col_widths=[5, 6, 6.5])
    for row in tbl.rows[1:]:
        row.height = Cm(0.9)

    pergunta(doc, "O MAC de casa e o de hoje são iguais? Por que você acha que isso aconteceu?")
    linhas(doc, 2)

    add_heading(doc, "3  Os endereços do computador (ipconfig e ipconfig /all)", level=2)
    tbl = add_table(doc, [
        ["Linha da tela", "O que apareceu"],
        ["Endereço IPv4", ""],
        ["Gateway Padrão", ""],
        ["Endereço Físico (MAC)", ""],
    ], col_widths=[5, 12.5])
    for row in tbl.rows[1:]:
        row.height = Cm(0.8)

    pergunta(doc, "O IP do computador e o do seu celular começam igual? O que isso quer dizer?")
    linhas(doc, 2)

    # PÁGINA 2 — Quem está na rede + intruso
    doc.add_page_break()
    cabecalho_documento(
        doc,
        "QUEM ESTÁ NA REDE?",
        f"Painel do roteador · {NOITE} · 2 de 3",
    )
    add_para(doc, "No navegador do celular, digite o endereço do gateway (192.168.0.1). Entre com a senha que o professor passar e ache a lista de aparelhos. Hoje NÃO se muda nada — só olhar e anotar.", italic=True, size=10, color=COR_CINZA)

    p = doc.add_paragraph()
    r = p.add_run("Nome do Wi-Fi da equipe: ____________________    Modelo do roteador: ____________________")
    r.font.size = Pt(11)
    p = doc.add_paragraph()
    r = p.add_run("Onde fica a lista no painel (o caminho): ______________________________________________")
    r.font.size = Pt(11)

    tbl = add_table(doc, [["#", "Nome que aparece na lista", "IP", "MAC", "De quem é?"]] +
                    [[str(i), "", "", "", ""] for i in range(1, 8)],
                    col_widths=[0.8, 5, 3.4, 4.8, 3.5])
    for row in tbl.rows[1:]:
        row.height = Cm(0.85)

    pergunta(doc, "Algum aparelho apareceu com nome esquisito ou repetido? Como vocês descobriram de quem era?")
    linhas(doc, 2)

    add_heading(doc, "Caça ao intruso", level=2)
    tbl = add_table(doc, [
        ["Tem intruso na nossa rede?", "☐ Sim   ☐ Não"],
        ["Nome que ele usa na lista", ""],
        ["IP do intruso", ""],
        ["MAC do intruso", ""],
    ], col_widths=[6.5, 11])
    for row in tbl.rows[1:]:
        row.height = Cm(0.8)
    pergunta(doc, "Como vocês chegaram nessa conclusão?")
    linhas(doc, 2)
    pergunta(doc, "Na casa de um cliente, o que você faria se achasse um aparelho que ninguém reconhece?")
    linhas(doc, 2)

    # PÁGINA 3 — Ping, sem-endereço, telas, chamado, tarefa
    doc.add_page_break()
    cabecalho_documento(
        doc,
        "PING, SEM-ENDEREÇO E LEITURA DE TELA",
        f"Diagnóstico · {NOITE} · 3 de 3",
    )

    add_heading(doc, "1  Estação A — Ping: quem responde?", level=2)
    tbl = add_table(doc, [
        ["Comando", "Respondeu?", "Tempo (ms) ou mensagem"],
        ["ping 192.168.0.1 (o roteador)", "☐ Sim   ☐ Não", ""],
        ["ping IP do celular do colega: ______________", "☐ Sim   ☐ Não", ""],
        ["O mesmo, com o Wi-Fi do celular desligado", "☐ Sim   ☐ Não", ""],
        ["ping 192.168.0.250 (ninguém tem)", "☐ Sim   ☐ Não", ""],
    ], col_widths=[7.5, 3.5, 6.5])
    for row in tbl.rows[1:]:
        row.height = Cm(0.8)

    add_heading(doc, "2  Estação B — O roteador mudo", level=2)
    p = doc.add_paragraph()
    r = p.add_run("A luz da porta acendeu?  ☐ Sim  ☐ Não        IP que apareceu no ipconfig: ________________________")
    r.font.size = Pt(11)
    pergunta(doc, "O que aconteceu com o celular? Qual o seu diagnóstico?")
    linhas(doc, 2)

    add_heading(doc, "3  Leia a tela: o que está acontecendo?", level=2)
    for nome, linhas_tela, _ in IPCONFIG:
        add_para(doc, nome, bold=True, size=10)
        bloco_tela(doc, linhas_tela)
        linhas(doc, 1)

    add_heading(doc, "4  Chamado", level=2)
    add_para(doc, "\"O computador diz 'conectado, sem internet'. O cabo é novo e a luz acende.\" O que você olha primeiro? E depois?", italic=True, color=COR_PRIMARIA, size=11)
    linhas(doc, 2)

    add_heading(doc, "TAREFA DA SEMANA", level=2, color=COR_DESTAQUE)
    add_bullet(doc, "Abra três sites que você usa no celular e anote o endereço completo da barra (com o https://)")
    add_bullet(doc, "1) ________________________________  2) ________________________________  3) ________________________________")
    add_bullet(doc, "Tem cadeado do lado do endereço?  ☐ Nos três   ☐ Em alguns   ☐ Em nenhum")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Traga esta folha na próxima aula · CEJA Itapiranga – SC")
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = COR_CINZA

    path = os.path.join(OUTDIR, "Folha do Aluno - Redes 6 - Endereco na rede.docx")
    doc.save(path)
    return path


# ============================================================
# RESUMO DO ALUNO
# ============================================================
def build_resumo():
    doc = Document()
    setup_page(doc, margem=Cm(1.8))

    cabecalho_documento(
        doc,
        "ENDEREÇO NA REDE: IP E MAC",
        f"Resumo do aluno · {NOITE} · Redes de Computadores · Curso FIC · CEJA Itapiranga",
    )

    add_heading(doc, "A ideia principal", level=2)
    add_para(doc, "Pra entregar cada coisa pro aparelho certo, a rede precisa de endereço. Todo aparelho tem dois: o MAC (o CPF, vem de fábrica) e o IP (o número do apartamento, quem dá é o roteador).")

    add_heading(doc, "A rede é um condomínio", level=2)
    add_table(doc, [["No condomínio…", "Na rede…"]] + [list(r) for r in CONDOMINIO], col_widths=[5.5, 10])

    add_heading(doc, "Como ler os endereços", level=2)
    add_bullet(doc, "IP: 4 números de 0 a 255 com ponto (ex.: 192.168.0.104). Em casa, os três primeiros são a rua (a rede) e o último é a casa (o aparelho)")
    add_bullet(doc, "MAC: 6 pares de letras e números (ex.: 3C-52-82-1A-7F-09). Os 3 primeiros pares dizem o fabricante")
    add_bullet(doc, "Mesmo IP em casas diferentes: normal. Mesmo IP na mesma casa: conflito. MAC repetido: não existe")
    add_bullet(doc, "Celular novo usa MAC aleatório (Android) ou endereço privado (iPhone): um MAC inventado pra cada Wi-Fi, por privacidade")

    add_heading(doc, "Onde achar", level=2)
    add_table(doc, [
        ["Aparelho", "Caminho"],
        ["Android", "Configurações → Wi-Fi → engrenagem da rede (MAC em Avançado/Privacidade)"],
        ["iPhone", "Ajustes → Wi-Fi → (i) ao lado da rede"],
        ["Windows", "Prompt de Comando → ipconfig (IP e gateway) · ipconfig /all (Endereço Físico = MAC)"],
        ["Roteador", "Navegador → endereço do gateway (ex.: 192.168.0.1) → lista de clientes / dispositivos"],
    ], col_widths=[3, 12.5])

    add_heading(doc, "O que a tela está dizendo", level=2)
    add_table(doc, [
        ["Você vê…", "Quer dizer", "Por onde começar"],
        ["IP 192.168.x.x e gateway preenchido", "Recebeu endereço normal", "O problema está em outro lugar"],
        ["IP 169.254.x.x, gateway vazio", "Ninguém deu endereço", "Roteador, DHCP, cabo na porta WAN em vez da LAN"],
        ["\"Mídia desconectada\"", "Nem conectou", "Físico: cabo e porta (método da Noite 5)"],
    ], col_widths=[5.5, 4.3, 5.7])

    add_heading(doc, "Ping: bater na porta", level=2)
    add_para(doc, "No Prompt: ping e o IP. Respondeu (com tempo em ms): o aparelho está lá e o caminho funciona. \"Esgotado o tempo limite\" ou \"Host de destino inacessível\": não está, o caminho está quebrado — ou é um aparelho que não responde ping por segurança (alguns celulares e o Windows fazem isso).")

    add_heading(doc, "No atendimento", level=2)
    add_bullet(doc, "\"Acho que o vizinho usa meu Wi-Fi\": entre no painel do roteador, conte os aparelhos da casa com o cliente. Sobrou? Troque a senha")
    add_bullet(doc, "\"A impressora sumiu depois que faltou luz\": ela ganhou outro IP. Ache o novo no painel pelo MAC")
    add_bullet(doc, "O nome que aparece na lista do roteador quem escolhe é o aparelho. Quem identifica é o MAC")

    add_heading(doc, "Três frases pra guardar", level=2)
    add_para(doc, "\"MAC é o CPF. IP é o número do apartamento.\"", bold=True, color=COR_PRIMARIA)
    add_para(doc, "\"169.254: ninguém deu endereço.\"", bold=True, color=COR_PRIMARIA)
    add_para(doc, "\"Quem está na rede? Pergunte ao roteador.\"", bold=True, color=COR_PRIMARIA)

    add_heading(doc, "Pra se testar em casa", level=2)
    add_bullet(doc, "Seu celular e o do seu vizinho podem ter o mesmo IP? E o mesmo MAC? Por quê?")
    add_bullet(doc, "O computador do cliente mostra 169.254. O cabo está bom. O que você suspeita?")
    add_bullet(doc, "Qual o endereço do gateway na sua casa? (Dica: ipconfig, ou detalhes do Wi-Fi no celular)")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Este material é seu. Guarde e volte quando precisar. · CEJA Itapiranga – SC")
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = COR_CINZA

    path = os.path.join(OUTDIR, "Resumo do Aluno - Redes 6 - Endereco na rede.docx")
    doc.save(path)
    return path


# ============================================================
# MAIN
# ============================================================
if __name__ == "__main__":
    print("Gerando Roteiro...")
    print("  →", build_roteiro())
    print("Gerando Slides...")
    print("  →", build_slides())
    print("Gerando Folha do Aluno...")
    print("  →", build_folha())
    print("Gerando Resumo do Aluno...")
    print("  →", build_resumo())
    print("\nPronto.")
