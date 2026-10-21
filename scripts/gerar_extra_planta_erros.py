"""Atividade extra — Caça aos erros na planta (noite sem energia no laboratório).

Duas plantas desenhadas com 12 erros e 3 acertos cada, revisando as Noites 1 a 8:
posição do roteador, cabos, energia, endereços, segurança do Wi-Fi e repetidor.
Não precisa de computador nem de tomada: imprime, ou manda o PNG no grupo da
turma e o laudo vai no caderno. Sem role-play e sem duelo.
Gera: folha do aluno, gabarito do professor e as duas plantas em PNG.
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT, WD_SECTION

from common import (
    COR_PRIMARIA, COR_ALERTA, COR_CINZA,
    add_heading, add_para, add_bullet, add_table, add_hr,
    setup_page, cabecalho_documento,
)
import plantas_erros

OUTDIR = "/home/nedel/curso"
NOME = "Redes - Extra - Planta com erros"
PNG_A = os.path.join(OUTDIR, f"Planta A - {NOME}.png")
PNG_B = os.path.join(OUTDIR, f"Planta B - {NOME}.png")


# Gabarito: (onde, erro, por que é problema, como corrigir)
ERROS_A = [
    ("A3/A4", "Roteador no canto da casa (área de serviço)", "O sinal sai em todas as direções; do canto, metade vai pra rua e os quartos ficam sem", "Levar o roteador pro centro da casa (sala); a ONU fica onde a fibra chega e um cabo leva até ele"),
    ("A3", "Roteador no chão, atrás da máquina de lavar", "Chão, metal e água seguram o sinal; perto da máquina ainda tem umidade", "Alto, entre 1,5 e 2 m, livre, longe de eletrodoméstico e metal"),
    ("A4", "Benjamim: ONU e roteador junto com a máquina de lavar", "Motor da máquina dá pico e queda; benjamim esquenta e solta", "Tomada própria ou régua boa com proteção; nobreak se puder"),
    ("A3", "ONU e roteador com o mesmo endereço 192.168.1.1", "Dois aparelhos com o mesmo número: o painel some, a internet falha", "Mudar o endereço do painel do roteador (ex.: 192.168.10.1) — Noite 8"),
    ("C3", "Nome e senha de fábrica (TP-Link_3F2A / 12345678, a da etiqueta)", "Senha curta e conhecida; quem vê a etiqueta ou o bilhete entra", "Nome novo sem dado pessoal e senha longa nova; anotar na ficha de entrega"),
    ("C3", "Segurança WEP", "WEP é tranca velha, quebrada em minutos", "WPA2-PSK com AES (ou WPA3)"),
    ("C3", "Visitas usam a mesma senha da casa", "Visita enxerga impressora, câmera, PC; a senha se espalha", "Rede de convidados com senha própria"),
    ("B1/B2", "Cabo da TV no mesmo conduíte da fiação 220 V", "A fiação elétrica atrapalha o sinal do cabo; e não é permitido dividir", "Conduíte ou canaleta separada só pra rede"),
    ("B4/C4", "Cabo solto no chão, por baixo da porta", "É pisado, a porta amassa, alguém tropeça; o cabo quebra por dentro", "Canaleta no rodapé ou passar pelo forro"),
    ("A1/A2", "Cabo da câmera: cabo interno exposto ao sol e chuva", "A capa de cabo interno resseca e racha com o sol; entra água", "Cabo de uso externo (capa preta, proteção UV) ou dentro de conduíte — Noite 4"),
    ("D4", "Impressora com endereço automático", "Cada vez ganha um número; os PCs procuram no número velho", "Reserva de endereço pelo MAC — Noite 8"),
    ("F3", "Repetidor no quarto do fundo, onde chega 1 barra", "Repetidor só repete o que recebe: 1 barra vira 1 barra repetida", "Repetidor no meio do caminho, onde pega 3 de 4 barras; melhor ainda: cabo até um segundo ponto de acesso"),
]
ACERTOS_A = [
    ("A3", "Senha do painel trocada pelo técnico"),
    ("D3", "Cabo do PC: CAT5e, 15 m, testado com 8 fios OK (o cabo está certo; o caminho dele é que está errado)"),
    ("D3", "Cópia de segurança da configuração salva no PC"),
]

ERROS_B = [
    ("B1 e F1", "Dois aparelhos entregando endereço (DHCP ligado nos dois)", "Cada aparelho ganha número de um, uns funcionam, outros não; conflito", "Desligar o DHCP do roteador da loja (vira só ponto de acesso) ou ligar pela porta WAN"),
    ("A2", "Wi-Fi dos clientes é a mesma rede do caixa e das câmeras", "Qualquer cliente enxerga o PC do caixa, a maquininha e as câmeras", "Rede de convidados só pros clientes"),
    ("F2", "Painel com admin/admin, nunca trocado, etiqueta colada", "Qualquer um conectado entra no painel e muda tudo", "Senha de administrador nova, anotada na ficha de entrega"),
    ("C1", "Switch no forro, em cima da lâmpada, sem ventilação", "Esquenta, trava, queima; ninguém acha pra reiniciar", "Lugar ventilado e acessível: caixa ou mini-rack na parede"),
    ("D1/D2", "Cabo de rede no mesmo eletroduto do ar-condicionado", "Fiação de motor atrapalha o sinal; risco elétrico", "Eletroduto ou canaleta separada"),
    ("B3/C3", "Cabo atravessando o corredor no chão", "O carrinho amassa, quebra por dentro, o caixa cai no meio da venda", "Canaleta de piso, passar pelo forro ou pela parede"),
    ("A3", "Conector com a capa fora do RJ45 (fios aparecendo)", "Os fios soltam com o tempo, mau contato — Noite 4/5", "Cortar e crimpar de novo, com a capa dentro do conector; testar"),
    ("B3", "PC, maquininha e freezer na mesma régua, sem nobreak", "Freezer liga e desliga o motor (pico); faltou luz, o caixa para", "Freezer em tomada própria; nobreak para PC, maquininha, ONU, roteador e switch"),
    ("B4", "Câmera do pátio: emenda com fita isolante, pendurada na calha", "Emenda pega umidade e oxida; calha é água e sol", "Cabo inteiro, de uso externo, preso com abraçadeira ou em conduíte"),
    ("F4", "Cabo de 120 m direto do switch até a câmera do depósito", "O cabo de rede vai até 100 m; passou disso, falha ou nem conecta", "Switch no depósito (cabo até 100 m) ou fibra/rádio até lá"),
    ("D4", "Repetidor dentro do galpão de chapa metálica", "A chapa barra o sinal: o repetidor não recebe nada de fora", "Levar cabo até o depósito e colocar um ponto de acesso lá dentro"),
    ("E2/F2", "PC e impressora com a mesma reserva 192.168.0.50", "Dois aparelhos com o mesmo número: um dos dois some da rede", "Cada reserva com um número diferente"),
]
ACERTOS_B = [
    ("B1", "Roteador da loja na parede, a 2 m de altura, no meio da loja"),
    ("B1", "DNS 8.8.8.8 e 1.1.1.1 entregue pelo roteador"),
    ("E2", "PC do escritório com cabo CAT6 testado"),
]


def campo(doc, texto, size=11, space_after=4, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(texto)
    r.font.size = Pt(size)
    r.bold = bold
    return p


def linha(doc, n=1):
    for _ in range(n):
        p = campo(doc, "_" * 100, space_after=0)
        p.runs[0].font.color.rgb = COR_CINZA


def secao(doc, paisagem):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    larg, alt = Cm(21.0), Cm(29.7)
    sec.orientation = WD_ORIENT.LANDSCAPE if paisagem else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = (alt, larg) if paisagem else (larg, alt)
    m = Cm(1.2) if paisagem else Cm(1.8)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = m
    return sec


def pagina_planta(doc, png):
    secao(doc, paisagem=True)
    doc.add_picture(png, width=Cm(25.5))


def tabela_laudo(doc, n):
    tbl = add_table(doc, [["Nº", "Onde (grade)", "O que está errado", "Por que é problema", "Como o técnico corrige"]] +
                    [[str(i), "", "", "", ""] for i in range(1, n + 1)],
                    col_widths=[0.9, 1.9, 5.0, 4.7, 4.9])
    for row in tbl.rows[1:]:
        row.height = Cm(1.25)
    return tbl


def tabela_acertos(doc):
    tbl = add_table(doc, [["Onde (grade)", "O que está CERTO", "Por que está certo"]] + [["", "", ""] for _ in range(3)],
                    col_widths=[2.3, 7.5, 7.6])
    for row in tbl.rows[1:]:
        row.height = Cm(1.0)


# ============================================================
# FOLHA DO ALUNO
# ============================================================
def build_folha():
    doc = Document()
    setup_page(doc, margem=Cm(1.8))
    cabecalho_documento(doc, "CAÇA AOS ERROS NA PLANTA",
                        "Redes de Computadores · Atividade extra · Curso FIC · CEJA Itapiranga – SC")
    campo(doc, "Nome: _______________________________________________     Dupla/equipe: ____________", space_after=8)

    add_heading(doc, "O serviço")
    add_para(doc, "Você foi chamado pra olhar a rede de dois clientes. Outro \"técnico\" já passou por lá antes e deixou tudo do jeito dele. Seu trabalho é achar o que está errado, explicar por que é problema e dizer como se corrige — do jeito que você vai explicar pro cliente.")
    add_para(doc, "Cada planta tem 12 ERROS e 3 coisas CERTAS. Cuidado: nem tudo que parece estranho está errado.", bold=True, color=COR_PRIMARIA)

    add_heading(doc, "Como fazer", level=2)
    add_bullet(doc, "Use a grade da planta pra dizer onde está o erro: letra da coluna + número da linha (ex.: C3)")
    add_bullet(doc, "Leia tudo: as caixas azuis são os aparelhos, as amarelas são bilhetes, etiquetas e anotações")
    add_bullet(doc, "Olhe as linhas: azul é cabo de rede, verde é a fibra, vermelha tracejada é fiação elétrica")
    add_bullet(doc, "Planta A: primeiro sozinho (15 min), depois compare com a dupla e completem juntos")
    add_bullet(doc, "Planta B: em equipe. É mais difícil: é um comércio")
    add_bullet(doc, "Um sintoma (\"a internet trava\") não é erro: é a reclamação. O erro é a causa")

    add_heading(doc, "Do que já vimos, o que entra aqui", level=2)
    add_table(doc, [
        ["Assunto", "Noite"],
        ["Onde colocar o roteador, o que atrapalha o sinal", "3"],
        ["Tipo de cabo, cabo externo, limite de distância, caminho do cabo", "3, 4"],
        ["Conector bem crimpado e testado", "4, 5"],
        ["Endereço do painel, endereço repetido, reserva por MAC", "6, 8"],
        ["A agenda (DNS)", "7"],
        ["Nome, senha, WPA2, rede de convidados, cópia de segurança", "8"],
    ], col_widths=[14, 3.4])

    # PLANTA A
    pagina_planta(doc, PNG_A)
    secao(doc, paisagem=False)
    cabecalho_documento(doc, "LAUDO — PLANTA A", "Casa da Dona Lurdes · ache os 12 erros")
    tabela_laudo(doc, 12)
    add_heading(doc, "As 3 coisas que estão certas", level=2)
    tabela_acertos(doc)

    # PLANTA B
    pagina_planta(doc, PNG_B)
    secao(doc, paisagem=False)
    cabecalho_documento(doc, "LAUDO — PLANTA B", "Agropecuária do Seu Ademir · ache os 12 erros")
    tabela_laudo(doc, 12)
    add_heading(doc, "As 3 coisas que estão certas", level=2)
    tabela_acertos(doc)

    # FECHAMENTO
    doc.add_page_break()
    cabecalho_documento(doc, "O QUE ARRUMAR PRIMEIRO", "Prioridade, recado ao cliente e a planta corrigida")
    add_heading(doc, "1  Não dá pra arrumar tudo hoje", level=2)
    add_para(doc, "O cliente só tem tempo (e dinheiro) pra 3 correções agora. Escolha as 3 mais urgentes de cada planta e diga por quê.", size=10, italic=True, color=COR_CINZA)
    tbl = add_table(doc, [["", "Planta A — Dona Lurdes", "Planta B — Seu Ademir"],
                          ["1º", "", ""], ["2º", "", ""], ["3º", "", ""],
                          ["Por que essas?", "", ""]], col_widths=[2.4, 7.5, 7.5])
    for row in tbl.rows[1:]:
        row.height = Cm(1.3)

    add_heading(doc, "2  O recado pra Dona Lurdes", level=2)
    add_para(doc, "Escreva a mensagem de WhatsApp que você mandaria pra ela explicando o que achou. Ela não entende de rede: nada de sigla sem explicar. Máximo 6 linhas.", size=10, italic=True, color=COR_CINZA)
    linha(doc, 6)

    add_heading(doc, "3  A planta corrigida", level=2)
    add_para(doc, "Redesenhe a Planta A do jeito certo: onde fica a ONU, o roteador, o repetidor (se ainda precisar), por onde passam os cabos. Use o verso desta folha ou uma folha em branco.", size=10, italic=True, color=COR_CINZA)
    campo(doc, "Com o roteador no lugar certo, a Dona Lurdes ainda precisa do repetidor?  ☐ Sim  ☐ Não   Por quê?")
    linha(doc, 2)

    doc.save(os.path.join(OUTDIR, f"Folha do Aluno - {NOME}.docx"))


# ============================================================
# GABARITO DO PROFESSOR
# ============================================================
def build_gabarito():
    doc = Document()
    setup_page(doc)
    cabecalho_documento(doc, "CAÇA AOS ERROS NA PLANTA",
                        "Roteiro e gabarito do professor · Atividade extra (sala sem energia) · Redes de Computadores")

    add_heading(doc, "A ideia")
    add_para(doc, "Sem computador e sem tomada: duas plantas (uma casa e um comércio) montadas por um \"técnico\" anterior, com 12 erros e 3 acertos cada. A turma acha, explica e corrige. Revisa tudo das Noites 1 a 8 e serve de aquecimento pra Noite 10 (segurança).")
    add_para(doc, "Se não der pra imprimir: mande as duas imagens PNG no grupo da turma. Eles olham no celular (dá zoom) e fazem o laudo no caderno, com as mesmas colunas da folha. Desenhe a grade A–F × 1–4 no quadro pra correção.", italic=True, color=COR_ALERTA)

    add_heading(doc, "Tempo (cerca de 3 h)", level=2)
    add_table(doc, [
        ["Bloco", "Min", "Como"],
        ["Abertura", "10", "Explicar a grade e as cores. Combinar: sintoma não é erro"],
        ["Planta A — sozinho", "15", "Cada um caça em silêncio"],
        ["Planta A — dupla", "20", "Comparar e completar \"por que\" e \"como corrige\""],
        ["Correção A", "20", "Cada dupla diz um erro por vez, no quadro; a turma confirma ou contesta"],
        ["Planta B — equipe", "35", "Equipes de 3–4. Mais difícil: DHCP duplo, 120 m, reserva repetida"],
        ["Correção B", "25", "Igual à A. Puxar as 3 pegadinhas no fim"],
        ["Prioridade + recado", "25", "Página final, partes 1 e 2. Duas ou três duplas leem o recado"],
        ["Planta corrigida", "30", "Parte 3. Se tiver papel pardo, vira cartaz por equipe e cada uma apresenta"],
    ], col_widths=[4.5, 1.5, 11])

    add_heading(doc, "Onde dá pra esticar", level=2)
    add_bullet(doc, "Em cada erro: \"Quem já viu isso na vida real? Onde?\" — sempre aparece história")
    add_bullet(doc, "Discordância vale ouro: deixe uma dupla defender que algo NÃO é erro antes de dar a resposta")
    add_bullet(doc, "Na Planta B, pergunte o que o Seu Ademir perde em dinheiro com cada erro (caixa parado, câmera sem gravar)")
    add_bullet(doc, "No recado, leia um em voz alta e a turma \"traduz\" o que ainda está difícil pro cliente")
    add_bullet(doc, "Desafio extra: cada dupla cria a SUA planta com 5 erros escondidos e troca com outra dupla")

    for titulo, erros, acertos, nota in [
        ("GABARITO — PLANTA A · Casa da Dona Lurdes", ERROS_A, ACERTOS_A,
         "Notebook com 1 barra (E2) é sintoma, não erro: a causa é o roteador no canto (erro 1). Aceite se o aluno ligar os dois. Smart TV no cabo também está certo — aceite como acerto extra."),
        ("GABARITO — PLANTA B · Agropecuária do Seu Ademir", ERROS_B, ACERTOS_B,
         "\"O coletor do estoque não conecta\" é sintoma do erro 11. A placa com a senha dos clientes não é erro em si: o erro é ser a mesma rede do caixa (erro 2). Se alguém disser que a senha é fraca, aceite como observação."),
    ]:
        doc.add_page_break()
        add_heading(doc, titulo)
        tbl = add_table(doc, [["Nº", "Onde", "Erro", "Por que é problema", "Como corrige"]] +
                        [[str(i), o, e, p, c] for i, (o, e, p, c) in enumerate(erros, 1)],
                        col_widths=[0.8, 1.6, 4.4, 5.0, 5.6])
        for row in tbl.rows:
            for cell in row.cells:
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(9)
        add_heading(doc, "As 3 pegadinhas (estão certas)", level=2)
        add_table(doc, [["Onde", "O que está certo"]] + [list(a) for a in acertos], col_widths=[2, 15.4])
        add_para(doc, nota, italic=True, size=10, color=COR_CINZA)

    add_heading(doc, "Prioridade esperada (não é única)", level=2)
    add_bullet(doc, "Dona Lurdes: senha/WEP/convidados (segurança), roteador pro centro e pro alto, tirar do benjamim")
    add_bullet(doc, "Seu Ademir: rede de convidados (clientes vendo o caixa), régua com freezer/nobreak, DHCP duplo. Aceite qualquer escolha bem justificada")

    doc.save(os.path.join(OUTDIR, f"Gabarito - {NOME}.docx"))


if __name__ == "__main__":
    plantas_erros.planta_a(PNG_A)
    plantas_erros.planta_b(PNG_B)
    build_folha()
    build_gabarito()
    print("ok")
