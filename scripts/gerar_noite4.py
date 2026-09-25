"""Gera os 4 arquivos da Noite 4 do curso de Redes.

Noite 4 — Cabo de rede na prática: categorias, capas e crimpagem
Padrão prática-pesada. Aluno passa a maior parte da noite com o alicate na mão.
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from common import (
    COR_PRIMARIA, COR_ALERTA, COR_DESTAQUE, COR_CINZA,
    add_heading, add_para, add_bullet, add_table, add_hr,
    setup_page, cabecalho_documento,
    PPalette, new_presentation, add_slide, add_text, add_title_block,
    add_slide_cover, add_footer, add_table_slide, add_bar,
)
from pptx.util import Inches as PInches, Pt as PPt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTDIR = "/home/nedel/curso"
NOITE = "Noite 4"
FOOTER = f"Redes de Computadores · {NOITE}"


# ============================================================
# ROTEIRO
# ============================================================
def build_roteiro():
    doc = Document()
    setup_page(doc)

    cabecalho_documento(
        doc,
        "CABO NA MÃO — CRIMPAGEM",
        f"Roteiro do professor · Redes de Computadores · {NOITE} · Curso FIC · CEJA Itapiranga – SC",
    )

    # A IDEIA
    add_heading(doc, "A IDEIA DESTA NOITE")
    add_para(doc, "Na Noite 3, os alunos aprenderam a reconhecer o cabo. Hoje eles fazem o cabo. Vai ter dedo doendo, vai ter conector desperdiçado, vai ter cabo torto — é parte da noite. Todo técnico de campo crimpou cabo errado no início. O que separa o profissional é a quantidade de conectores que ele já queimou até parar de errar.")
    add_para(doc, "Um cabo bom por aluno no fim da noite = objetivo cumprido. Se um aluno tirar dois cabos bons de primeira, é raríssima exceção — celebre, mas não use como padrão.")

    add_heading(doc, "O que se resolve nesta noite", level=2)
    add_para(doc, "Cada aluno crimpa e testa seu próprio cabo. Sai daqui sabendo o padrão T568B de cor, e reconhecendo os 6 erros mais comuns quando vir o cabo mal-feito de outro técnico.")

    add_heading(doc, "Esta é a quarta de quinze noites", level=2)
    add_para(doc, "É a noite mais suja do curso e a mais divertida. Você vai gastar conector — reserve pelo menos 4-5 por aluno pra ter margem. Se der pra ter um testador de cabo, o efeito \"deu certo!\" que ele produz vale ouro pra motivação da turma.")

    add_hr(doc)

    # ANTES DA AULA
    add_heading(doc, "ANTES DA AULA")

    add_heading(doc, "Material da mesa (pesado)", level=2)
    add_bullet(doc, "Alicate crimpador — se possível, 3 ou 4 pra rodar entre equipes (com um só, o rodízio consome tempo mas funciona)")
    add_bullet(doc, "Rolo de cabo Cat5e ou Cat6, uns 20-30 m (cada aluno vai usar ~40 cm — reserva conta perdas)")
    add_bullet(doc, "Conectores RJ45 — 50 a 100 unidades pra 15 alunos (reserva 4-5 por pessoa, esperando erros)")
    add_bullet(doc, "Tesoura ou estilete")
    add_bullet(doc, "Régua ou trena")
    add_bullet(doc, "Testador de cabo (LED 8 pinos) — se tiver, é o melhor investimento pra essa noite")
    add_bullet(doc, "Capas protetoras (opcional)")

    add_heading(doc, "Prepare antes", level=2)
    add_bullet(doc, "Crimpe UM cabo perfeito antes da aula. Vai ser sua referência (\"o cabo modelo\").")
    add_bullet(doc, "Crimpe 3 cabos com defeitos intencionais (destrançado demais, fios em tamanhos diferentes, capa fora do conector). Esses são material didático — mostre no bloco de erros.")
    add_bullet(doc, "Confira o testador com o cabo modelo antes da aula começar. Testador ruim vira pesadelo em sala.")
    add_bullet(doc, "Prepare uma bancada com jornal ou toalha embaixo — cai cabinho pra todo lado.")

    add_heading(doc, "Imprima", level=2)
    add_bullet(doc, "Folha do aluno (3 páginas), uma por pessoa")
    add_bullet(doc, "Resumo do aluno, uma por pessoa (pra levar pra casa)")

    add_heading(doc, "Segurança rápida", level=2)
    add_bullet(doc, "Manga curta ou dobrada. Cabo suja o braço.")
    add_bullet(doc, "Alicate crimpador aperta forte. Não deixe dedo no lugar errado — vai pinçar.")
    add_bullet(doc, "Fio pequeno no chão fica invisível e pode furar pé. Recolha ao fim.")

    add_hr(doc)

    # QUADRO DE TEMPO
    add_heading(doc, "QUADRO DE TEMPO — 210 MINUTOS")
    add_table(doc, [
        ["Bloco", "Tempo", "O que acontece"],
        ["1. Retomada + cabos que trouxeram", "10 min", "Recap Noite 3 + análise dos cabos que apareceram"],
        ["2. Anatomia do cabo + ferramentas", "20 min", "Por dentro do par trançado, alicate, estilete"],
        ["3. Padrão T568B (e T568A pra saber)", "20 min", "Ordem das cores, cabo direto × cruzado"],
        ["4. Demonstração do professor", "20 min", "Passo a passo ao vivo, comentando cada erro possível"],
        ["INTERVALO", "15 min", ""],
        ["5. PRÁTICA — Cada aluno crimpa", "55 min", "Todo mundo com cabo na mão, professor circula"],
        ["6. Teste e correção", "25 min", "Testador roda pela sala, quem falhou refaz"],
        ["7. Análise coletiva de erros", "20 min", "Você pega 3-4 cabos ruins e analisa com a turma"],
        ["8. Duelo relâmpago crimpagem", "15 min", "Time A × Time B com perguntas sobre cores e erros"],
        ["9. Fechamento + tarefa", "10 min", "Quiz de saída e anúncio Noite 5 (teste de cabo)"],
    ], col_widths=[6, 2, 8.5])
    add_para(doc, "155 min de atividade em bancada (74% do total). A parte expositiva desta noite é curta de propósito — a mão faz o resto.", italic=True, color=COR_CINZA)

    add_hr(doc)

    # BLOCO 1
    add_heading(doc, "BLOCO 1 — RETOMADA + CABOS QUE TROUXERAM · 10 MIN")

    add_heading(doc, "Recap relâmpago da Noite 3 — 4 min", level=2)
    add_para(doc, "Perguntas rápidas, resposta em voz alta:")
    add_bullet(doc, "Quais são os três meios de transmissão? (cabo, fibra, onda)")
    add_bullet(doc, "Cat5e aguenta quantos metros? (100 m)")
    add_bullet(doc, "Cabo par trançado tem quantos fios por dentro? (8, em 4 pares)")
    add_bullet(doc, "Por que os pares são trançados? (proteger contra interferência)")
    add_bullet(doc, "Qual é o conector do cabo de rede? (RJ45)")

    add_heading(doc, "Vitrine dos cabos que apareceram — 6 min", level=2)
    add_para(doc, "Peça pra galera colocar em cima da mesa os cabos que trouxeram (ou as fotos). Você comenta 3-4:")
    add_bullet(doc, "\"Esse aqui, olhem: Cat6 UTP. Marca X. Dá pra ver a metragem?\"")
    add_bullet(doc, "\"Esse conector aqui já entortou. Vamos ver por quê no bloco de erros.\"")
    add_bullet(doc, "\"Quem trouxe cabo com conector quebrado? Deixa aqui na mesa, vou pedir emprestado.\"")
    add_para(doc, "Não faça auditoria de tarefa. Quem trouxe é reconhecido, quem não trouxe não é constrangido.")

    add_hr(doc)

    # BLOCO 2
    add_heading(doc, "BLOCO 2 — ANATOMIA DO CABO + FERRAMENTAS · 20 MIN")

    add_heading(doc, "Por dentro do cabo — 8 min", level=2)
    add_para(doc, "Corte um pedaço curto de cabo na hora, na frente deles. Ou tenha um pedaço já cortado com as pontas abertas.")
    add_bullet(doc, "A capa externa protege o resto")
    add_bullet(doc, "Um fio de nylon fininho no meio ajuda a puxar a capa (dá pra ver, é resistente)")
    add_bullet(doc, "8 fios coloridos, formando 4 pares")
    add_bullet(doc, "Os pares: laranja/branco+laranja, verde/branco+verde, azul/branco+azul, marrom/branco+marrom")
    add_bullet(doc, "Cada par é trançado sobre si mesmo — vocês viram isso na Noite 3")
    add_para(doc, "\"O nome 'branco' na verdade é 'branco listrado com a cor do par'. Prestem atenção nas listras — é como você reconhece qual é o par de qual fio.\"", italic=True, color=COR_PRIMARIA)

    add_heading(doc, "O alicate crimpador — 7 min", level=2)
    add_para(doc, "Mostre o alicate. Aponte cada parte:")
    add_bullet(doc, "Boca de 8 pinos (RJ45) — o furo grande, onde entra o conector")
    add_bullet(doc, "Boca de 6 pinos (RJ11) — o furo menor, pra telefone (não vamos usar)")
    add_bullet(doc, "Lâmina de corte — na base ou lateral, pra cortar cabo no tamanho")
    add_bullet(doc, "Decapador — buraquinho pra passar o cabo, gira e tira só a capa (nem todo alicate tem)")
    add_para(doc, "\"O alicate tem tudo que você precisa: cortar o cabo, tirar a capa e crimpar o conector. Se tiver que trocar de ferramenta no meio da instalação, você atrasa. Aprenda a usar tudo do mesmo alicate.\"", italic=True)

    add_heading(doc, "Outras ferramentas — 5 min", level=2)
    add_bullet(doc, "Estilete: alternativa pro decapador do alicate. Precisa cuidado — cortou o fio de dentro, o cabo tá perdido.")
    add_bullet(doc, "Régua ou trena: pra medir tamanho da ponta a decapar")
    add_bullet(doc, "Testador: veremos depois")

    add_hr(doc)

    # BLOCO 3
    add_heading(doc, "BLOCO 3 — T568B (E T568A PRA SABER) · 20 MIN")

    add_heading(doc, "Existem dois padrões — 3 min", level=2)
    add_para(doc, "\"O cabo de rede tem duas normas de cor: T568A e T568B. As duas funcionam. A diferença é qual par vai onde. No Brasil, e na maioria do mundo, a gente usa T568B. Vou ensinar B primeiro; depois mostro A rapidinho pra vocês saberem que existe.\"", italic=True)

    add_heading(doc, "O padrão T568B — 10 min", level=2)
    add_para(doc, "Escreva no quadro, grande:")
    add_table(doc, [
        ["Pino", "Cor"],
        ["1", "Laranja / Branco"],
        ["2", "Laranja"],
        ["3", "Verde / Branco"],
        ["4", "Azul"],
        ["5", "Azul / Branco"],
        ["6", "Verde"],
        ["7", "Marrom / Branco"],
        ["8", "Marrom"],
    ], col_widths=[2, 6])
    add_para(doc, "\"Truque de memorização: começa laranja, azul no meio, marrom no fim. Verde 'quebra' a ordem do azul — o verde/branco fica no 3 e o verde no 6. Se você fixar essa quebra, decora o resto.\"", italic=True, color=COR_PRIMARIA)
    add_para(doc, "Peça pra galera escrever no caderno. Peça pra dizer em voz alta 3 vezes. É de decorar mesmo, não tem jeito. Vai render o resto da vida profissional.")

    add_heading(doc, "T568A — 3 min", level=2)
    add_para(doc, "\"Padrão A troca o par verde com o par laranja. Fica assim:\"", italic=True)
    add_table(doc, [
        ["Pino", "Cor"],
        ["1", "Verde / Branco"],
        ["2", "Verde"],
        ["3", "Laranja / Branco"],
        ["4", "Azul"],
        ["5", "Azul / Branco"],
        ["6", "Laranja"],
        ["7", "Marrom / Branco"],
        ["8", "Marrom"],
    ], col_widths=[2, 6])
    add_para(doc, "\"Se quiser lembrar de A: o padrão A é comum nos EUA em rede de telefonia unificada. Você não vai usar. Mas se aparecer um cabo com verde no lugar do laranja, agora sabe o que é.\"", italic=True)

    add_heading(doc, "Cabo direto × cabo cruzado — 4 min", level=2)
    add_bullet(doc, "Cabo direto: T568B nas duas pontas (ou A nas duas). É o que você vai fazer hoje. Serve pra quase tudo.")
    add_bullet(doc, "Cabo cruzado: T568B numa ponta e T568A na outra. Servia pra ligar dois computadores direto sem roteador. Roteadores modernos fazem essa correção sozinhos (auto MDI-X). Vocês quase nunca vão precisar fazer cabo cruzado — mas precisam saber que existe.")

    add_hr(doc)

    # BLOCO 4
    add_heading(doc, "BLOCO 4 — DEMONSTRAÇÃO DO PROFESSOR · 20 MIN")

    add_heading(doc, "Instrução geral", level=2)
    add_para(doc, "Todo mundo em volta da bancada. Você faz um cabo do zero, comentando cada passo. Vai devagar. Deixe erro acontecer se acontecer — ensina mais que o passo perfeito.")

    add_heading(doc, "Os 12 passos, um por um", level=2)

    add_para(doc, "Passo 1 — Cortar o cabo no tamanho", bold=True)
    add_para(doc, "\"Meço com folga. Se preciso de 3 metros, corto 3,20. Sobra é melhor que faltar 5 cm.\"", italic=True)

    add_para(doc, "Passo 2 — Decapar 3 cm da ponta", bold=True)
    add_para(doc, "\"Aperto pouco. Gira, gira, puxa. Não pode cortar o fio de dentro. Se a lâmina do alicate marca todos os fios, tá apertando demais.\"", italic=True)

    add_para(doc, "Passo 3 — Separar os 4 pares", bold=True)
    add_para(doc, "\"Puxo os pares separando um do outro. Ainda não destrancei — só afastei.\"", italic=True)

    add_para(doc, "Passo 4 — Destrançar CADA par SÓ O NECESSÁRIO", bold=True)
    add_para(doc, "\"Destranço 1,5 cm de cada par. Nem mais nem menos. Se destrançar demais, o cabo perde qualidade. Esse é o erro número um.\"", italic=True, color=COR_ALERTA)

    add_para(doc, "Passo 5 — Ordenar as cores no padrão T568B", bold=True)
    add_para(doc, "\"Laranja/branco, laranja, verde/branco, azul, azul/branco, verde, marrom/branco, marrom. Digo em voz alta enquanto ordeno.\"", italic=True)

    add_para(doc, "Passo 6 — Cortar todos os fios no mesmo tamanho", bold=True)
    add_para(doc, "\"1,5 cm da capa. Corte reto. Se um fio ficar mais curto, ele não vai encostar no fundo do conector — e o cabo não funciona.\"", italic=True, color=COR_ALERTA)

    add_para(doc, "Passo 7 — Encaixar no conector RJ45", bold=True)
    add_para(doc, "\"O conector vai com o clipe pra baixo (ou pra você, depende de como está segurando). Os fios entram na sequência que você já ordenou.\"", italic=True)

    add_para(doc, "Passo 8 — Empurrar até o fundo", bold=True)
    add_para(doc, "\"Os fios têm que encostar no fundo do conector. A capa do cabo tem que entrar dentro do conector também — isso segura o cabo depois.\"", italic=True)

    add_para(doc, "Passo 9 — Conferir antes de crimpar", bold=True)
    add_para(doc, "\"Olho contra a luz. Vejo se cada cor tá no pino certo. Vejo se todos os fios encostaram no fundo. Se algo estiver errado, ainda dá pra retirar e fazer de novo. Depois de crimpar, o conector foi.\"", italic=True, color=COR_PRIMARIA)

    add_para(doc, "Passo 10 — Colocar no alicate crimpador", bold=True)
    add_para(doc, "\"Encaixo o conector na boca de 8 pinos. Firmeza — ele tem só uma posição certa.\"", italic=True)

    add_para(doc, "Passo 11 — Apertar com força", bold=True)
    add_para(doc, "\"Aperto até o alicate travar. Vou ouvir um clique. Se apertar de leve, o pino não fura o fio — vai dar mau contato.\"", italic=True, color=COR_ALERTA)

    add_para(doc, "Passo 12 — Testar", bold=True)
    add_para(doc, "\"Se tiver testador, uso. Se não, ligo entre dois aparelhos e vejo se conectam. Só considero pronto depois do teste — cabo que não passou no teste é lixo.\"", italic=True, color=COR_PRIMARIA)

    add_para(doc, "\"Vou fazer mais um agora — quem quiser, faz junto comigo, passo a passo.\"", italic=True)

    add_hr(doc)
    add_heading(doc, "INTERVALO · 15 MIN", color=COR_CINZA)
    add_hr(doc)

    # BLOCO 5
    add_heading(doc, "BLOCO 5 — PRÁTICA · CADA ALUNO CRIMPA · 55 MIN")

    add_heading(doc, "Distribuição do material — 5 min", level=2)
    add_bullet(doc, "Cada aluno pega ~40 cm de cabo")
    add_bullet(doc, "Cada aluno pega 2 conectores RJ45 (um pra sair errado, um pra sair certo — é a expectativa mínima)")
    add_bullet(doc, "Alicates ficam pra circular — se são poucos, faça 3-4 estações de crimpagem")

    add_heading(doc, "Ordem sugerida por aluno — 45 min", level=2)
    add_bullet(doc, "Decapa uma ponta (5 min de cuidado)")
    add_bullet(doc, "Separa, destrança, ordena cores, corta tamanho (10 min)")
    add_bullet(doc, "Encaixa no conector e confere (5 min — se tá tudo certo, chama você antes de crimpar)")
    add_bullet(doc, "Crimpa (2 min)")
    add_bullet(doc, "Repete pra outra ponta (repete tudo — 20 min)")
    add_bullet(doc, "Espera o testador chegar (3 min)")

    add_heading(doc, "Enquanto isso, você circula — 45 min", level=2)
    add_para(doc, "Papel do professor:")
    add_bullet(doc, "Passa em cada aluno pelo menos duas vezes")
    add_bullet(doc, "Antes de o aluno crimpar, olha a sequência de cor com ele (essa é a hora de corrigir sem gastar conector)")
    add_bullet(doc, "Se o aluno destrançou demais, mande refazer aquela ponta — não vale apertar e torcer pra dar certo")
    add_bullet(doc, "Elogia o cabo bem feito. Diga alto pra turma inteira ouvir: \"olha o cabo do Fulano, capa entrou direitinho\"")
    add_bullet(doc, "Se um aluno errar 3 vezes seguidas, pare, respire, faça de novo do zero DEVAGAR com ele. Não vale ficar rebimbocando o mesmo cabo de 20 cm.")

    add_heading(doc, "Regra do conector", level=2)
    add_para(doc, "\"Cada aluno tem direito a 4 conectores. Depois disso, pergunta pra mim antes de pegar outro. Não é castigo — é pra te fazer parar e ver o que tá errando antes de gastar peça à toa.\"", italic=True, color=COR_PRIMARIA)

    add_hr(doc)

    # BLOCO 6
    add_heading(doc, "BLOCO 6 — TESTE E CORREÇÃO · 25 MIN")

    add_heading(doc, "Se você tem testador — 20 min", level=2)
    add_para(doc, "Passa pelas mesas. Cada aluno testa seu cabo. LED aceso em sequência = perfeito. LED faltando em algum pino = fio não encostou. LED em posição errada = trocou cor.")
    add_para(doc, "Anote no quadro:")
    add_bullet(doc, "Quantos passaram na primeira")
    add_bullet(doc, "Quantos precisaram refazer")
    add_bullet(doc, "Comemore. Não faça vergonha de ninguém.")

    add_heading(doc, "Se você NÃO tem testador — 20 min", level=2)
    add_para(doc, "Alternativa 1: ligue o cabo entre 2 computadores/roteadores. Se o LED da porta acender, tá bom. Se não acender, algum fio não fez contato.")
    add_para(doc, "Alternativa 2: ligue entre roteador e computador. Se o computador receber IP (Windows: Painel de rede → Ver detalhes), o cabo tá bom.")

    add_heading(doc, "Quem falhou refaz — 5 min", level=2)
    add_para(doc, "Quem passou de primeira ajuda quem tá refazendo. Vira monitor por 5 min. Fixa o conteúdo pros dois lados.")

    add_hr(doc)

    # BLOCO 7
    add_heading(doc, "BLOCO 7 — ANÁLISE COLETIVA DE ERROS · 20 MIN")

    add_heading(doc, "A rodada dos cabos ruins", level=2)
    add_para(doc, "Pegue 3-4 cabos que não passaram no teste (peça permissão pros donos). Mostre pra turma um por vez. Antes de dizer o que tá errado, pergunte à turma. Deixe eles descobrirem.")
    add_para(doc, "Erros que provavelmente vão aparecer:")

    add_para(doc, "Erro 1 — Destrançou demais", bold=True, color=COR_ALERTA)
    add_para(doc, "Como reconhecer: você vê os pares desmontados por mais de 2 cm dentro do cabo. Consequência: sinal fica ruim, cabo passa no teste curto mas falha em 30 metros.")

    add_para(doc, "Erro 2 — Fios em tamanhos diferentes", bold=True, color=COR_ALERTA)
    add_para(doc, "Como reconhecer: olhe o conector contra a luz. Se algum fio parece mais curto que os outros, ele não tá encostando no fundo. Consequência: um pino sem sinal, cabo passa parcial no teste.")

    add_para(doc, "Erro 3 — Cor trocada", bold=True, color=COR_ALERTA)
    add_para(doc, "Como reconhecer: olhe a sequência das cores por transparência. Se não é laranja/br, laranja, verde/br, azul... algo tá fora. Consequência: cabo não funciona ou funciona só pra alguns usos.")

    add_para(doc, "Erro 4 — Capa não entrou no conector", bold=True, color=COR_ALERTA)
    add_para(doc, "Como reconhecer: você vê os fios coloridos aparecendo fora do conector. Consequência: primeiro puxão que der, o cabo solta do conector. Reinstalação garantida.")

    add_para(doc, "Erro 5 — Não apertou o suficiente", bold=True, color=COR_ALERTA)
    add_para(doc, "Como reconhecer: puxe cada fio de leve. Se sair, não apertou. Consequência: mau contato intermitente — pior tipo de defeito, porque parece resolver e volta.")

    add_para(doc, "Erro 6 — Cortou fio interno com o estilete", bold=True, color=COR_ALERTA)
    add_para(doc, "Como reconhecer: um fio partido dentro da capa, geralmente perto do conector. Consequência: pino morto. Só descobre no teste.")

    add_para(doc, "\"Todos esses eu já fiz. Você vai fazer. O objetivo é reconhecer pra evitar da próxima.\"", italic=True, color=COR_PRIMARIA)

    add_hr(doc)

    # BLOCO 8
    add_heading(doc, "BLOCO 8 — DUELO RELÂMPAGO · 15 MIN")
    add_para(doc, "Turma em dois times. Resposta em 5 segundos. Alterna.")

    add_heading(doc, "12 perguntas de reserva", level=2)
    add_bullet(doc, "Pino 1 do padrão T568B tem que cor? — Laranja / Branco")
    add_bullet(doc, "Pino 8 é qual cor? — Marrom")
    add_bullet(doc, "Quantos fios tem dentro do cabo? — 8")
    add_bullet(doc, "Quantos pares? — 4")
    add_bullet(doc, "Se destrançar demais o par, o que acontece? — Cabo perde qualidade / sinal ruim")
    add_bullet(doc, "Padrão B ou padrão A é o usado no Brasil? — B")
    add_bullet(doc, "Diferença entre A e B? — Troca o par verde com o par laranja")
    add_bullet(doc, "Cabo direto tem o mesmo padrão nas duas pontas? — Sim")
    add_bullet(doc, "Pra que serve cabo cruzado? — Ligar dois computadores direto sem roteador")
    add_bullet(doc, "Quanto decapar da capa antes de ordenar as cores? — Cerca de 3 cm")
    add_bullet(doc, "O que segura o cabo depois de crimpado, pra não soltar? — A capa dentro do conector")
    add_bullet(doc, "Como você sabe que apertou o suficiente? — Ouviu o clique do alicate travando")

    add_hr(doc)

    # BLOCO 9
    add_heading(doc, "BLOCO 9 — FECHAMENTO E TAREFA · 10 MIN")

    add_heading(doc, "Quiz de saída", level=2)
    add_bullet(doc, "Diga a sequência do T568B começando pelo pino 1")
    add_bullet(doc, "Se o cabo passa no teste mas cai depois de instalado, o erro mais provável foi qual?")
    add_bullet(doc, "Por que não pode destrançar demais na ponta?")
    add_bullet(doc, "O que você faria se percebesse que a cor tá trocada, antes de crimpar?")
    add_bullet(doc, "Cabo cruzado ainda faz sentido hoje?")

    add_heading(doc, "A tarefa", level=2)
    add_para(doc, "\"Vocês vão levar pra casa o cabo de vocês (o bom, o que passou no teste). Se derem uma volta com ele, guardem enrolado. Na próxima noite, cada um traz seu cabo — vamos testar em situações diferentes: cabo dobrado, cabo esticado, cabo perto de motor. Quem trouxer o cabo, participa direto do experimento. Quem esquecer, participa como observador.\"", italic=True)
    add_para(doc, "\"Bônus: quem quiser praticar em casa, leva mais um conector pra crimpar sozinho. Traga o cabo pronto — a gente compara com o feito na aula.\"", italic=True, bold=True)

    add_heading(doc, "Anuncie a próxima", level=2)
    add_para(doc, "\"Na próxima noite: testando cabos e resolvendo cabos ruins. A gente vai pegar cabos de propósito estragados — alguns que eu vou trazer, alguns que vocês já viram hoje — e vamos aprender a diagnosticar sem chutar. Vai ter caso real que aparece no atendimento.\"", italic=True)

    add_hr(doc)

    # SE PERGUNTAREM
    add_heading(doc, "SE PERGUNTAREM")

    add_para(doc, "\"E se eu não quiser decorar cor?\"", bold=True)
    add_para(doc, "> \"Você vai precisar por causa do serviço. Todo conector já tem os pinos, todo alicate espera a sequência certa. Não tem atalho. Mas depois do primeiro mês, decora sozinho.\"", italic=True)

    add_para(doc, "\"Cabo já pronto é melhor que crimpado por mim?\"", bold=True)
    add_para(doc, "> \"Cabo de fábrica geralmente é mais confiável (a máquina não erra) e mais bonito. Mas você não controla o comprimento. Numa instalação real, o comprimento perfeito compensa o pequeno risco do cabo crimpado à mão. E você economiza dinheiro do cliente.\"", italic=True)

    add_para(doc, "\"Posso emendar cabo?\"", bold=True)
    add_para(doc, "> \"Emenda de cabo par trançado dá pra fazer com peça chamada \"emenda RJ45 fêmea-fêmea\". Funciona, mas cada emenda perde um pouco de sinal. Cabo novo é sempre melhor. Só emende se realmente não puder passar cabo novo.\"", italic=True)

    add_para(doc, "\"O conector tem lado?\"", bold=True)
    add_para(doc, "> \"Tem. O clipe (a alavanquinha que trava) fica de um lado só. Você sempre orienta o conector com o clipe pra baixo (ou pra você) — assim quando encaixar na tomada, o clipe trava.\"", italic=True)

    add_para(doc, "\"Cabo torto funciona igual cabo reto?\"", bold=True)
    add_para(doc, "> \"Se as cores estão certas e o crimp tá firme, funciona igual. Bonito ou feio é estética. Mas cabo feio no cliente conta contra você como profissional.\"", italic=True)

    add_hr(doc)

    # DEPOIS DA AULA
    add_heading(doc, "DEPOIS DA AULA")
    add_bullet(doc, "Anote quantos cabos passaram no teste de primeira. É seu indicador da noite.")
    add_bullet(doc, "Quem crimpou bem sem esforço aparente — puxe pra ajudar nas próximas noites (monitor natural).")
    add_bullet(doc, "Quem errou nos primeiros 4 conectores — verifique se o alicate não estava com problema.")
    add_bullet(doc, "Se sobrou conector, guarde. Se acabou, reponha antes da Noite 5.")

    add_hr(doc)

    # DESCRITIVO
    add_heading(doc, "DESCRITIVO PARA O DIÁRIO")
    add_para(doc, "Desenvolvi atividade prática de montagem e conectorização de cabos par trançado, apresentando os padrões de pinagem T568A e T568B, com ênfase no padrão T568B utilizado como referência no Brasil. Trabalhei anatomia interna do cabo (composição em quatro pares trançados) e apresentei o instrumental básico de conectorização: alicate crimpador, ferramenta de decapagem e testador de cabos. Coordenei demonstração completa da sequência de conectorização (corte, decapagem, separação de pares, destrançamento controlado, ordenação de cores, corte de comprimento, encaixe no conector, inspeção e crimpagem), seguida de atividade prática individual em que cada estudante confeccionou e testou seu próprio cabo de rede. Realizei análise coletiva dos erros mais comuns identificados nos cabos confeccionados (destrançamento excessivo, fios em comprimentos desiguais, cores trocadas, capa externa fora do conector, crimpagem incompleta, corte acidental de condutor interno) como recurso de fixação e de preparação para atividades futuras de diagnóstico.")

    add_hr(doc)

    # CAIXA DE FERRAMENTAS DIDÁTICAS
    add_heading(doc, "CAIXA DE FERRAMENTAS DIDÁTICAS", color=COR_DESTAQUE)
    add_para(doc, "Se sobrar tempo, ou pra usar em outra noite:", italic=True, color=COR_CINZA)

    add_para(doc, "1. Contest de crimpagem", bold=True)
    add_para(doc, "Quem crimpar 2 cabos bons no tempo dado ganha um brinde simbólico. Alonga a prática e cria motivação sem competição tóxica. 20-30 min.")

    add_para(doc, "2. Cabo cruzado", bold=True)
    add_para(doc, "Depois do cabo direto, quem terminou faz um cabo cruzado (A numa ponta, B na outra). Testa. Prova que a diferença é só 4 pinos. 15-20 min.")

    add_para(doc, "3. Cabo transparente", bold=True)
    add_para(doc, "Se você conseguir conectores transparentes (existem), o aluno crimpa um e vê as cores por dentro. Fica ótimo pra fotografar e pra usar como referência depois. 10 min extras por aluno.")

    add_para(doc, "4. Autopsia do cabo defeituoso", bold=True)
    add_para(doc, "Pegue um cabo que falhou. Corte o conector com o alicate e a turma investiga por dentro pra descobrir o que deu errado. Ensina diagnóstico. 15 min.")

    add_para(doc, "5. Sabotagem controlada", bold=True)
    add_para(doc, "Você crimpa um cabo bom e depois dá uma sabotada minúscula (aperta demais a capa, dobra num ponto). Turma tem que descobrir. Prepara terreno pra Noite 5. 15 min.")

    add_para(doc, "6. Reprise do padrão", bold=True)
    add_para(doc, "Ao vivo, você escreve as cores do T568B com a turma dizendo em coro. Depois escreve o T568A. Depois só a diferença entre os dois. Fixa por repetição. 5 min sempre que a turma ficar dispersa.")

    path = os.path.join(OUTDIR, "Roteiro - Redes 4 - Cabo na mao e crimpagem.docx")
    doc.save(path)
    return path


# ============================================================
# SLIDES
# ============================================================
def build_slides():
    prs = new_presentation()
    P, D, W, G, K, R, LBG = PPalette.P, PPalette.D, PPalette.W, PPalette.G, PPalette.K, PPalette.R, PPalette.LBG

    # 1: Capa
    add_slide_cover(prs, "NOITE 4", "CABO", "NA MÃO",
                    "Padrão T568B, alicate e crimpagem",
                    "Redes de Computadores · Curso FIC · CEJA Itapiranga – SC")

    # 2: Retomada + Objetivo
    s = add_slide(prs)
    add_title_block(s, "A META DE HOJE", "Um cabo por aluno, crimpado por você, testado e funcionando")
    add_text(s, [
        "Não é aula pra decorar teoria.",
        "É aula pra sujar a mão e sentir a ferramenta.",
        "",
        "Vai ter dedo doendo. Vai ter conector desperdiçado.",
        "Vai ter cabo torto — isso é NORMAL.",
        "",
        "Todo técnico crimpou cabo errado no início.",
        "O que separa o profissional é a quantidade de conectores que ele já queimou.",
    ], 0.7, 2.8, 11.9, 4.0, size=18, color=K)
    add_footer(s, FOOTER)

    # 3: Pergunta da noite
    s = add_slide(prs)
    add_title_block(s, "A PERGUNTA DA NOITE", "Você conseguiria fazer o cabo?")
    add_text(s, [
        "O cabo que passa entre seu roteador e sua TV,",
        "aquele com o conector transparente na ponta —",
        "",
        "você conseguiria fazer um igual, do zero?",
        "",
        "Hoje o objetivo é: pode dizer que sim.",
    ], 0.7, 2.8, 11.9, 4.0, size=22, color=K)
    add_footer(s, FOOTER)

    # 4: Anatomia do cabo
    s = add_slide(prs)
    add_title_block(s, "POR DENTRO DO CABO", "8 fios em 4 pares — e uma capa que protege tudo")
    add_text(s, [
        "•  Capa externa: protege o resto",
        "•  Fio de nylon fininho no meio: ajuda a puxar a capa",
        "•  8 fios coloridos, em 4 pares",
        "•  Pares: laranja, verde, azul, marrom",
        "•  Cada par é trançado sobre si — não é enfeite",
    ], 0.7, 2.8, 11.9, 3.5, size=18, color=K)
    add_text(s, "\"Branco\" na verdade é branco listrado com a cor do par. Prestem atenção nas listras.", 0.7, 6.4, 11.9, 0.5, size=16, italic=True, color=D, align=PP_ALIGN.CENTER)
    add_footer(s, FOOTER)

    # 5: Ferramentas
    s = add_slide(prs)
    add_title_block(s, "AS FERRAMENTAS", "O alicate crimpador faz quase tudo")
    add_text(s, [
        "Alicate crimpador",
        "  • Boca de 8 pinos (RJ45)",
        "  • Boca de 6 pinos (RJ11) — não usamos hoje",
        "  • Lâmina de corte",
        "  • Decapador (nem todo alicate tem)",
        "",
        "Outras ferramentas",
        "  • Estilete ou tesoura (alternativa ao decapador)",
        "  • Régua ou trena",
        "  • Testador de cabo (a estrela do teste)",
    ], 0.7, 2.8, 11.9, 4.2, size=16, color=K)
    add_footer(s, FOOTER)

    # 6: T568B
    s = add_slide(prs)
    add_title_block(s, "O PADRÃO T568B", "Este é o que a gente usa. Decore.")
    rows = [
        ["Pino", "Cor"],
        ["1", "Laranja / Branco"],
        ["2", "Laranja"],
        ["3", "Verde / Branco"],
        ["4", "Azul"],
        ["5", "Azul / Branco"],
        ["6", "Verde"],
        ["7", "Marrom / Branco"],
        ["8", "Marrom"],
    ]
    add_table_slide(s, rows, top=2.8, height=4.0, col_widths=[3, 8.9], size=15)
    add_footer(s, FOOTER)

    # 7: Truque de memória
    s = add_slide(prs)
    add_title_block(s, "TRUQUE DE MEMÓRIA", "Como decorar sem sofrer")
    add_text(s, [
        "•  Começa laranja, termina marrom",
        "•  Azul fica no meio",
        "•  Verde \"quebra\" a ordem: verde/branco no 3, verde no 6",
        "",
        "Se você fixar essa quebra, decora o resto.",
        "",
        "Diga em voz alta 3 vezes agora:",
        "  laranja/br, laranja, verde/br, azul, azul/br, verde, marrom/br, marrom",
    ], 0.7, 2.8, 11.9, 4.2, size=17, color=K)
    add_footer(s, FOOTER)

    # 8: T568A
    s = add_slide(prs)
    add_title_block(s, "T568A", "Existe. Você não vai usar. Mas precisa reconhecer.")
    add_text(s, [
        "Padrão A troca o par verde com o par laranja.",
        "",
        "Fica: verde/br, verde, laranja/br, azul, azul/br, laranja, marrom/br, marrom",
        "",
        "Comum nos EUA em rede unificada com telefonia.",
        "Aqui você quase não vê. Mas se aparecer, agora sabe o que é.",
    ], 0.7, 2.8, 11.9, 3.5, size=18, color=K)
    add_footer(s, FOOTER)

    # 9: Cabo direto x cruzado
    s = add_slide(prs)
    add_title_block(s, "DIRETO × CRUZADO", "Dois cabos diferentes com a mesma aparência")
    add_text(s, [
        "CABO DIRETO",
        "  Mesmo padrão nas duas pontas (B-B ou A-A)",
        "  É o que você vai fazer hoje",
        "  Serve pra 99% dos casos",
        "",
        "CABO CRUZADO",
        "  Uma ponta B, outra A",
        "  Servia pra ligar dois computadores direto (sem roteador)",
        "  Roteadores modernos fazem essa correção sozinhos",
        "  Você quase nunca vai precisar — mas precisa saber que existe",
    ], 0.7, 2.8, 11.9, 4.2, size=16, color=K)
    add_footer(s, FOOTER)

    # 10: 12 passos - visão geral
    s = add_slide(prs)
    add_title_block(s, "OS 12 PASSOS", "Do rolo de cabo ao cabo testado")
    add_text(s, [
        "1.  Cortar o cabo no tamanho",
        "2.  Decapar 3 cm da ponta",
        "3.  Separar os 4 pares",
        "4.  Destrançar cada par SÓ o necessário",
        "5.  Ordenar cores no T568B",
        "6.  Cortar todos os fios no mesmo tamanho",
        "7.  Encaixar no conector",
        "8.  Empurrar até o fundo (com a capa dentro do conector)",
        "9.  CONFERIR antes de crimpar",
        "10. Colocar no alicate crimpador",
        "11. Apertar com força até ouvir o clique",
        "12. Testar",
    ], 0.7, 2.8, 11.9, 4.2, size=14, color=K)
    add_footer(s, FOOTER)

    # 11: Decapar (detalhe)
    s = add_slide(prs)
    add_title_block(s, "DETALHE · DECAPAR", "3 cm — nem mais nem menos")
    add_text(s, [
        "•  Aperto pouco o alicate",
        "•  Gira, gira, puxa",
        "•  Não pode cortar o fio de dentro",
        "",
        "Como saber se apertou demais?",
        "  Se a lâmina marcou todos os fios, tá apertando forte demais.",
    ], 0.7, 2.8, 11.9, 3.5, size=18, color=K)
    add_footer(s, FOOTER)

    # 12: Destrançar
    s = add_slide(prs)
    add_title_block(s, "DETALHE · DESTRANÇAR", "SÓ o necessário — este é o ERRO Nº 1")
    add_text(s, [
        "•  Destrance 1,5 cm de cada par",
        "•  NÃO mais que isso",
        "",
        "Por que importa tanto?",
        "  A trança protege contra interferência.",
        "  Destrançou demais → cabo perde qualidade.",
        "  Pode passar no teste de bancada e falhar em 30 metros.",
    ], 0.7, 2.8, 11.9, 3.5, size=18, color=K)
    add_footer(s, FOOTER)

    # 13: Ordenar cores
    s = add_slide(prs)
    add_title_block(s, "DETALHE · ORDENAR CORES", "Padrão T568B")
    add_text(s, [
        "Sequência (diga em voz alta enquanto ordena):",
        "",
        "  laranja/br  →  laranja  →  verde/br  →  azul",
        "  →  azul/br  →  verde  →  marrom/br  →  marrom",
        "",
        "Se você errar aqui, tudo depois vai errado.",
        "Não tenha pressa neste passo.",
    ], 0.7, 2.8, 11.9, 4.0, size=17, color=K)
    add_footer(s, FOOTER)

    # 14: Cortar tamanho igual
    s = add_slide(prs)
    add_title_block(s, "DETALHE · CORTAR TAMANHO IGUAL", "1,5 cm da capa — todos iguais")
    add_text(s, [
        "•  Corte reto, uma única passada do alicate",
        "•  Todos os 8 fios do mesmo tamanho",
        "",
        "Se um fio ficar mais curto que os outros,",
        "ele não vai encostar no fundo do conector.",
        "",
        "Resultado: um pino sem sinal.",
    ], 0.7, 2.8, 11.9, 3.5, size=18, color=K)
    add_footer(s, FOOTER)

    # 15: Encaixar no conector
    s = add_slide(prs)
    add_title_block(s, "DETALHE · ENCAIXAR", "O clipe fica pra baixo")
    add_text(s, [
        "•  Conector com o clipe pra baixo",
        "•  Fios entram na ordem que você já organizou",
        "•  Empurre até TODOS encostarem no fundo",
        "•  A CAPA tem que entrar dentro do conector",
        "",
        "A capa dentro do conector é o que segura o cabo depois.",
        "Se ela ficou de fora, o cabo solta ao primeiro puxão.",
    ], 0.7, 2.8, 11.9, 4.0, size=17, color=K)
    add_footer(s, FOOTER)

    # 16: Conferir antes de crimpar
    s = add_slide(prs)
    add_title_block(s, "DETALHE · CONFERIR", "Este é o passo mais importante da noite")
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PInches(0.7), PInches(2.8), PInches(11.9), PInches(3.5))
    box.fill.solid()
    box.fill.fore_color.rgb = W
    box.line.color.rgb = P
    box.line.width = Emu(25000)
    add_text(s, [
        "Antes de crimpar:",
        "",
        "  ✓  Olho o conector contra a luz",
        "  ✓  Vejo se cada cor tá no pino certo",
        "  ✓  Vejo se todos os fios encostaram no fundo",
        "  ✓  Vejo se a capa entrou dentro do conector",
        "",
        "Se algo estiver errado, tira e faz de novo.",
        "Depois de crimpar, o conector foi.",
    ], 1.0, 3.0, 11.5, 3.0, size=17, color=K, bold=False)
    add_footer(s, FOOTER)

    # 17: Crimpar
    s = add_slide(prs)
    add_title_block(s, "DETALHE · CRIMPAR", "Aperto firme até o clique")
    add_text(s, [
        "•  Encaixo o conector na boca de 8 pinos do alicate",
        "•  Firmeza — o conector tem só uma posição certa",
        "•  Aperto até o alicate travar",
        "•  Vou ouvir um clique",
        "",
        "Se apertar de leve, os pinos não furam os fios.",
        "Resultado: mau contato. Cabo parece pronto e não funciona.",
    ], 0.7, 2.8, 11.9, 4.0, size=17, color=K)
    add_footer(s, FOOTER)

    # 18: PRÁTICA
    s = add_slide(prs)
    add_title_block(s, "PRÁTICA · 55 MIN", "Cada aluno faz seu cabo")
    add_text(s, [
        "•  40 cm de cabo por pessoa",
        "•  2 conectores RJ45 (esperando um erro)",
        "•  Alicate roda entre as bancadas",
        "",
        "Antes de crimpar, chame o professor",
        "     — a hora de corrigir cor é ANTES do clique.",
        "",
        "Regra do conector",
        "     — 4 conectores por aluno.",
        "     — Depois disso, pergunta antes de pegar outro.",
    ], 0.7, 2.8, 11.9, 4.2, size=16, color=K)
    add_footer(s, FOOTER)

    # 19: Testar
    s = add_slide(prs)
    add_title_block(s, "COMO TESTAR", "Se o teste passou, tá pronto")
    add_text(s, [
        "Com testador de cabo (LED 8 pinos)",
        "  Todos os LEDs em sequência = perfeito",
        "  LED faltando = fio não encostou (refaz)",
        "  LED fora de ordem = trocou cor (refaz)",
        "",
        "Sem testador",
        "  Liga entre dois computadores. Aparece rede? Passou.",
        "  Liga entre roteador e computador. Recebeu IP? Passou.",
        "",
        "Cabo que não passou no teste é lixo — não instale.",
    ], 0.7, 2.8, 11.9, 4.2, size=15, color=K)
    add_footer(s, FOOTER)

    # 20: Erros 1-3
    s = add_slide(prs)
    add_title_block(s, "OS 6 ERROS MAIS COMUNS · 1 A 3", "Todos vocês vão cometer pelo menos um")
    add_text(s, [
        "1.  Destrançou demais",
        "     Vê pares desmontados por mais de 2 cm dentro do cabo",
        "     Consequência: cabo passa no teste curto, falha em 30 m",
        "",
        "2.  Fios em tamanhos diferentes",
        "     Um fio parece mais curto que os outros contra a luz",
        "     Consequência: um pino sem sinal, cabo parcial",
        "",
        "3.  Cor trocada",
        "     Sequência de cores não bate com T568B",
        "     Consequência: cabo não funciona ou funciona só pra alguns usos",
    ], 0.7, 2.8, 11.9, 4.2, size=15, color=K)
    add_footer(s, FOOTER)

    # 21: Erros 4-6
    s = add_slide(prs)
    add_title_block(s, "OS 6 ERROS MAIS COMUNS · 4 A 6", "Os mais chatos")
    add_text(s, [
        "4.  Capa não entrou no conector",
        "     Vê fios coloridos aparecendo fora do conector",
        "     Consequência: cabo solta ao primeiro puxão",
        "",
        "5.  Não apertou o suficiente",
        "     Puxa o fio de leve e ele sai",
        "     Consequência: MAU CONTATO INTERMITENTE (o pior tipo)",
        "",
        "6.  Cortou fio interno com o estilete",
        "     Fio partido dentro da capa perto do conector",
        "     Consequência: pino morto, só descobre no teste",
    ], 0.7, 2.8, 11.9, 4.2, size=15, color=K)
    add_footer(s, FOOTER)

    # 22: Duelo
    s = add_slide(prs)
    add_title_block(s, "DUELO RELÂMPAGO · 15 MIN", "Time A × Time B")
    add_text(s, [
        "As perguntas de hoje:",
        "     •  Cores do T568B por pino",
        "     •  Diferenças entre A e B",
        "     •  Direto vs cruzado",
        "     •  Reconhecimento de erros",
        "     •  Reconhecimento de ferramentas",
        "",
        "Resposta em 5 segundos. Alterna entre times.",
        "Empate = pergunta bônus.",
    ], 0.7, 2.8, 11.9, 4.0, size=17, color=K)
    add_footer(s, FOOTER)

    # 23: Resumo
    s = add_slide(prs)
    add_title_block(s, "O RESUMO DA NOITE", "Cinco ideias que ficam")
    add_text(s, [
        "•  Padrão T568B é o que a gente usa. Decore.",
        "•  Destrançar demais é o erro nº 1. Máximo 1,5 cm por par.",
        "•  Fios do mesmo tamanho. Todos encostando no fundo.",
        "•  A capa DENTRO do conector segura o cabo.",
        "•  Cabo que não passou no teste é lixo.",
    ], 0.7, 2.8, 11.9, 4.0, size=19, color=K)
    add_footer(s, FOOTER)

    # 24: Fechando
    s = add_slide(prs)
    add_title_block(s, "FECHANDO", "Cinco perguntas para conferir")
    add_text(s, [
        "1. Diga o T568B começando pelo pino 1.",
        "2. Se o cabo passa no teste mas cai depois de instalado — qual foi o erro?",
        "3. Por que não pode destrançar demais na ponta?",
        "4. Cor trocada — o que fazer se você perceber antes de crimpar?",
        "5. Cabo cruzado ainda faz sentido hoje? Quando?",
    ], 0.7, 2.8, 11.9, 4.0, size=17, color=K)
    add_footer(s, FOOTER)

    # 25: O que fica
    s = add_slide(prs)
    add_title_block(s, "O QUE FICA DESTA NOITE", "Você tem um cabo feito por você")
    add_text(s, [
        "•  Você conhece o padrão T568B de cor",
        "•  Você sabe usar o alicate crimpador",
        "•  Você reconhece os 6 erros mais comuns",
        "•  Você tem um cabo testado, feito com suas mãos",
        "•  Na próxima noite, esse cabo volta — pra testes mais duros.",
    ], 0.7, 2.8, 11.9, 4.0, size=18, color=K)
    add_footer(s, FOOTER)

    # 26: Para a próxima noite
    s = add_slide(prs)
    add_title_block(s, "PARA A PRÓXIMA NOITE", "Testar cabos e resolver cabos ruins")
    add_text(s, [
        "Tarefa desta semana:",
        "     •  Guarde o seu cabo bom (o que passou no teste)",
        "     •  Traga na próxima aula",
        "     •  Vamos testar em situações diferentes:",
        "         cabo dobrado, cabo esticado, cabo perto de motor",
        "",
        "Bônus:",
        "     Quem quiser praticar em casa, leva mais um conector.",
        "     Traga o cabo pronto — vamos comparar.",
        "",
        "Na Noite 5: vamos aprender a diagnosticar sem chutar.",
    ], 0.7, 2.8, 11.9, 4.2, size=15, color=K)
    add_footer(s, FOOTER)

    path = os.path.join(OUTDIR, "Aula - Redes 4 - Cabo na mao e crimpagem.pptx")
    prs.save(path)
    return path


# ============================================================
# FOLHA DO ALUNO (3 páginas)
# ============================================================
def build_folha():
    doc = Document()
    setup_page(doc, margem=Cm(1.8))

    # PÁGINA 1 — Padrão de cores + 12 passos
    cabecalho_documento(
        doc,
        "SEU COMPANHEIRO DE BANCADA",
        f"Redes de Computadores · {NOITE} · Curso FIC · CEJA Itapiranga · 1 de 3",
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Nome: ")
    r.font.size = Pt(11)
    r.bold = True
    r = p.add_run("_________________________________________________     Data: ____ / ____ / ______")
    r.font.size = Pt(11)

    add_para(doc, "Esta folha é sua durante a crimpagem. Consulte quando esquecer a ordem, marque quando conseguir, anote onde errou.", italic=True, size=10, color=COR_CINZA)

    add_heading(doc, "1  Padrão T568B (o que você vai usar)", level=2)
    add_table(doc, [
        ["Pino", "Cor do fio", "Marque quando decorar"],
        ["1", "Laranja / Branco", "☐"],
        ["2", "Laranja", "☐"],
        ["3", "Verde / Branco", "☐"],
        ["4", "Azul", "☐"],
        ["5", "Azul / Branco", "☐"],
        ["6", "Verde", "☐"],
        ["7", "Marrom / Branco", "☐"],
        ["8", "Marrom", "☐"],
    ], col_widths=[2, 8, 5])

    add_para(doc, "Truque: começa laranja, azul no meio, marrom no fim. O verde \"quebra\" a ordem.", italic=True, size=10, color=COR_DESTAQUE, bold=True)

    add_heading(doc, "2  Padrão T568A (só pra saber)", level=2)
    add_para(doc, "Padrão A troca o par verde com o par laranja. Pinos 4, 5, 7, 8 são iguais nos dois.", size=11)
    add_table(doc, [
        ["Pino", "Cor T568A"],
        ["1", "Verde / Branco"],
        ["2", "Verde"],
        ["3", "Laranja / Branco"],
        ["6", "Laranja"],
    ], col_widths=[2, 6])

    add_heading(doc, "3  Os 12 passos", level=2)
    passos = [
        "Cortar o cabo no tamanho",
        "Decapar 3 cm da ponta",
        "Separar os 4 pares",
        "Destrançar cada par (SÓ 1,5 cm)",
        "Ordenar as cores no T568B",
        "Cortar todos os fios no mesmo tamanho (1,5 cm da capa)",
        "Encaixar no conector (clipe pra baixo)",
        "Empurrar até o fundo (capa entra no conector)",
        "CONFERIR contra a luz",
        "Encaixar no alicate crimpador",
        "Apertar até ouvir o clique",
        "Testar",
    ]
    tbl = add_table(doc, [["#", "Passo", "Feito"]] + [[str(i+1), passo, "☐"] for i, passo in enumerate(passos)],
                    col_widths=[1.2, 12, 2])

    # PÁGINA 2 — Diário da crimpagem
    doc.add_page_break()
    cabecalho_documento(
        doc,
        "DIÁRIO DA CRIMPAGEM",
        f"Suas tentativas · {NOITE} · 2 de 3",
    )

    add_para(doc, "Preencha uma linha para cada tentativa. Não é vergonha ter várias — é como se aprende.", italic=True, size=10, color=COR_CINZA)

    add_heading(doc, "Ponta A do cabo", level=2)
    tbl = add_table(doc, [
        ["Tentativa", "Resultado do teste", "O que estava errado", "O que aprendeu"],
        ["1", "", "", ""],
        ["2", "", "", ""],
        ["3", "", "", ""],
        ["4", "", "", ""],
    ], col_widths=[2, 4, 5.5, 5.5])
    for row in tbl.rows[1:]:
        row.height = Cm(1.2)

    add_heading(doc, "Ponta B do cabo", level=2)
    tbl = add_table(doc, [
        ["Tentativa", "Resultado do teste", "O que estava errado", "O que aprendeu"],
        ["1", "", "", ""],
        ["2", "", "", ""],
        ["3", "", "", ""],
        ["4", "", "", ""],
    ], col_widths=[2, 4, 5.5, 5.5])
    for row in tbl.rows[1:]:
        row.height = Cm(1.2)

    add_heading(doc, "Seu cabo passou no teste?", level=2)
    p = doc.add_paragraph()
    r = p.add_run("(   ) Sim, de primeira    (   ) Sim, depois de refazer    (   ) Ainda não — vou levar pra tentar em casa")
    r.font.size = Pt(11)

    add_heading(doc, "Alguma pergunta que ficou?", level=2)
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("_" * 100)
        r.font.size = Pt(11)
        r.font.color.rgb = COR_CINZA

    # PÁGINA 3 — Erros comuns + tarefa
    doc.add_page_break()
    cabecalho_documento(
        doc,
        "OS 6 ERROS MAIS COMUNS",
        f"Reconhecer pra evitar · {NOITE} · 3 de 3",
    )

    add_para(doc, "Marque com X o(s) erro(s) que aconteceram no seu cabo. Sinceridade ajuda a fixar o aprendizado.", italic=True, size=10, color=COR_CINZA)

    erros = [
        ("1. Destrançou demais",
         "Vê os pares desmontados por mais de 2 cm dentro do cabo",
         "Cabo passa no teste curto mas falha em distância maior"),
        ("2. Fios em tamanhos diferentes",
         "Um fio parece mais curto contra a luz do que os outros",
         "Um pino fica sem sinal, cabo funciona pela metade"),
        ("3. Cor trocada",
         "A sequência de cor não bate com o T568B",
         "Cabo não funciona, ou funciona só pra alguns usos"),
        ("4. Capa não entrou no conector",
         "Vê os fios coloridos aparecendo fora do conector",
         "Cabo solta ao primeiro puxão que der"),
        ("5. Não apertou o suficiente",
         "Puxa um fio de leve e ele sai",
         "Mau contato intermitente — pior tipo de defeito"),
        ("6. Cortou fio com o estilete",
         "Fio partido dentro da capa, perto do conector",
         "Pino morto, só descobre depois no teste"),
    ]
    tbl = add_table(doc, [["Aconteceu?", "Erro", "Como reconhecer", "O que causa"]] +
                    [["☐", tit, rec, cau] for tit, rec, cau in erros],
                    col_widths=[2, 3.5, 5.8, 5.7])
    for row in tbl.rows[1:]:
        row.height = Cm(1.5)

    doc.add_paragraph()
    add_heading(doc, "TAREFA DA SEMANA", level=2, color=COR_DESTAQUE)
    add_bullet(doc, "Guarde o seu cabo bom (aquele que passou no teste). Não jogue fora.")
    add_bullet(doc, "Traga na Noite 5 — vamos testar em situações diferentes: dobrado, esticado, perto de motor.")
    add_bullet(doc, "Bônus: leva um conector extra. Pratique em casa. Traga o cabo pronto pra comparar com o da aula.")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Traga esta folha na próxima aula · CEJA Itapiranga – SC")
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = COR_CINZA

    path = os.path.join(OUTDIR, "Folha do Aluno - Redes 4 - Cabo na mao e crimpagem.docx")
    doc.save(path)
    return path


# ============================================================
# RESUMO DO ALUNO
# ============================================================
def build_resumo():
    doc = Document()
    setup_page(doc, margem=Cm(1.8))

    cabecalho_documento(
        doc,
        "CABO NA MÃO — CRIMPAGEM",
        f"Resumo do aluno · {NOITE} · Redes de Computadores · Curso FIC · CEJA Itapiranga",
    )

    add_heading(doc, "A ideia principal", level=2)
    add_para(doc, "O cabo de rede que o cliente compra pronto foi feito da mesma forma que o seu de hoje: com um alicate crimpador, oito fios coloridos e um conector RJ45. Quem sabe fazer, sabe consertar. Quem sabe consertar, atende mais chamados com menos gasto de peça.")

    add_heading(doc, "Padrão T568B — o que a gente usa", level=2)
    add_table(doc, [
        ["Pino", "Cor"],
        ["1", "Laranja / Branco"],
        ["2", "Laranja"],
        ["3", "Verde / Branco"],
        ["4", "Azul"],
        ["5", "Azul / Branco"],
        ["6", "Verde"],
        ["7", "Marrom / Branco"],
        ["8", "Marrom"],
    ], col_widths=[3, 12.5])
    add_para(doc, "Truque: começa laranja, azul no meio, marrom no fim. O verde \"quebra\" a ordem — verde/br no 3 e verde no 6.", italic=True, color=COR_DESTAQUE, bold=True)

    add_heading(doc, "T568A — só pra reconhecer", level=2)
    add_para(doc, "Padrão A troca o par verde com o par laranja. Pinos 4, 5, 7 e 8 são iguais nos dois. Cabo direto = mesmo padrão nas duas pontas (B-B ou A-A). Cabo cruzado = B numa ponta e A na outra (quase não usa mais hoje).")

    add_heading(doc, "Anatomia do cabo", level=2)
    add_bullet(doc, "Capa externa protegendo tudo")
    add_bullet(doc, "Fio de nylon fino no meio, ajuda a puxar a capa quando você decapa")
    add_bullet(doc, "8 fios coloridos, em 4 pares (laranja, verde, azul, marrom)")
    add_bullet(doc, "Cada par é trançado — a trança protege contra interferência")

    add_heading(doc, "Os 12 passos, resumidos", level=2)
    add_bullet(doc, "1-3: Cortar o cabo, decapar 3 cm, separar os 4 pares")
    add_bullet(doc, "4-6: Destrançar só 1,5 cm, ordenar cores T568B, cortar fios no mesmo tamanho")
    add_bullet(doc, "7-9: Encaixar no conector, empurrar até o fundo (a capa entra dentro), CONFERIR contra a luz")
    add_bullet(doc, "10-12: Encaixar no alicate, apertar até o clique, testar")

    add_heading(doc, "Os 6 erros mais comuns", level=2)
    add_table(doc, [
        ["Erro", "Consequência"],
        ["Destrançou demais", "Cabo passa curto, falha em distância maior"],
        ["Fios em tamanhos diferentes", "Um pino sem sinal, cabo parcial"],
        ["Cor trocada", "Cabo não funciona ou funciona só pra alguns usos"],
        ["Capa não entrou no conector", "Cabo solta ao primeiro puxão"],
        ["Não apertou o suficiente", "Mau contato intermitente (o pior tipo)"],
        ["Cortou fio com o estilete", "Pino morto, só descobre no teste"],
    ], col_widths=[5, 10.5])

    add_heading(doc, "Como testar", level=2)
    add_bullet(doc, "Com testador de LED: todos os LEDs em sequência = perfeito. LED faltando = fio não encosta. Fora de ordem = cor trocada.")
    add_bullet(doc, "Sem testador: liga entre dois computadores/roteadores e vê se conectam. Se receber IP, o cabo tá bom.")
    add_bullet(doc, "Cabo que não passou no teste é lixo. Nunca instale.")

    add_heading(doc, "Três frases pra guardar", level=2)
    add_para(doc, "\"Destrançar demais é o erro número um. Máximo 1,5 cm por par.\"", bold=True, color=COR_PRIMARIA)
    add_para(doc, "\"A capa DENTRO do conector é o que segura o cabo depois.\"", bold=True, color=COR_PRIMARIA)
    add_para(doc, "\"Conferir antes de crimpar economiza conector, tempo e paciência.\"", bold=True, color=COR_PRIMARIA)

    add_heading(doc, "Pra se testar em casa", level=2)
    add_bullet(doc, "Diga o T568B começando pelo pino 1, sem olhar.")
    add_bullet(doc, "Se um cabo funciona logo depois de instalado mas cai depois de uma semana, qual erro provavelmente foi cometido?")
    add_bullet(doc, "Cabo direto usa qual sequência em cada ponta?")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Este material é seu. Guarde e volte quando precisar. · CEJA Itapiranga – SC")
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = COR_CINZA

    path = os.path.join(OUTDIR, "Resumo do Aluno - Redes 4 - Cabo na mao e crimpagem.docx")
    doc.save(path)
    return path


# ============================================================
# MAIN
# ============================================================
if __name__ == "__main__":
    print("Gerando Roteiro...")
    print("  →", build_roteiro())
    print("Gerando Slides...")
    print("  →", build_slides())
    print("Gerando Folha do Aluno...")
    print("  →", build_folha())
    print("Gerando Resumo do Aluno...")
    print("  →", build_resumo())
    print("\nPronto.")
