"""Gera o Plano Mensal 3 de Informática (curso de Redes) — 16/10/2026 a 15/11/2026.

Segue o layout do modelo da SED/CEJA (Plano Mensal 2 — SETEMBRO-OUTUBRO 2026):
cabeçalho com brasão, campos em negrito na mesma linha, marcadores ● e Arial 12.
Conteúdo das Noites 7 a 11 do módulo de Redes (DNS, configuração do roteador,
cobertura, segurança e redes móveis).
"""
import os

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTDIR = "/home/nedel/curso"
BRASAO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "brasao_sc.jpeg")
FONTE = "Arial"

CAMPOS = [
    ("COMPONENTE CURRICULAR:", "Informática — Curso FIC em Montador e Reparador de Computadores"),
    ("PROFESSOR:", "Joni César Eidt"),
    ("TEMA GERADOR:", "Projeto de Vida: Construindo o Futuro"),
    ("RECORTE DO COMPONENTE:", "Configurar — do roteador de fábrica à rede pronta, segura e móvel"),
    ("TURMA:", "Ensino Médio – Noturno (Multisseriado)"),
    ("QUESTÃO MOTRIZ:", "Se um vizinho me entregar um roteador novo na caixa, o que eu preciso configurar "
                        "para deixar a internet da casa dele boa, segura e funcionando em todos os cômodos?"),
]

OBJETIVOS = [
    "Explicar a função do serviço de nomes (DNS) na tradução do endereço de um site em endereço de rede;",
    "Ler e interpretar o endereço de um site, identificando o domínio responsável e reconhecendo links fraudulentos;",
    "Distinguir o que a conexão segura (HTTPS) garante e o que ela não garante ao usuário;",
    "Configurar um roteador doméstico a partir das configurações de fábrica: senha de administração, nome "
    "de rede, senha e criptografia, endereço de acesso, distribuição automática e reserva de endereços, "
    "servidor de nomes e rede de convidados;",
    "Salvar e restaurar a configuração do equipamento e registrar o serviço em ficha de entrega;",
    "Medir a intensidade do sinal sem fio e a velocidade entregue, e propor soluções para áreas de "
    "cobertura fraca;",
    "Diferenciar as faixas de 2,4 GHz e 5 GHz quanto a alcance e velocidade;",
    "Aplicar e auditar práticas de segurança na configuração da rede sem fio e orientar o usuário final;",
    "Reconhecer as redes móveis (4G e 5G) e usar o celular e o roteador móvel como ponto de acesso;",
    "Reconhecer na configuração, na manutenção e na orientação ao cliente uma frente concreta de "
    "prestação de serviço.",
]

INTERACAO_INTRO = ("O componente segue integrado ao tema gerador pelo viés do trabalho como projeto de "
                   "vida, e dialoga com as demais áreas nos seguintes pontos:")

INTERACAO = [
    ("LINGUAGENS — ", "leitura crítica de links e de mensagens, identificando as estratégias de "
                      "persuasão usadas em golpes digitais, produção de cartilha de orientação para a "
                      "família e registro técnico do serviço em ficha de entrega (EM13LGG701 e EM13LGG702)."),
    ("MATEMÁTICA — ", "unidades de velocidade de transmissão e de consumo de dados, e comparação entre o "
                      "plano contratado, a velocidade medida perto e longe do roteador e a franquia móvel "
                      "(EM13MAT101)."),
    ("CIÊNCIAS HUMANAS — ", "segurança digital e golpes como questão social, e o uso ético e crítico "
                            "das tecnologias na vida em família e em comunidade (EM13CHS106)."),
    ("CIÊNCIAS DA NATUREZA — ", "propagação das ondas de rádio, faixas de frequência, atenuação por "
                                "paredes e interferência entre redes (EM13CNT103)."),
]

HABILIDADE_IF = ("(EMIFFTP05) Selecionar e mobilizar intencionalmente recursos criativos para resolver "
                 "problemas reais relacionados à produtividade, à colaboração e/ou à comunicação.")

CONTEUDOS = [
    "Nomes na rede: serviço de nomes (DNS); estrutura do endereço de um site (protocolo, domínio e "
    "caminho); conexão segura (HTTPS) e certificado; links fraudulentos e engenharia social.",
    "Configuração do roteador: senha de administração; tela de conexão com o provedor (IP dinâmico e "
    "PPPoE); nome de rede, senha e criptografia WPA2 e WPA3; alteração do endereço de acesso e gateway; "
    "faixa do servidor DHCP e reserva de endereço por MAC; servidores DNS; rede de convidados; cópia "
    "de segurança e restauração.",
    "Cobertura sem fio: intensidade de sinal; teste de velocidade; posicionamento do roteador; faixas "
    "de 2,4 GHz e 5 GHz; canais e interferência de redes vizinhas; repetidor, sistema mesh e ponto de "
    "acesso cabeado.",
    "Segurança em rede: senhas fortes e frases-senha; senha padrão de equipamentos; WPS; firmware; "
    "riscos da rede aberta, da rede pública e da rede falsa; auditoria da configuração; orientação ao "
    "cliente.",
    "Redes móveis: 4G e 5G; sinal e cobertura da operadora; roteamento de internet pelo celular "
    "(hotspot); roteador móvel e roteador com chip; franquia e consumo de dados.",
    "Registro do serviço: ficha de entrega da configuração ao cliente.",
]

METODOLOGIA = [
    "Predominância de atividade prática (mais de sessenta por cento do encontro) e exposição dialogada "
    "em blocos curtos, sempre seguidos de prática;",
    "A situação-problema antecede o conteúdo: cada encontro abre por algo que o estudante traz de casa "
    "— os endereços de sites que usa, a etiqueta do próprio roteador, a velocidade medida no cômodo "
    "mais longe;",
    "Encenação com material físico para tornar visível o que é invisível na rede, como a agenda viva "
    "do serviço de nomes;",
    "Execução em bancada com roteadores reais, partindo das configurações de fábrica, com passo a "
    "passo impresso e rodízio de quem opera o equipamento;",
    "Verificação objetiva permanente: o aparelho conecta ou não conecta, o endereço reservado aparece "
    "ou não aparece, a rede de convidados enxerga ou não enxerga a rede da casa;",
    "Estações rotativas com laudo: classificação de links, medição de sinal e auditoria de roteador;",
    "Conferência cruzada entre equipes: cada equipe valida a configuração e a ficha de entrega de outra;",
    "Trabalho em equipes com rodízio de funções e tutoria entre pares;",
    "Retomada nos minutos iniciais de cada encontro e atendimento individual nos minutos finais;",
    "Registro do serviço executado em ficha técnica, no mesmo formato adotado nas unidades anteriores.",
]

RECURSOS = [
    "Laboratório de informática e projetor;",
    "Roteadores domésticos para bancada, cabos de rede e régua de tomadas;",
    "Celulares do professor e dos estudantes, com aplicativo de análise de sinal sem fio, teste de "
    "velocidade e compartilhamento de internet (hotspot);",
    "Fatura ou aplicativo do provedor de internet, trazidos pelos próprios estudantes;",
    "Cartolinas e canetões para a produção de cartilha e cartazes;",
    "Material impresso produzido pelo professor: roteiros, folhas do aluno, cartões de links, "
    "passo a passo de configuração, checklist de auditoria e fichas de entrega.",
]

AVALIACAO = [
    "Portfólio profissional: ficha de entrega da configuração do roteador e laudos de cobertura e de "
    "auditoria de segurança;",
    "Produto verificável: roteador configurado que conecta os aparelhos da equipe, restaurado a partir "
    "da cópia de segurança e validado por outra equipe somente com a ficha de entrega;",
    "Atividades práticas de aula, com ficha de acompanhamento por encontro (clínica de links, mapa de "
    "sinal e laudo de auditoria);",
    "Produções coletivas: cartilha anti-golpe para a família e cartaz de diagnóstico de cobertura;",
    "Avaliação prática integradora aplicada em 06/11: configuração completa e auditoria de segurança "
    "de um roteador;",
    "Participação e frequência, com atenção permanente ao risco de evasão;",
    "Recomposição imediata das aprendizagens não atingidas, com repetição do procedimento até a "
    "execução correta.",
]


def _run(p, texto, bold=False):
    r = p.add_run(texto)
    r.bold = bold
    r.font.name = FONTE
    r.font.size = Pt(12)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONTE)
    return r


def _para(doc, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    p.alignment = align
    return p


def campo(doc, rotulo, valor):
    p = _para(doc, space_after=10)
    _run(p, rotulo, bold=True)
    _run(p, "  " + valor)


def titulo_secao(doc, texto, space_before=10):
    p = _para(doc, space_before=space_before)
    p.paragraph_format.keep_with_next = True
    _run(p, texto, bold=True)


def bullet(doc, texto, negrito=None):
    p = _para(doc)
    pf = p.paragraph_format
    pf.left_indent = Cm(1.05)
    pf.first_line_indent = Cm(-0.55)
    pf.tab_stops.add_tab_stop(Cm(1.05))
    _run(p, "●\t")
    if negrito:
        _run(p, negrito, bold=True)
    _run(p, texto)


def sem_bordas(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement(f'w:{lado}')
        b.set(qn('w:val'), 'nil')
        borders.append(b)
    tbl_pr.append(borders)


def cabecalho(doc):
    t = doc.add_table(rows=1, cols=2)
    sem_bordas(t)
    larguras = (Cm(4.3), Cm(12.6))
    for cell, w in zip(t.rows[0].cells, larguras):
        cell.width = w
    c_logo, c_txt = t.rows[0].cells
    p = c_logo.paragraphs[0]
    p.paragraph_format.left_indent = Cm(1.25)
    p.add_run().add_picture(BRASAO, width=Cm(2.5))
    linhas = [
        ("ESTADO DE SANTA CATARINA", True),
        ("Secretaria de Estado da Educação", False),
        ("Centro de Educação de Jovens e Adultos – CEJA", False),
        ("Itapiranga – SC Fone: 3472 5032", False),
    ]
    for i, (txt, b) in enumerate(linhas):
        p = c_txt.paragraphs[0] if i == 0 else c_txt.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(10) if i == 0 else Pt(0)
        _run(p, txt, bold=b)


def gerar():
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.1)
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
    normal = doc.styles['Normal']
    normal.font.name = FONTE
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONTE)

    cabecalho(doc)
    _para(doc, space_after=10)

    p = _para(doc, space_after=10)
    _run(p, "PLANO MENSAL", bold=True)
    for rotulo, valor in CAMPOS:
        campo(doc, rotulo, valor)

    titulo_secao(doc, "OBJETIVOS E METAS:", space_before=0)
    for t in OBJETIVOS:
        bullet(doc, t)

    titulo_secao(doc, "INTERAÇÃO COM AS ÁREAS DO CONHECIMENTO:")
    p = _para(doc)
    _run(p, INTERACAO_INTRO)
    for area, texto in INTERACAO:
        bullet(doc, texto, negrito=area)

    titulo_secao(doc, "HABILIDADE DO ITINERÁRIO FORMATIVO:")
    p = _para(doc)
    _run(p, HABILIDADE_IF)

    titulo_secao(doc, "CONTEÚDOS DO PERÍODO:")
    for t in CONTEUDOS:
        bullet(doc, t)

    titulo_secao(doc, "CAMINHO METODOLÓGICO:")
    for t in METODOLOGIA:
        bullet(doc, t)

    titulo_secao(doc, "RECURSOS:")
    for t in RECURSOS:
        bullet(doc, t)

    _para(doc)
    titulo_secao(doc, "AVALIAÇÃO E MONITORAMENTO:")
    for t in AVALIACAO:
        bullet(doc, t)

    path = os.path.join(OUTDIR, "Plano Mensal 3 - Informatica - OUTUBRO-NOVEMBRO 2026.docx")
    doc.save(path)
    print(path)


if __name__ == "__main__":
    gerar()
