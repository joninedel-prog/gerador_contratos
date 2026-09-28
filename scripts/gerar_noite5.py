"""Gera os 4 arquivos da Noite 5 do curso de Redes.

Noite 5 — Testando cabos e resolvendo cabos ruins
Padrão prática-pesada. Os alunos testam o cabo que crimparam na Noite 4,
depois diagnosticam em equipe uma "clínica" de cabos sabotados pelo professor,
consertam, e fecham com o método de 4 passos aplicado a chamados.
Funciona com ou sem testador de cabo (plano B: LED da porta + velocidade do link + ping).
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from common import (
    COR_PRIMARIA, COR_ALERTA, COR_DESTAQUE, COR_CINZA,
    add_heading, add_para, add_bullet, add_table, add_hr,
    setup_page, cabecalho_documento,
    PPalette, new_presentation, add_slide, add_text, add_title_block,
    add_slide_cover, add_footer, add_table_slide, add_bar,
)
from pptx.enum.text import PP_ALIGN

OUTDIR = "/home/nedel/curso"
NOITE = "Noite 5"
FOOTER = f"Redes de Computadores · {NOITE}"


# Os cabos da clínica — usados no roteiro (preparo e gabarito) e na folha
CLINICA = [
    ("C1", "Pino aberto",
     "Numa ponta, corte o fio 3 (verde/branco) uns 3 mm mais curto antes de crimpar.",
     "LED 3 não acende. Na porta: sem luz, sem conexão.",
     "Refazer a ponta: cortar o conector, decapar de novo, crimpar."),
    ("C2", "Par partido — o \"funciona, mas lento\"",
     "Numa ponta, corte o fio 7 (marrom/branco) curto antes de crimpar.",
     "LED 7 não acende. Na porta: conecta, mas o computador mostra 100 Mbps em vez de 1 Gbps.",
     "Refazer a ponta. Lição: cabo que \"funciona\" pode estar ruim."),
    ("C3", "Cor trocada",
     "Numa ponta, troque de lugar o fio 2 (laranja) com o fio 3 (verde/branco).",
     "LEDs 2 e 3 acendem fora de ordem. Na porta: pode nem conectar, ou conectar com falhas.",
     "Refazer a ponta conferindo o T568B contra a luz."),
    ("C4", "Mau contato",
     "Crimpe com pouca força e deixe a capa fora do conector.",
     "Às vezes passa, às vezes não. No teste do balanço, o ping perde pacote.",
     "Refazer a ponta com a capa dentro e aperto até o clique."),
    ("C5", "Defeito no meio do cabo",
     "Pontas perfeitas. No meio do cabo, esmague um ponto com o alicate ou dê um corte na capa que pegue 1-2 fios.",
     "LED(s) faltando, mas os conectores estão perfeitos.",
     "Procurar o ponto machucado. Cabo curto: trocar. Cabo longo: emenda só se não der pra trocar."),
    ("C6", "Cabo cruzado (pegadinha)",
     "T568B numa ponta, T568A na outra.",
     "LEDs 1↔3 e 2↔6 trocados. Na porta: funciona normal (aparelho moderno se ajusta sozinho).",
     "Nada. Diferente não é defeito — mas etiquete o cabo."),
    ("C7", "Cabo bom (controle)",
     "Um cabo perfeito, igual aos outros por fora.",
     "Tudo em ordem.",
     "Nada. Nem todo cabo suspeito é culpado."),
]


# ============================================================
# ROTEIRO
# ============================================================
def build_roteiro():
    doc = Document()
    setup_page(doc)

    cabecalho_documento(
        doc,
        "TESTANDO CABOS E RESOLVENDO CABOS RUINS",
        f"Roteiro do professor · Redes de Computadores · {NOITE} · Curso FIC · CEJA Itapiranga – SC",
    )

    # A IDEIA
    add_heading(doc, "A IDEIA DESTA NOITE")
    add_para(doc, "Na Noite 4 eles fizeram o cabo. Hoje eles aprendem a desconfiar dele — e a provar se ele é culpado ou inocente. O técnico iniciante troca peça até funcionar. O técnico bom testa, descobre o defeito e troca só o que precisa.")
    add_para(doc, "A frase da noite: \"Cabo não se chuta. Cabo se testa.\"", bold=True, color=COR_PRIMARIA)

    add_heading(doc, "O que se resolve nesta noite", level=2)
    add_para(doc, "Cada aluno sai sabendo testar um cabo de três jeitos (testador, luz da porta e velocidade do link), reconhecendo os defeitos mais comuns pelo resultado do teste, e usando um método de 4 passos pra achar o culpado sem chutar.")

    add_heading(doc, "Esta é a quinta de quinze noites", level=2)
    add_para(doc, "Fecha o bloco de cabos (Noites 3, 4 e 5). A Noite 6 entra em IP e MAC — hoje já aparece o ping como ferramenta, sem explicar o que é IP. Diga só \"é um número que o roteador tem; na próxima noite a gente entende\".")

    add_hr(doc)

    # ANTES DA AULA
    add_heading(doc, "ANTES DA AULA")

    add_heading(doc, "Material da mesa", level=2)
    add_bullet(doc, "Os cabos que os alunos crimparam na Noite 4 (eles devem trazer) + alguns de reserva seus")
    add_bullet(doc, "Os 7 cabos da clínica (C1 a C7), preparados por você — ver abaixo")
    add_bullet(doc, "Testador de cabo de LED (8 pinos) — se tiver. Custa pouco e vale muito nesta noite")
    add_bullet(doc, "Roteador com portas Gigabit (o Archer C5 da bancada) — IMPORTANTE: o TL-WR840N tem portas de 100 Mbps e não serve pro teste de velocidade")
    add_bullet(doc, "1 ou 2 computadores/notebooks com porta de rede (os da sala servem)")
    add_bullet(doc, "Alicate, conectores RJ45 (uns 20 de reserva pros consertos) e um pedaço de cabo do rolo")
    add_bullet(doc, "Fita crepe e caneta pra numerar os cabos")
    add_bullet(doc, "Régua de tomadas")

    add_heading(doc, "Prepare a clínica de cabos (em casa, antes)", level=2)
    add_para(doc, "Faça 7 cabos de 1 a 2 metros. Por fora, todos têm que parecer iguais. Numere com fita crepe (C1 a C7). Guarde este gabarito com você — os alunos NÃO recebem:")
    add_table(doc, [["Cabo", "Defeito", "Como preparar", "O que o teste mostra"]] +
              [[c, d, prep, teste] for c, d, prep, teste, _ in CLINICA],
              col_widths=[1.3, 3.2, 6, 6])
    add_para(doc, "Teste os 7 no roteador e no testador ANTES da aula. Anote a velocidade que o computador mostra em cada um. Se o C2 der \"sem conexão\" em vez de 100 Mbps, tudo bem — algumas placas de rede não baixam a velocidade, simplesmente desistem. Adapte a fala.", italic=True, color=COR_ALERTA)

    add_heading(doc, "Deixe o computador pronto", level=2)
    add_bullet(doc, "Descubra onde aparece a velocidade do link. Windows: Painel de Controle → Central de Rede e Compartilhamento → clique em \"Ethernet\" → a janela de Status mostra \"Velocidade: 1,0 Gbps\"")
    add_bullet(doc, "Atalho: no PowerShell, o comando Get-NetAdapter mostra a coluna LinkSpeed")
    add_bullet(doc, "Descubra o endereço do roteador pro ping: no Prompt, ipconfig → \"Gateway Padrão\" (nos TP-Link costuma ser 192.168.0.1)")
    add_bullet(doc, "Deixe o Wi-Fi do computador DESLIGADO. Senão o ping passa pelo Wi-Fi e engana o teste do cabo")

    add_heading(doc, "Imprima", level=2)
    add_bullet(doc, "Folha do aluno (3 páginas), uma por pessoa")
    add_bullet(doc, "Resumo do aluno, uma por pessoa (pra levar pra casa)")

    add_hr(doc)

    # QUADRO DE TEMPO
    add_heading(doc, "QUADRO DE TEMPO — 210 MINUTOS")
    add_table(doc, [
        ["Bloco", "Tempo", "O que acontece"],
        ["1. Retomada + cabos da Noite 4", "10 min", "Recap relâmpago, cabos na mesa"],
        ["2. Três jeitos de testar um cabo", "20 min", "Demonstração: testador, luz da porta, velocidade do link"],
        ["3. PRÁTICA — Teste o seu cabo", "25 min", "Cada aluno testa o cabo que crimpou, dos 3 jeitos"],
        ["4. O que cada resultado quer dizer", "15 min", "Tabela de defeitos: pino aberto, cor trocada, par partido…"],
        ["INTERVALO", "15 min", ""],
        ["5. PRÁTICA — Clínica de cabos", "45 min", "Equipes diagnosticam os cabos C1 a C7 e escrevem o laudo"],
        ["6. PRÁTICA — Conserto e reteste", "25 min", "Cada equipe conserta os cabos que condenou"],
        ["7. Método de 4 passos + chamados", "25 min", "Olhar, testar, isolar, resolver — aplicado em role-play"],
        ["8. Duelo relâmpago", "15 min", "Time A × Time B"],
        ["9. Fechamento + tarefa", "15 min", "Quiz de saída e anúncio da Noite 6"],
    ], col_widths=[6, 2, 8.5])
    add_para(doc, "Blocos 3, 5, 6, 7 e 8: 135 min de atividade do aluno (64% do total). A exposição está quebrada em dois pedaços curtos, sempre com prática entre eles.", italic=True, color=COR_CINZA)

    add_hr(doc)

    # BLOCO 1
    add_heading(doc, "BLOCO 1 — RETOMADA + CABOS DA NOITE 4 · 10 MIN")
    add_heading(doc, "Recap relâmpago — 5 min", level=2)
    add_bullet(doc, "Diga o T568B do pino 1 ao 8 (em coro)")
    add_bullet(doc, "Qual o erro número um da crimpagem? (destrançar demais)")
    add_bullet(doc, "O que segura o cabo depois de crimpado? (a capa dentro do conector)")
    add_bullet(doc, "Cabo que não passou no teste serve pra quê? (pra nada — não se instala)")

    add_heading(doc, "Cabos na mesa — 5 min", level=2)
    add_para(doc, "Peça pra quem trouxe o cabo da Noite 4 colocar na mesa, com o nome numa fita crepe. Quem esqueceu forma dupla com quem trouxe.")
    add_para(doc, "\"Semana passada vocês fizeram o cabo e alguns testaram. Hoje a pergunta é outra: se um cliente te liga dizendo que a internet do computador tá ruim, como você PROVA que o cabo é o culpado — ou que é inocente?\"", italic=True, color=COR_PRIMARIA)

    add_hr(doc)

    # BLOCO 2
    add_heading(doc, "BLOCO 2 — TRÊS JEITOS DE TESTAR UM CABO · 20 MIN")
    add_para(doc, "Demonstração. Turma em volta da bancada. Use o cabo C7 (o bom) pra mostrar o normal, e só depois mostre um defeito.")

    add_heading(doc, "Jeito 1 — O testador de LED — 7 min", level=2)
    add_bullet(doc, "Duas peças: a principal (com bateria) e a remota. Uma ponta do cabo em cada")
    add_bullet(doc, "Liga e observa: os LEDs acendem um por um, de 1 a 8, nas duas peças, na mesma ordem")
    add_bullet(doc, "Mostre o C1: \"o LED 3 pulou. O fio 3 não está chegando do outro lado.\"")
    add_para(doc, "\"O testador conta uma história: qual fio sai de onde e chega aonde. Se a ordem das duas peças é igual e nenhum LED falta, a ligação está certa.\"", italic=True)
    add_para(doc, "Sem testador: pule pro Jeito 2 e diga que existe, como é e quanto custa. Mostre uma foto no celular.", italic=True, color=COR_CINZA)

    add_heading(doc, "Jeito 2 — A luz da porta — 5 min", level=2)
    add_bullet(doc, "Liga o cabo entre o roteador e o computador")
    add_bullet(doc, "A luz da porta LAN do roteador acende (e pisca quando passa dado). Muitos computadores também têm luz na própria porta")
    add_bullet(doc, "Luz apagada = não tem conexão elétrica. Mas luz acesa NÃO garante cabo perfeito")
    add_para(doc, "\"A luz da porta diz se o cabo conectou. Não diz se ele está bom.\"", italic=True, color=COR_ALERTA)

    add_heading(doc, "Jeito 3 — A velocidade do link — 8 min", level=2)
    add_bullet(doc, "No computador, abre a janela de Status da Ethernet e mostra: \"Velocidade: 1,0 Gbps\"")
    add_bullet(doc, "Troca pelo cabo C2. A luz acende, a internet funciona… e a velocidade cai pra 100 Mbps")
    add_para(doc, "\"Esse é o defeito mais traiçoeiro que existe. O cliente não liga dizendo 'meu cabo está ruim'. Ele liga dizendo 'contratei 500 Mega e o teste dá 94'. E o técnico que não sabe disso troca o roteador, reclama do provedor… e o problema era um fio.\"", italic=True, color=COR_PRIMARIA)
    add_para(doc, "Por quê: a conexão de 100 Mbps usa só 2 pares (laranja e verde). A de 1 Gbps precisa dos 4. Faltou um fio do azul ou do marrom, a placa de rede baixa pra 100 e segue funcionando.")
    add_para(doc, "Liga com a Noite 2: \"Lembram do TL-WR840N? As portas dele são de 100. Nele, esse teste nunca vai mostrar 1 Gbps — nem com cabo perfeito. Por isso hoje a gente usa o Archer.\"", italic=True)

    add_hr(doc)

    # BLOCO 3
    add_heading(doc, "BLOCO 3 — PRÁTICA · TESTE O SEU CABO · 25 MIN")
    add_para(doc, "Cada aluno (ou dupla) testa o cabo que crimpou na Noite 4 e preenche a página 1 da folha.")
    add_bullet(doc, "Monte 2 ou 3 estações: testador numa, roteador + computador na outra. A turma roda")
    add_bullet(doc, "Cada um anota os 3 resultados: testador, luz da porta, velocidade")
    add_bullet(doc, "Quem não trouxe cabo testa um dos seus de reserva")
    add_bullet(doc, "Anote no quadro: quantos deram 1 Gbps, quantos 100 Mbps, quantos sem conexão")
    add_para(doc, "Onde dilatar: se aparecer cabo de aluno com 100 Mbps, é ouro. Pare a turma e investigue junto, ao vivo, antes de mostrar a clínica.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)

    # BLOCO 4
    add_heading(doc, "BLOCO 4 — O QUE CADA RESULTADO QUER DIZER · 15 MIN")
    add_para(doc, "Desenhe no quadro e peça pra turma completar a coluna do meio antes de você falar:")
    add_table(doc, [
        ["O teste mostra…", "Nome do defeito", "Causa mais comum"],
        ["Um LED não acende", "Pino aberto", "Fio curto no conector, crimp fraco, fio partido"],
        ["LEDs fora de ordem", "Cor trocada", "Erro na sequência na hora de crimpar"],
        ["Dois LEDs acendem juntos", "Curto", "Fios encostando (conector amassado, cabo esmagado)"],
        ["1↔3 e 2↔6 trocados", "Cabo cruzado", "Uma ponta em B, outra em A (não é defeito)"],
        ["Conecta, mas em 100 Mbps", "Par partido", "Faltou fio do par azul ou marrom"],
        ["Funciona, mas cai quando mexe", "Mau contato", "Capa fora do conector, crimp fraco, clipe quebrado"],
    ], col_widths=[5, 3.5, 8])

    add_heading(doc, "Três avisos de campo", level=2)
    add_bullet(doc, "Testador simples não pega tudo. Um cabo pode passar no testador e falhar no uso. Por isso o técnico testa também no uso (velocidade, ping).")
    add_bullet(doc, "Nem todo defeito está na ponta. Cabo esmagado por porta, cadeira, grampo ou móvel dá defeito no meio.")
    add_bullet(doc, "Cabo de alumínio: tem cabo barato vendido como \"cabo de rede\" que é alumínio pintado de cobre (CCA). Raspe o fio com o estilete: se por dentro é prateado, não é cobre. Quebra fácil e perde sinal em distância.")

    add_hr(doc)
    add_heading(doc, "INTERVALO · 15 MIN", color=COR_CINZA)
    add_hr(doc)

    # BLOCO 5
    add_heading(doc, "BLOCO 5 — PRÁTICA · CLÍNICA DE CABOS · 45 MIN")
    add_para(doc, "Turma em equipes de 3. Cada equipe recebe um cabo da clínica por vez e tem ~6 minutos pra dar o laudo. Depois passa o cabo pra próxima equipe (rodízio). A página 2 da folha é o laudo.")

    add_heading(doc, "Regras da clínica", level=2)
    add_bullet(doc, "Primeiro OLHA (contra a luz, as duas pontas, o cabo inteiro passando pela mão). Só depois TESTA")
    add_bullet(doc, "Testa no testador (se tiver) e no roteador + computador, olhando a velocidade")
    add_bullet(doc, "Escreve: o que o teste mostrou, o que viu, o diagnóstico e como consertaria")
    add_bullet(doc, "Não pode cortar nem consertar ainda — só diagnosticar")

    add_heading(doc, "Gabarito (só com você)", level=2)
    add_table(doc, [["Cabo", "Defeito", "Conserto"]] +
              [[c, d, cons] for c, d, _, _, cons in CLINICA],
              col_widths=[1.3, 5, 10.2])

    add_heading(doc, "Fechamento da clínica — 10 min", level=2)
    add_para(doc, "Cabo por cabo, pergunte a cada equipe o laudo. Só depois revele. Faça questão de parar em três:")
    add_bullet(doc, "C2: \"quem achou que estava bom porque a luz acendeu?\" — é a lição principal da noite")
    add_bullet(doc, "C6: \"quem condenou o cruzado?\" — diferente não é defeito")
    add_bullet(doc, "C7: \"quem condenou o bom?\" — técnico que troca peça boa perde dinheiro e credibilidade")

    add_hr(doc)

    # BLOCO 6
    add_heading(doc, "BLOCO 6 — PRÁTICA · CONSERTO E RETESTE · 25 MIN")
    add_para(doc, "Cada equipe fica com um ou dois cabos condenados (C1 a C5) e conserta. Regra: só vale se passar no reteste.")
    add_bullet(doc, "Ponta ruim: corta o conector (perde ~2 cm), decapa de novo, crimpa")
    add_bullet(doc, "Defeito no meio (C5): corta o pedaço machucado e faz dois cabos menores — ou declara o cabo perdido")
    add_bullet(doc, "Reteste obrigatório: testador + velocidade. Só está consertado se der 1 Gbps")
    add_bullet(doc, "Quem consertou primeiro vira monitor de outra equipe")
    add_para(doc, "Teste do balanço (se sobrar tempo neste bloco): no Prompt, ping -t e o endereço do roteador. Com o C4 ligado, balance o conector. Cada \"Esgotado o tempo limite\" é um pacote perdido. Ctrl+C para. Troca pelo C7 e balança de novo: nada se perde.", italic=True, color=COR_PRIMARIA)

    add_hr(doc)

    # BLOCO 7
    add_heading(doc, "BLOCO 7 — MÉTODO DE 4 PASSOS + CHAMADOS · 25 MIN")

    add_heading(doc, "O método — 5 min", level=2)
    add_table(doc, [
        ["Passo", "O que fazer"],
        ["1. OLHAR", "Conector quebrado? Capa fora? Cabo esmagado, dobrado, grampeado?"],
        ["2. TESTAR", "Testador, luz da porta, velocidade. Anota o que viu"],
        ["3. ISOLAR", "Troca UMA coisa por vez: um cabo bom conhecido no lugar, outra porta, outro aparelho"],
        ["4. RESOLVER", "Refaz a ponta, troca o cabo ou passa pro próximo suspeito. E testa de novo"],
    ], col_widths=[3, 13.5])
    add_para(doc, "\"Todo técnico deveria andar com um cabo que ele SABE que está bom. É o cabo-testemunha. Trocou pelo cabo-testemunha e o problema sumiu? O culpado era o cabo. Não sumiu? O cabo é inocente — procura em outro lugar.\"", italic=True, color=COR_PRIMARIA)

    add_heading(doc, "Chamados em role-play — 20 min", level=2)
    add_para(doc, "Duplas: um é cliente, o outro técnico. O técnico tem que usar o método e dizer o que faria em cada passo. Troca os papéis a cada chamado. Os chamados estão na página 3 da folha:")
    chamados = [
        ("\"Contratei 300 Mega. No celular dá 280, no computador do cabo dá 94.\"",
         "Velocidade do link em 100 Mbps. Par partido no cabo, ou porta/placa de 100 Mbps. Testa com cabo-testemunha."),
        ("\"A internet do computador cai toda vez que eu arrasto a cadeira.\"",
         "Cabo passando embaixo da cadeira, esmagado ou com mau contato. Olhar o trajeto; teste do balanço."),
        ("\"Mudei o móvel de lugar e agora o computador diz 'cabo de rede desconectado'.\"",
         "Conector puxado ou clipe quebrado. Olhar a ponta; refazer ou trocar."),
        ("\"O técnico passou um cabo de 40 metros pelo forro e desde então a internet está ruim.\"",
         "Suspeitar do cabo novo: grampeado, esmagado, CCA, ou ponta mal crimpada. Testar e raspar o fio."),
    ]
    for ch, resp in chamados:
        add_para(doc, ch, bold=True)
        add_para(doc, "> " + resp, italic=True, color=COR_CINZA)

    add_hr(doc)

    # BLOCO 8
    add_heading(doc, "BLOCO 8 — DUELO RELÂMPAGO · 15 MIN")
    add_para(doc, "Turma em dois times. Resposta em 5 segundos. Alterna.")
    add_heading(doc, "12 perguntas de reserva", level=2)
    add_bullet(doc, "No testador, o LED 5 não acendeu. Qual o defeito? — Pino aberto no fio 5")
    add_bullet(doc, "Luz da porta acesa garante cabo bom? — Não")
    add_bullet(doc, "Conexão de 100 Mbps usa quantos pares? — 2 (laranja e verde)")
    add_bullet(doc, "E a de 1 Gbps? — Os 4")
    add_bullet(doc, "Cabo conecta em 100 em vez de 1000. Suspeito número um? — Par azul ou marrom com problema")
    add_bullet(doc, "Cabo cruzado é defeito? — Não")
    add_bullet(doc, "O que é o cabo-testemunha? — Um cabo que você sabe que está bom, pra comparar")
    add_bullet(doc, "Como descobrir se o cabo é de alumínio? — Raspar o fio: prateado por dentro")
    add_bullet(doc, "Internet cai quando mexe no cabo. Nome do defeito? — Mau contato")
    add_bullet(doc, "Qual o primeiro passo do método? — Olhar")
    add_bullet(doc, "Isolar é trocar quantas coisas por vez? — Uma")
    add_bullet(doc, "Distância máxima de um cabo de rede? — 100 metros")

    add_hr(doc)

    # BLOCO 9
    add_heading(doc, "BLOCO 9 — FECHAMENTO E TAREFA · 15 MIN")
    add_heading(doc, "Quiz de saída", level=2)
    add_bullet(doc, "Quais os três jeitos de testar um cabo?")
    add_bullet(doc, "Por que \"a luz acendeu\" não encerra o atendimento?")
    add_bullet(doc, "O cliente reclama de velocidade só no computador do cabo. O que você olha primeiro?")
    add_bullet(doc, "Quais são os 4 passos do método?")

    add_heading(doc, "A tarefa", level=2)
    add_para(doc, "\"Hoje o ping apareceu com um número estranho: 192.168.0.1. Semana que vem a gente entende que número é esse. A tarefa: no celular, entrem em Configurações → Sobre o telefone → Status (ou Informações do Wi-Fi) e anotem dois números: o endereço IP e o endereço MAC do Wi-Fi. Não precisa entender, só anotar.\"", italic=True)

    add_heading(doc, "Anuncie a próxima", level=2)
    add_para(doc, "\"Na próxima noite: endereço na rede. Todo aparelho tem um endereço, igual casa tem. Vocês vão descobrir o endereço do próprio celular e como o roteador acha cada aparelho no meio de vinte.\"", italic=True)

    add_hr(doc)

    # SE PERGUNTAREM
    add_heading(doc, "SE PERGUNTAREM")
    add_para(doc, "\"Qual testador comprar?\"", bold=True)
    add_para(doc, "> \"Pra começar, o de LED simples resolve 90% dos casos e é barato. Testador que mede distância até o defeito e certifica cabo existe, mas é caro — é pra empresa que faz rede grande.\"", italic=True)
    add_para(doc, "\"Posso emendar o cabo que estragou no meio?\"", bold=True)
    add_para(doc, "> \"Pode, com emenda RJ45 fêmea-fêmea: crimpa as duas pontas e liga na emenda. Funciona, mas é mais um ponto pra dar problema. Se der pra trocar o cabo, troca.\"", italic=True)
    add_para(doc, "\"Por que o cabo deu 1 Gbps mas a internet continua lenta?\"", bold=True)
    add_para(doc, "> \"Porque aí o cabo é inocente. Pode ser Wi-Fi, roteador, plano, o próprio site. O cabo foi descartado — isso também é diagnóstico.\"", italic=True)
    add_para(doc, "\"Cabo velho estraga sozinho?\"", bold=True)
    add_para(doc, "> \"O cobre dentro dura muito. O que estraga é o conector (clipe quebra, pino oxida em lugar úmido) e o cabo que sofre: esmagado, dobrado, roído, no sol.\"", italic=True)
    add_para(doc, "\"E se passar de 100 metros?\"", bold=True)
    add_para(doc, "> \"O sinal enfraquece. Pode até conectar, mas com erro e velocidade baixa. Acima de 100 m, coloca um switch no meio ou usa fibra.\"", italic=True)

    add_hr(doc)

    # DEPOIS DA AULA
    add_heading(doc, "DEPOIS DA AULA")
    add_bullet(doc, "Guarde os cabos da clínica com a etiqueta. Eles servem de novo na Noite 14 (diagnóstico completo)")
    add_bullet(doc, "Anote quantas equipes acertaram o C2 sem ajuda. É o indicador da noite")
    add_bullet(doc, "Os cabos consertados pelos alunos ficam na caixa de \"cabos testados\" da bancada")
    add_bullet(doc, "Reponha conectores se baixou de 30")

    add_hr(doc)

    # DESCRITIVO
    add_heading(doc, "DESCRITIVO PARA O DIÁRIO")
    add_para(doc, "Desenvolvi atividade prática de teste e diagnóstico de cabos de par trançado. Apresentei três formas de verificação: testador de continuidade com indicação por LED, observação do indicador luminoso de enlace nas portas dos equipamentos e verificação da velocidade de enlace negociada pelo adaptador de rede (100 Mbps e 1 Gbps), relacionando a queda de velocidade à falha em pares não utilizados pelo padrão Fast Ethernet. Os estudantes testaram os cabos confeccionados na aula anterior e, em equipes, realizaram diagnóstico de um conjunto de cabos com defeitos preparados (pino aberto, par interrompido, inversão de condutores, mau contato, dano no meio do cabo, cabo cruzado e cabo em perfeito estado como controle), registrando laudo técnico e executando a correção com reteste. Apresentei método de diagnóstico em quatro etapas (inspeção visual, teste, isolamento por substituição e resolução), aplicado em simulações de atendimento ao cliente.")

    add_hr(doc)

    # CAIXA DE FERRAMENTAS DIDÁTICAS
    add_heading(doc, "CAIXA DE FERRAMENTAS DIDÁTICAS", color=COR_DESTAQUE)
    add_para(doc, "Se sobrar tempo, ou pra usar em outra noite:", italic=True, color=COR_CINZA)

    add_para(doc, "1. Sabotagem entre equipes", bold=True)
    add_para(doc, "Cada equipe faz um cabo com UM defeito escondido e entrega pra outra diagnosticar. Quem esconder melhor e quem descobrir mais rápido ganham. 25-30 min.")
    add_para(doc, "2. Teste do balanço cronometrado", bold=True)
    add_para(doc, "ping -t com o C4. Cada aluno balança o conector por 30 segundos e conta os pacotes perdidos. Depois refaz a ponta e repete. Mostra em número o que é mau contato. 15 min.")
    add_para(doc, "3. Raspe o fio", bold=True)
    add_para(doc, "Traga um pedaço de cabo CCA (alumínio cobreado) e um de cobre. A turma raspa os dois e compara a cor por dentro, e dobra os dois várias vezes pra ver qual quebra primeiro. 10 min.")
    add_para(doc, "4. Caça ao defeito no meio", bold=True)
    add_para(doc, "Passe um cabo longo (10 m+) com um ponto esmagado escondido embaixo de fita. Aluno tem que achar passando o cabo pela mão, centímetro por centímetro. 10 min.")
    add_para(doc, "5. O TL-WR840N contra o Archer", bold=True)
    add_para(doc, "Liga o mesmo cabo bom nos dois roteadores e compara a velocidade: 100 × 1000. Retoma a Noite 2 e mostra que o limite pode estar no aparelho, não no cabo. 10 min.")

    path = os.path.join(OUTDIR, "Roteiro - Redes 5 - Testando cabos.docx")
    doc.save(path)
    return path


# ============================================================
# SLIDES
# ============================================================
def build_slides():
    prs = new_presentation()
    P, D, K, R = PPalette.P, PPalette.D, PPalette.K, PPalette.R

    def lista(kicker, titulo, linhas, subtitulo=None, size=20, nota=None, nota_cor=None):
        s = add_slide(prs)
        add_title_block(s, kicker, titulo, subtitulo)
        add_text(s, linhas, 0.7, 2.8, 11.9, 3.8, size=size, color=K)
        if nota:
            add_text(s, nota, 0.7, 6.4, 11.9, 0.5, size=16, italic=True, color=nota_cor or D, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    def frase(kicker, texto, cor=None):
        s = add_slide(prs)
        add_text(s, kicker, 0.7, 1.6, 11.9, 0.5, size=16, bold=True, color=D, align=PP_ALIGN.CENTER)
        add_text(s, texto.split("\n"), 0.9, 2.4, 11.5, 3.0, size=36, bold=True, color=cor or P, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    def tabela(kicker, titulo, rows, col_widths=None, subtitulo=None, size=16, nota=None):
        s = add_slide(prs)
        add_title_block(s, kicker, titulo, subtitulo)
        add_table_slide(s, rows, top=2.6, height=3.6, col_widths=col_widths, size=size)
        if nota:
            add_text(s, nota, 0.7, 6.5, 11.9, 0.5, size=15, italic=True, color=D, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    # 1: Capa
    add_slide_cover(prs, "NOITE 5", "CABO NÃO", "SE CHUTA",
                    "Testando cabos e resolvendo cabos ruins",
                    "Redes de Computadores · Curso FIC · CEJA Itapiranga – SC")

    # 2: Recap
    lista("RETOMADA", "Da Noite 4 pra hoje", [
        "•  Diga o T568B, do pino 1 ao 8",
        "•  Qual o erro número um da crimpagem?",
        "•  O que segura o cabo depois de crimpado?",
        "",
        "Coloque seu cabo da Noite 4 na mesa, com seu nome na fita.",
    ])

    # 3: Pergunta da noite
    frase("A PERGUNTA DA NOITE", "A internet do cliente está ruim.\nComo você PROVA que o cabo\né o culpado — ou que é inocente?")

    # 4: Três jeitos
    lista("TRÊS JEITOS DE TESTAR", "Do mais simples ao mais esperto", [
        "1.  Testador de LED — conta fio por fio",
        "2.  Luz da porta — diz se conectou",
        "3.  Velocidade do link — diz se conectou BEM",
    ], size=26)

    # 5: Testador
    lista("JEITO 1", "O testador de LED", [
        "•  Duas peças: principal e remota, uma em cada ponta",
        "•  Os LEDs acendem de 1 a 8, nas duas, na mesma ordem",
        "•  LED faltando: o fio não chega do outro lado",
        "•  LED fora de ordem: cor trocada",
        "•  Dois LEDs juntos: fios encostando (curto)",
    ], nota="O testador conta qual fio sai de onde e chega aonde.")

    # 6: Luz da porta
    lista("JEITO 2", "A luz da porta", [
        "•  Liga o cabo entre o roteador e o computador",
        "•  Acendeu: tem conexão elétrica",
        "•  Piscando: está passando dado",
        "•  Apagada: não conectou",
    ], nota="A luz diz se o cabo conectou. Não diz se ele está bom.", nota_cor=R)

    # 7: Velocidade
    lista("JEITO 3", "A velocidade do link", [
        "Windows: Central de Rede → Ethernet → Status → Velocidade",
        "",
        "•  1,0 Gbps — os 4 pares funcionando",
        "•  100 Mbps — alguma coisa está limitando",
        "•  Sem conexão — nem conectou",
    ])

    # 8: 2 pares x 4 pares
    tabela("POR QUE CAI PRA 100?", "100 Mbps usa 2 pares. 1 Gbps usa os 4.", [
        ["Par", "Pinos", "100 Mbps", "1 Gbps"],
        ["Laranja", "1 e 2", "usa", "usa"],
        ["Verde", "3 e 6", "usa", "usa"],
        ["Azul", "4 e 5", "—", "usa"],
        ["Marrom", "7 e 8", "—", "usa"],
    ], col_widths=[3, 3, 2.95, 2.95], size=18,
        nota="Faltou um fio do azul ou do marrom? Funciona… em 100. O cliente nem percebe que é o cabo.")

    # 9: Frase
    frase("O DEFEITO MAIS TRAIÇOEIRO", "\"Contratei 500 Mega\ne o teste dá 94.\"")

    # 10: Lembra da Noite 2
    lista("LEMBRA DA NOITE 2?", "O limite pode estar no aparelho", [
        "•  TL-WR840N: portas de 100 Mbps",
        "•  Archer C5: portas Gigabit (1000 Mbps)",
        "",
        "No 840N, o teste nunca mostra 1 Gbps — nem com cabo perfeito.",
        "Antes de condenar o cabo, confira a porta.",
    ])

    # 11: Prática
    lista("PRÁTICA · 25 MIN", "Teste o SEU cabo", [
        "1.  Testador (se tiver) — anote os LEDs",
        "2.  Roteador + computador — a luz da porta acendeu?",
        "3.  Status da Ethernet — qual a velocidade?",
        "",
        "Anote na página 1 da folha.",
    ])

    # 12: Tabela de defeitos
    tabela("O QUE O TESTE ESTÁ DIZENDO", "Resultado → defeito", [
        ["O teste mostra…", "Defeito", "Causa comum"],
        ["Um LED não acende", "Pino aberto", "Fio curto, crimp fraco"],
        ["LEDs fora de ordem", "Cor trocada", "Sequência errada"],
        ["Dois LEDs juntos", "Curto", "Fios encostando"],
        ["1↔3 e 2↔6", "Cabo cruzado", "B numa ponta, A na outra"],
        ["Conecta em 100 Mbps", "Par partido", "Azul ou marrom falhando"],
        ["Cai quando mexe", "Mau contato", "Capa fora, crimp fraco"],
    ], col_widths=[4, 3, 4.9], size=15)

    # 13: Avisos
    lista("TRÊS AVISOS DE CAMPO", "O que o livro não conta", [
        "•  Testador simples não pega tudo. Teste também no uso.",
        "•  Nem todo defeito está na ponta: porta, cadeira e grampo esmagam cabo.",
        "•  Cabo de alumínio (CCA): raspe o fio. Prateado por dentro = não é cobre.",
    ], size=19)

    # 14: Intervalo
    frase("", "INTERVALO\n15 minutos", cor=K)

    # 15: Clínica
    lista("PRÁTICA · 45 MIN", "Clínica de cabos", [
        "•  Equipes de 3. Sete cabos: C1 a C7",
        "•  ~6 minutos por cabo, depois passa pra próxima equipe",
        "•  Primeiro OLHA. Depois TESTA",
        "•  Escreve o laudo na página 2 da folha",
        "•  Ainda não pode consertar — só diagnosticar",
    ], nota="Atenção: nem todo cabo da mesa está com defeito.")

    # 16: Conserto
    lista("PRÁTICA · 25 MIN", "Conserto e reteste", [
        "•  Ponta ruim: corta o conector, decapa de novo, crimpa",
        "•  Defeito no meio: corta o pedaço ruim ou declara o cabo perdido",
        "•  Só está consertado se passar no reteste: testador + 1 Gbps",
    ])

    # 17: Teste do balanço
    lista("TESTE DO BALANÇO", "Pegando o mau contato no flagra", [
        "Prompt de Comando:   ping -t 192.168.0.1",
        "",
        "•  Balance o conector com o ping rodando",
        "•  \"Esgotado o tempo limite\" = pacote perdido",
        "•  Ctrl+C para",
    ], nota="O 192.168.0.1 é o endereço do roteador. Na Noite 6 a gente entende.")

    # 18: Método
    tabela("O MÉTODO", "Olhar · Testar · Isolar · Resolver", [
        ["Passo", "O que fazer"],
        ["1. OLHAR", "Conector, capa, trajeto do cabo"],
        ["2. TESTAR", "Testador, luz, velocidade — anota"],
        ["3. ISOLAR", "Troca UMA coisa por vez"],
        ["4. RESOLVER", "Refaz, troca ou passa pro próximo suspeito"],
    ], col_widths=[3, 8.9], size=18)

    # 19: Cabo-testemunha
    frase("O CABO-TESTEMUNHA", "Ande sempre com um cabo\nque você SABE que está bom.")

    # 20: Chamados
    lista("ROLE-PLAY · 20 MIN", "Chamados", [
        "•  \"No celular dá 280 Mega, no computador do cabo dá 94.\"",
        "•  \"A internet cai toda vez que eu arrasto a cadeira.\"",
        "•  \"Mudei o móvel e agora diz 'cabo de rede desconectado'.\"",
        "•  \"Passaram um cabo de 40 m pelo forro e desde então tá ruim.\"",
    ], size=18, nota="Técnico: diga o que faria em cada um dos 4 passos.")

    # 21: Duelo
    lista("DUELO RELÂMPAGO", "Time A × Time B", [
        "•  5 segundos pra responder",
        "•  Errou, passa pro outro time",
        "•  Vale ponto quem explicar o porquê",
    ], size=24)

    # 22: O que fica
    lista("O QUE FICA", "Três frases pra guardar", [
        "\"Cabo não se chuta. Cabo se testa.\"",
        "",
        "\"Luz acesa não é cabo bom. Olhe a velocidade.\"",
        "",
        "\"Troque uma coisa por vez.\"",
    ], size=24)

    # 23: Próxima
    lista("PARA A PRÓXIMA NOITE", "Endereço na rede: IP e MAC", [
        "Uma tarefa simples:",
        "",
        "•  No celular: Configurações → Sobre o telefone → Status",
        "•  Anote o endereço IP e o endereço MAC do Wi-Fi",
        "•  Não precisa entender — só anotar",
    ])

    path = os.path.join(OUTDIR, "Aula - Redes 5 - Testando cabos.pptx")
    prs.save(path)
    return path


# ============================================================
# FOLHA DO ALUNO
# ============================================================
def linhas(doc, n):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("_" * 100)
        r.font.size = Pt(11)
        r.font.color.rgb = COR_CINZA


def pergunta(doc, texto, size=11):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.font.size = Pt(size)
    r.bold = True


def build_folha():
    doc = Document()
    setup_page(doc, margem=Cm(1.8))

    # PÁGINA 1 — Três jeitos de testar + teste do próprio cabo
    cabecalho_documento(
        doc,
        "CABO NÃO SE CHUTA, SE TESTA",
        f"Redes de Computadores · {NOITE} · Curso FIC · CEJA Itapiranga · 1 de 3",
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Nome: ")
    r.font.size = Pt(11)
    r.bold = True
    r = p.add_run("_________________________________________________     Data: ____ / ____ / ______")
    r.font.size = Pt(11)

    add_heading(doc, "1  Três jeitos de testar", level=2)
    add_table(doc, [
        ["Jeito", "Como faz", "O que ele te diz"],
        ["Testador de LED", "Uma ponta em cada peça, liga e olha os LEDs de 1 a 8", "Qual fio chega e onde chega"],
        ["Luz da porta", "Liga entre roteador e computador", "Se conectou (não se está bom)"],
        ["Velocidade do link", "Status da Ethernet no computador", "Se conectou BEM: 1 Gbps ou 100 Mbps"],
    ], col_widths=[3.8, 7.2, 6.5])

    add_heading(doc, "2  Por que o cabo cai pra 100 Mbps", level=2)
    add_para(doc, "Complete com \"usa\" ou \"—\".", italic=True, size=10, color=COR_CINZA)
    tbl = add_table(doc, [
        ["Par", "Pinos", "Conexão de 100 Mbps", "Conexão de 1 Gbps"],
        ["Laranja", "1 e 2", "", ""],
        ["Verde", "3 e 6", "", ""],
        ["Azul", "4 e 5", "", ""],
        ["Marrom", "7 e 8", "", ""],
    ], col_widths=[3, 2.5, 6, 6])
    for row in tbl.rows[1:]:
        row.height = Cm(0.7)

    add_heading(doc, "3  Teste o seu cabo da Noite 4", level=2)
    tbl = add_table(doc, [
        ["Teste", "Resultado", "Está bom?"],
        ["Testador: quais LEDs acenderam, e em que ordem?", "", "☐ Sim   ☐ Não"],
        ["Luz da porta do roteador", "☐ Acesa   ☐ Piscando   ☐ Apagada", "☐ Sim   ☐ Não"],
        ["Velocidade no computador", "☐ 1 Gbps   ☐ 100 Mbps   ☐ Sem conexão", "☐ Sim   ☐ Não"],
    ], col_widths=[6, 7, 4.5])
    for row in tbl.rows[1:]:
        row.height = Cm(1.0)

    pergunta(doc, "Seu cabo ficou reprovado em algum teste? Qual defeito você acha que ele tem?")
    linhas(doc, 3)

    # PÁGINA 2 — Clínica de cabos
    doc.add_page_break()
    cabecalho_documento(
        doc,
        "CLÍNICA DE CABOS · LAUDO DA EQUIPE",
        f"Diagnóstico em equipe · {NOITE} · 2 de 3",
    )
    add_para(doc, "Primeiro OLHE (as duas pontas contra a luz e o cabo inteiro pela mão). Depois TESTE. Ainda não conserte. Atenção: nem todo cabo da mesa está com defeito.", italic=True, size=10, color=COR_CINZA)

    p = doc.add_paragraph()
    r = p.add_run("Equipe: ______________________________________________________________")
    r.font.size = Pt(11)

    tbl = add_table(doc, [["Cabo", "O que o teste mostrou", "O que você viu no cabo", "Diagnóstico", "Como consertar"]] +
                    [[c, "", "", "", ""] for c, *_ in CLINICA],
                    col_widths=[1.3, 4.5, 4, 3.7, 4])
    for row in tbl.rows[1:]:
        row.height = Cm(1.75)

    add_heading(doc, "Depois do gabarito", level=2)
    p = doc.add_paragraph()
    r = p.add_run("Quantos a equipe acertou?  ______ de 7        Qual enganou vocês? ______        Por quê?")
    r.font.size = Pt(11)
    linhas(doc, 2)

    # PÁGINA 3 — Método + chamados + tarefa
    doc.add_page_break()
    cabecalho_documento(
        doc,
        "O MÉTODO E OS CHAMADOS",
        f"Atendimento · {NOITE} · 3 de 3",
    )

    add_heading(doc, "1  Os 4 passos", level=2)
    add_table(doc, [
        ["Passo", "O que fazer"],
        ["1. OLHAR", "Conector quebrado? Capa fora? Cabo esmagado, dobrado, grampeado?"],
        ["2. TESTAR", "Testador, luz da porta, velocidade. Anote o que viu"],
        ["3. ISOLAR", "Troque UMA coisa por vez: cabo-testemunha, outra porta, outro aparelho"],
        ["4. RESOLVER", "Refaça a ponta, troque o cabo ou passe pro próximo suspeito. Teste de novo"],
    ], col_widths=[3, 14.5])

    add_heading(doc, "2  Chamados", level=2)
    add_para(doc, "Em dupla: um é o cliente, o outro o técnico. Escreva o que o técnico faria.", italic=True, size=10, color=COR_CINZA)
    chamados = [
        "\"Contratei 300 Mega. No celular dá 280, no computador do cabo dá 94.\"",
        "\"A internet do computador cai toda vez que eu arrasto a cadeira.\"",
        "\"Mudei o móvel de lugar e agora o computador diz 'cabo de rede desconectado'.\"",
    ]
    for i, ch in enumerate(chamados, 1):
        add_para(doc, f"Chamado {i}:  {ch}", italic=True, color=COR_PRIMARIA, size=11)
        linhas(doc, 2)

    add_heading(doc, "TAREFA DA SEMANA", level=2, color=COR_DESTAQUE)
    add_bullet(doc, "No celular: Configurações → Sobre o telefone → Status (ou Informações do Wi-Fi)")
    add_bullet(doc, "Anote aqui o endereço IP: ___________________   e o endereço MAC: ___________________")
    add_bullet(doc, "Não precisa entender ainda — na Noite 6 a gente descobre o que esses números são")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Traga esta folha na próxima aula · CEJA Itapiranga – SC")
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = COR_CINZA

    path = os.path.join(OUTDIR, "Folha do Aluno - Redes 5 - Testando cabos.docx")
    doc.save(path)
    return path


# ============================================================
# RESUMO DO ALUNO
# ============================================================
def build_resumo():
    doc = Document()
    setup_page(doc, margem=Cm(1.8))

    cabecalho_documento(
        doc,
        "TESTANDO CABOS E RESOLVENDO CABOS RUINS",
        f"Resumo do aluno · {NOITE} · Redes de Computadores · Curso FIC · CEJA Itapiranga",
    )

    add_heading(doc, "A ideia principal", level=2)
    add_para(doc, "O técnico iniciante troca peça até funcionar. O técnico bom testa, descobre o defeito e troca só o que precisa. Cabo não se chuta: cabo se testa.")

    add_heading(doc, "Três jeitos de testar", level=2)
    add_bullet(doc, "Testador de LED: os LEDs acendem de 1 a 8, nas duas peças, na mesma ordem. Mostra qual fio chega e onde")
    add_bullet(doc, "Luz da porta: acesa = conectou. Mas luz acesa NÃO garante cabo bom")
    add_bullet(doc, "Velocidade do link (Windows: Central de Rede → Ethernet → Status): 1 Gbps = os 4 pares funcionando")

    add_heading(doc, "Por que o cabo cai pra 100 Mbps", level=2)
    add_para(doc, "A conexão de 100 Mbps usa só 2 pares (laranja e verde). A de 1 Gbps usa os 4. Se um fio do par azul ou marrom falha, o computador baixa pra 100 e continua funcionando — e o cliente reclama de internet lenta sem saber que é o cabo. Antes de condenar o cabo, confira se a porta do roteador é Gigabit.")

    add_heading(doc, "O que o teste está dizendo", level=2)
    add_table(doc, [
        ["O teste mostra…", "Defeito", "Causa comum"],
        ["Um LED não acende", "Pino aberto", "Fio curto no conector, crimp fraco, fio partido"],
        ["LEDs fora de ordem", "Cor trocada", "Sequência errada na crimpagem"],
        ["Dois LEDs acendem juntos", "Curto", "Fios encostando"],
        ["1↔3 e 2↔6 trocados", "Cabo cruzado", "Não é defeito: B numa ponta, A na outra"],
        ["Conecta em 100 Mbps", "Par partido", "Azul ou marrom falhando"],
        ["Cai quando mexe", "Mau contato", "Capa fora do conector, crimp fraco, clipe quebrado"],
    ], col_widths=[4.5, 3.2, 7.8])

    add_heading(doc, "O método de 4 passos", level=2)
    add_bullet(doc, "1. OLHAR — conector, capa, trajeto do cabo")
    add_bullet(doc, "2. TESTAR — testador, luz, velocidade. Anote")
    add_bullet(doc, "3. ISOLAR — troque UMA coisa por vez (cabo-testemunha, outra porta, outro aparelho)")
    add_bullet(doc, "4. RESOLVER — refaça a ponta ou troque o cabo. E teste de novo")

    add_heading(doc, "Avisos de campo", level=2)
    add_bullet(doc, "Testador simples não pega tudo. Teste também no uso (velocidade, ping)")
    add_bullet(doc, "Defeito pode estar no meio do cabo: porta, cadeira, grampo e móvel esmagam cabo")
    add_bullet(doc, "Cabo de alumínio (CCA): raspe o fio. Prateado por dentro = não é cobre. Evite")
    add_bullet(doc, "Teste do balanço: ping -t no endereço do roteador e balance o conector. Pacote perdido = mau contato")

    add_heading(doc, "Três frases pra guardar", level=2)
    add_para(doc, "\"Cabo não se chuta. Cabo se testa.\"", bold=True, color=COR_PRIMARIA)
    add_para(doc, "\"Luz acesa não é cabo bom. Olhe a velocidade.\"", bold=True, color=COR_PRIMARIA)
    add_para(doc, "\"Troque uma coisa por vez.\"", bold=True, color=COR_PRIMARIA)

    add_heading(doc, "Pra se testar em casa", level=2)
    add_bullet(doc, "O cliente tem 500 Mega, mas no computador do cabo o teste dá 94. Qual seu primeiro suspeito?")
    add_bullet(doc, "Por que o cabo cruzado não é defeito?")
    add_bullet(doc, "O que é o cabo-testemunha e pra que serve?")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Este material é seu. Guarde e volte quando precisar. · CEJA Itapiranga – SC")
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = COR_CINZA

    path = os.path.join(OUTDIR, "Resumo do Aluno - Redes 5 - Testando cabos.docx")
    doc.save(path)
    return path


# ============================================================
# MAIN
# ============================================================
if __name__ == "__main__":
    print("Gerando Roteiro...")
    print("  →", build_roteiro())
    print("Gerando Slides...")
    print("  →", build_slides())
    print("Gerando Folha do Aluno...")
    print("  →", build_folha())
    print("Gerando Resumo do Aluno...")
    print("  →", build_resumo())
    print("\nPronto.")
