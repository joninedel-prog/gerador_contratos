"""Gera os 4 arquivos da Noite 8 do curso de Redes.

Noite 8 — Rede local em casa I: planejando a instalação
Padrão prática-pesada, sem role-play e sem duelo. Parte da tarefa da Noite 7
(planta da casa desenhada à mão). A turma faz a galeria das plantas, preenche a
ficha de visita técnica da própria casa, passa por uma clínica de 8 plantas com
roteador mal posicionado, mede na escola o caminho real de um cabo (trena),
projeta a rede de uma das casas na malha quadriculada, orça com a tabela de
preços de referência e confere o projeto de outra dupla por checklist.
Não precisa de internet.
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL

from common import (
    COR_PRIMARIA, COR_ALERTA, COR_DESTAQUE, COR_CINZA,
    add_heading, add_para, add_bullet, add_table, add_hr,
    setup_page, cabecalho_documento,
    PPalette, new_presentation, add_slide, add_text, add_title_block,
    add_slide_cover, add_footer, add_table_slide,
)
from pptx.enum.text import PP_ALIGN

OUTDIR = "/home/nedel/curso"
NOITE = "Noite 8"
FOOTER = f"Redes de Computadores · {NOITE}"


# A visita técnica: o que perguntar e o que olhar
VISITA = [
    ("Quem usa", "Quantas pessoas moram? Alguém trabalha ou estuda de casa?"),
    ("O que liga", "Celulares, TV, computador, videogame, câmera, impressora. Onde cada um fica?"),
    ("Parado ou anda", "O aparelho fica sempre no mesmo lugar (TV, PC) ou anda pela casa (celular)?"),
    ("Onde a internet chega", "Por onde entra o cabo do provedor? Onde está o modem/ONU?"),
    ("A casa", "Térrea ou sobrado? Alvenaria ou madeira? Tem laje? Galpão, garagem, área externa?"),
    ("Onde dói", "Em quais cômodos o Wi-Fi é fraco ou cai? A que horas fica lento?"),
    ("Por onde passa", "Tem forro, rodapé, canaleta ou eletroduto vazio por onde o cabo pode ir?"),
    ("Tomadas", "Tem tomada perto de onde o roteador vai ficar?"),
]

# Regras de posição do roteador
REGRAS_ROTEADOR = [
    ("No centro", "O sinal se espalha pra todos os lados, igual luz. No canto, metade vai pra rua"),
    ("No alto", "Em cima de móvel ou na parede, a 1,5 m ou mais. No chão, o sinal bate nos móveis"),
    ("À vista", "Nunca dentro de armário, rack fechado ou gaveta"),
    ("Longe de metal e água", "Geladeira, micro-ondas, espelho grande, aquário, caixa d'água, quadro de luz"),
    ("Perto de tomada", "E com espaço pra ventilar: roteador esquenta"),
]

# Clínica de plantas: título, queixa do cliente, grade de cômodos, textos, problema, solução
# Na grade, células com o mesmo código viram um cômodo só.
PLANTAS = [
    {
        "titulo": "Onde o técnico deixou",
        "queixa": "\"Nos quartos dos fundos o Wi-Fi não pega. Na sala é ótimo.\"",
        "grade": [["Q1", "BW", "Q2"], ["CO", "CO", "CO"], ["SA", "CZ", "CZ"]],
        "texto": {"Q1": "QUARTO 1", "BW": "BANHEIRO", "Q2": "QUARTO 2", "CO": "CORREDOR",
                  "SA": "SALA\n► a fibra entra aqui\n■ ROTEADOR (ao lado da porta)", "CZ": "COZINHA"},
        "problema": "Roteador no canto da casa, onde o técnico do provedor deixou a fibra. Os quartos ficam no lado oposto",
        "solucao": "Levar um cabo de rede da ONU pelo rodapé até o corredor e colocar o roteador lá, no alto. Uns 8 m de cabo",
    },
    {
        "titulo": "Escondido no rack",
        "queixa": "\"Comprei roteador novo e continua fraco até na cozinha, que é do lado.\"",
        "grade": [["Q1", "Q2", "BW"], ["SA", "SA", "CZ"], ["SA", "SA", "CZ"]],
        "texto": {"Q1": "QUARTO 1", "Q2": "QUARTO 2", "BW": "BANHEIRO",
                  "SA": "SALA\n■ ROTEADOR dentro do rack fechado,\nno chão, atrás da TV", "CZ": "COZINHA"},
        "problema": "Roteador fechado no armário, no chão e atrás da TV (metal). O sinal sai fraco de lá",
        "solucao": "Tirar do rack e colocar em cima dele ou na parede, no alto, à vista. Custo zero",
    },
    {
        "titulo": "O sobrado",
        "queixa": "\"No quarto do meu filho, lá em cima, não funciona nada.\"",
        "grade": [["SUP", "SUP", "SUP"], ["SA", "ES", "CZ"], ["SA", "GA", "GA"]],
        "texto": {"SUP": "ANDAR DE CIMA (laje de concreto no meio)\nQUARTO DO FILHO · QUARTO · BANHEIRO",
                  "SA": "SALA (térreo)\n■ ROTEADOR em cima da estante", "ES": "ESCADA",
                  "CZ": "COZINHA", "GA": "GARAGEM"},
        "problema": "A laje de concreto entre os andares segura o sinal. O roteador está no térreo, longe da escada",
        "solucao": "Subir um cabo pela escada até o andar de cima e pôr ali um segundo aparelho (ponto de acesso ou kit mesh)",
    },
    {
        "titulo": "O galpão",
        "queixa": "\"Quero câmera e celular funcionando na oficina, lá nos fundos do terreno.\"",
        "grade": [["CA", "CA", "PA"], ["CA", "CA", "PA"], ["PA", "PA", "GL"]],
        "texto": {"CA": "CASA\n■ ROTEADOR na sala", "PA": "PÁTIO\n(30 m entre a casa e o galpão)",
                  "GL": "GALPÃO /\nOFICINA"},
        "problema": "Trinta metros de pátio e duas paredes: o Wi-Fi da casa não chega. Repetidor no meio do pátio é gambiarra",
        "solucao": "Cabo próprio pra área externa, dentro de eletroduto, até o galpão (o cabo vai até 100 m). Lá, um roteador como ponto de acesso",
    },
    {
        "titulo": "A cozinha",
        "queixa": "\"Toda vez que esquento comida, a internet cai.\"",
        "grade": [["Q1", "Q2", "Q3"], ["CO", "CO", "CO"], ["SA", "SA", "CZ"]],
        "texto": {"Q1": "QUARTO 1", "Q2": "QUARTO 2", "Q3": "QUARTO 3", "CO": "CORREDOR",
                  "SA": "SALA",
                  "CZ": "COZINHA\n■ ROTEADOR em cima da\ngeladeira, do lado\ndo micro-ondas"},
        "problema": "O micro-ondas atrapalha o Wi-Fi quando está ligado, e a geladeira (metal) bloqueia o sinal",
        "solucao": "Tirar da cozinha: corredor, no alto. Na Noite 10 a gente vê por que o micro-ondas atrapalha",
    },
    {
        "titulo": "O gamer do fundo",
        "queixa": "\"O sinal no meu quarto é bom, mas o jogo online trava.\"",
        "grade": [["Q1", "BW", "Q2"], ["CO", "CO", "CO"], ["SA", "SA", "CZ"]],
        "texto": {"Q1": "QUARTO 1", "BW": "BANHEIRO",
                  "Q2": "QUARTO DO GAMER\nvideogame e PC no Wi-Fi", "CO": "CORREDOR\n■ ROTEADOR no alto",
                  "SA": "SALA", "CZ": "COZINHA"},
        "problema": "A posição do roteador está boa. Mas videogame e PC ficam parados e precisam de conexão firme",
        "solucao": "Cabo do roteador até o quarto, pelo rodapé com canaleta. O que fica parado vai no cabo",
    },
    {
        "titulo": "A área de serviço",
        "queixa": "\"A sala é o pior lugar da casa pra internet.\"",
        "grade": [["SA", "SA", "Q1"], ["SA", "SA", "Q2"], ["CZ", "AS", "BW"]],
        "texto": {"SA": "SALA", "Q1": "QUARTO 1", "Q2": "QUARTO 2", "CZ": "COZINHA",
                  "AS": "ÁREA DE SERVIÇO\n■ ROTEADOR junto do\nquadro de luz, atrás\nda parede de azulejo",
                  "BW": "BANHEIRO"},
        "problema": "Roteador escondido na área de serviço, ao lado do quadro de luz, atrás de parede de azulejo e do tanque",
        "solucao": "Levar cabo da área de serviço até a parede entre a sala e os quartos e colocar o roteador lá, no alto",
    },
    {
        "titulo": "Nada a mexer",
        "queixa": "\"Todo dia, depois das oito da noite, a internet fica lenta.\"",
        "grade": [["Q1", "ES", "Q2"], ["CO", "CO", "CO"], ["SA", "SA", "CZ"]],
        "texto": {"Q1": "QUARTO 1", "ES": "ESCRITÓRIO\nPC no cabo", "Q2": "QUARTO 2",
                  "CO": "CORREDOR\n■ ROTEADOR no alto, na parede", "SA": "SALA\nTV no cabo",
                  "CZ": "COZINHA"},
        "problema": "A instalação está certa: roteador no centro e no alto, aparelhos parados no cabo. Lentidão em horário fixo é do plano ou do provedor",
        "solucao": "Não mexer na instalação. Teste de velocidade no PC do cabo, à noite, e falar com o provedor. Não vender o que o cliente não precisa",
    },
]

# Tabela de preços de referência (valores para a atividade, não de loja real)
PRECOS = [
    ("Cabo de rede Cat5e (interno)", "metro", 2.50),
    ("Cabo de rede Cat6 (interno)", "metro", 4.00),
    ("Cabo de rede para área externa", "metro", 5.00),
    ("Conector RJ45", "unidade", 1.00),
    ("Capa do conector", "unidade", 0.50),
    ("Canaleta (com adesivo)", "metro", 6.00),
    ("Eletroduto para área externa", "metro", 4.00),
    ("Tomada de rede (ponto na parede)", "unidade", 25.00),
    ("Roteador simples (100 Mbps)", "unidade", 120.00),
    ("Roteador dual band Gigabit", "unidade", 250.00),
    ("Repetidor de sinal", "unidade", 110.00),
    ("Kit mesh (2 aparelhos)", "kit", 600.00),
    ("Mão de obra: visita técnica", "visita", 50.00),
    ("Mão de obra: por ponto de cabo instalado", "ponto", 60.00),
    ("Mão de obra: configuração do roteador", "aparelho", 80.00),
]

# Checklist da banca
CHECKLIST = [
    "A entrada da internet (E) está marcada",
    "O roteador (R) está no centro da casa, no alto e fora de armário",
    "Os aparelhos que ficam parados (TV, PC, videogame) estão no cabo",
    "O caminho do cabo vai pelo rodapé, forro ou canaleta (nunca no meio do chão)",
    "A metragem tem subida, descida e sobra",
    "Conectores: 2 por cabo, mais reserva",
    "O orçamento tem mão de obra",
    "A dupla sabe explicar por que escolheu cada coisa",
]


def reais(v):
    return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def planta(doc, p, altura=1.2, largura=5.6, size=8):
    """Desenha a grade de cômodos; códigos iguais viram um cômodo só."""
    grade = p["grade"]
    n_l, n_c = len(grade), len(grade[0])
    tbl = doc.add_table(rows=n_l, cols=n_c)
    tbl.style = 'Table Grid'
    for row in tbl.rows:
        row.height = Cm(altura)
        for cell in row.cells:
            cell.width = Cm(largura)
    feitos = set()
    for i in range(n_l):
        for j in range(n_c):
            cod = grade[i][j]
            if cod in feitos:
                continue
            feitos.add(cod)
            pos = [(a, b) for a in range(n_l) for b in range(n_c) if grade[a][b] == cod]
            i2, j2 = max(a for a, _ in pos), max(b for _, b in pos)
            cell = tbl.cell(i, j)
            if (i2, j2) != (i, j):
                cell = cell.merge(tbl.cell(i2, j2))
            cell.text = ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(p["texto"][cod])
            r.font.size = Pt(size)
            if "■" in p["texto"][cod]:
                r.bold = True
                r.font.color.rgb = COR_ALERTA
    return tbl


# ============================================================
# ROTEIRO
# ============================================================
def build_roteiro():
    doc = Document()
    setup_page(doc)

    cabecalho_documento(
        doc,
        "REDE LOCAL EM CASA I: PLANEJANDO A INSTALAÇÃO",
        f"Roteiro do professor · Redes de Computadores · {NOITE} · Curso FIC · CEJA Itapiranga – SC",
    )

    add_heading(doc, "A IDEIA DESTA NOITE")
    add_para(doc, "Até aqui a turma aprendeu as peças: equipamentos, cabos, endereços, nomes. Hoje ela junta tudo pela primeira vez num serviço de verdade: planejar a rede de uma casa ANTES de pegar o alicate. Onde o roteador fica, o que vai no cabo, por onde o cabo passa, quanto material comprar e quanto cobrar.")
    add_para(doc, "É a noite em que o aluno passa de \"sei mexer\" pra \"sei fazer um orçamento\". E a matéria-prima é a casa dele: a planta que ele desenhou de tarefa.")
    add_para(doc, "A frase da noite: \"Roteador é lâmpada: no centro, no alto e à vista.\"", bold=True, color=COR_PRIMARIA)

    add_heading(doc, "O que se resolve nesta noite", level=2)
    add_para(doc, "Cada aluno sai sabendo fazer a visita técnica (o que perguntar e o que olhar), escolher onde o roteador fica, decidir o que vai no cabo e o que vai no Wi-Fi, medir o caminho real de um cabo, desenhar o projeto numa malha e montar a lista de materiais com o orçamento.")

    add_heading(doc, "Esta é a oitava de quinze noites", level=2)
    add_para(doc, "Abre o projeto da rede de casa, que segue na Noite 9 (configurar o roteador), na Noite 10 (cobertura de Wi-Fi) e na Noite 11 (segurança), e fecha no projeto final da Noite 15. Guarde os projetos das duplas: voltam todos.")
    add_para(doc, "Não entre em faixas de 2,4 e 5 GHz, canais nem potência de sinal: isso é a Noite 10. Hoje a regra é de olho: centro, alto, à vista, longe de metal e água.", italic=True, color=COR_ALERTA)

    add_hr(doc)

    add_heading(doc, "ANTES DA AULA")
    add_heading(doc, "Material da mesa", level=2)
    add_bullet(doc, "Fita crepe (muita): pra galeria das plantas e pra marcar os pontos de medição")
    add_bullet(doc, "Trena (uma por equipe, se der). Sem trena: barbante com um nó a cada metro, medido antes com uma régua")
    add_bullet(doc, "Réguas, lápis e lápis de cor ou canetinhas (duas cores por dupla)")
    add_bullet(doc, "Um roteador da bancada e um rolo ou pedaço de cabo de rede, só pra mostrar na mão no Bloco 4")
    add_bullet(doc, "Algumas folhas A4 em branco pra quem esqueceu a tarefa da planta")
    add_para(doc, "Não precisa de internet nem de computador hoje. Celular só como calculadora.", italic=True, color=COR_CINZA)

    add_heading(doc, "Marque os pontos de medição (Bloco 6)", level=2)
    add_para(doc, "Antes da turma chegar, cole na escola cinco pares de etiquetas de fita crepe, cada par com um ponto A e um ponto B, sempre em lugares onde o cabo não pode ir reto:")
    add_bullet(doc, "Par 1 — A: canto da sala de aula, perto da porta  ·  B: o canto oposto da mesma sala")
    add_bullet(doc, "Par 2 — A: a mesa do professor  ·  B: o fundo da sala, do outro lado (tem que contornar as carteiras pela parede)")
    add_bullet(doc, "Par 3 — A: dentro da sala  ·  B: no corredor, passando pela porta")
    add_bullet(doc, "Par 4 — A: perto de uma tomada baixa  ·  B: uma janela ou ponto alto (tem subida)")
    add_bullet(doc, "Par 5 — A: a sala de aula  ·  B: a sala vizinha (porta a porta pelo corredor)")
    add_para(doc, "Se a escola não deixar sair da sala, faça os cinco pares dentro dela, todos contornando móveis, porta ou quadro.", italic=True, color=COR_CINZA)

    add_heading(doc, "Imprima", level=2)
    add_bullet(doc, "Folha do aluno (4 páginas), uma por pessoa")
    add_bullet(doc, "Resumo do aluno, uma por pessoa (pra levar pra casa)")
    add_bullet(doc, "ANEXO 1 (clínica de plantas, 8 casas), 1 jogo por equipe")

    add_hr(doc)

    add_heading(doc, "QUADRO DE TEMPO — 210 MINUTOS")
    add_table(doc, [
        ["Bloco", "Tempo", "O que acontece"],
        ["1. Retomada + galeria das plantas", "20 min", "As plantas da tarefa na parede. Onde está o roteador de cada um?"],
        ["2. A visita técnica", "10 min", "O que o técnico pergunta e o que ele olha antes de orçar"],
        ["3. PRÁTICA — Ficha da minha casa", "20 min", "Cada um preenche a ficha de visita da própria casa. A dupla confere"],
        ["4. Onde fica o roteador", "15 min", "A lâmpada. Centro, alto, à vista. Cabo × Wi-Fi"],
        ["5. PRÁTICA — Clínica de plantas", "25 min", "Equipes analisam 8 casas: o problema, a solução, o que comprar"],
        ["INTERVALO", "15 min", ""],
        ["6. PRÁTICA — O cabo não anda reto", "20 min", "Equipes medem na escola o caminho real do cabo entre dois pontos"],
        ["7. PRÁTICA — O projeto da dupla", "30 min", "Uma das plantas na malha: entrada, roteador, cabos, metragem"],
        ["8. Lista de materiais e orçamento", "25 min", "5 min de explicação + 20 min de orçamento com a tabela de preços"],
        ["9. PRÁTICA — Banca de projetos", "20 min", "As duplas trocam de projeto e conferem pelo checklist"],
        ["10. Fechamento + tarefa", "10 min", "Quiz de saída e anúncio da Noite 9"],
    ], col_widths=[6.5, 2, 8])
    add_para(doc, "Blocos 3, 5, 6, 7, 9 e a parte prática do 8: 135 min de atividade do aluno (64% do total), mais a galeria do Bloco 1. A exposição está em três blocos curtos (30 min somados), cada um seguido de prática.", italic=True, color=COR_CINZA)

    add_hr(doc)

    # BLOCO 1
    add_heading(doc, "BLOCO 1 — RETOMADA + GALERIA DAS PLANTAS · 20 MIN")
    add_heading(doc, "Recap relâmpago — 5 min", level=2)
    add_bullet(doc, "O DNS faz o quê? (a agenda: nome vira número)")
    add_bullet(doc, "Quem é o dono de https://caixa.gov.br.consulta-fgts.com? (consulta-fgts.com)")
    add_bullet(doc, "Cadeado quer dizer que o site é bom? (não: cadeado mais dono certo)")
    add_bullet(doc, "Alguém mandou a cartilha pro grupo da família? O que responderam?")

    add_heading(doc, "A galeria — 15 min", level=2)
    add_para(doc, "Cada aluno cola a planta que desenhou de tarefa na parede, com fita crepe, numa linha só. Quem esqueceu desenha agora, em 5 minutos, numa folha em branco: vale rascunho. Depois a turma passa pela galeria olhando todas e, de volta ao lugar, responde no quadro:")
    add_bullet(doc, "\"Em quantas casas o roteador está no CANTO da casa?\" (quase todas)")
    add_bullet(doc, "\"E por que ele está ali?\" (porque foi ali que o técnico do provedor deixou: onde a fibra entra)")
    add_bullet(doc, "\"Os cômodos pintados, com Wi-Fi fraco, ficam perto ou longe do X?\" (longe)")
    add_bullet(doc, "\"Alguém tem a casa sem nenhum cômodo pintado? Onde está o roteador dele?\"")
    add_para(doc, "\"O técnico do provedor deixa o roteador onde é mais fácil pra ELE: perto da porta, onde a fibra entra. Não onde é melhor pra casa. Hoje vocês aprendem a fazer o contrário, e isso é serviço que se cobra.\"", italic=True, color=COR_PRIMARIA)
    add_para(doc, "Deixe a galeria na parede até o fim da aula: as plantas voltam no Bloco 7.", italic=True, color=COR_CINZA)

    add_hr(doc)

    # BLOCO 2
    add_heading(doc, "BLOCO 2 — A VISITA TÉCNICA · 10 MIN")
    add_para(doc, "\"Ninguém orça uma pintura sem ver a parede. Rede é igual: antes de falar preço, o técnico vai até a casa, pergunta e olha. Quem pula a visita erra a metragem, esquece o galpão, compra roteador que não precisava.\"", italic=True)
    add_table(doc, [["O quê", "O que perguntar ou olhar"]] + [list(r) for r in VISITA], col_widths=[4, 12.5])
    add_para(doc, "Ênfase em duas linhas: \"Parado ou anda\" decide cabo ou Wi-Fi (Bloco 4), e \"Onde dói\" é o que o cliente quer resolver. O resto é pra não errar o orçamento.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)

    # BLOCO 3
    add_heading(doc, "BLOCO 3 — PRÁTICA · A FICHA DA MINHA CASA · 20 MIN")
    add_para(doc, "Cada aluno faz a visita técnica da própria casa, de memória, na página 1 da folha. É a mesma ficha que o técnico leva.")
    add_bullet(doc, "12 min: cada um preenche sozinho a ficha da sua casa, incluindo a tabela de aparelhos (onde fica, parado ou anda)")
    add_bullet(doc, "8 min: em dupla, cada um lê a ficha do outro e faz duas perguntas que a ficha não respondeu. As respostas entram na ficha")
    add_para(doc, "Circule pelas duplas. As perguntas que mais aparecem: \"tem forro?\", \"onde fica a tomada?\", \"o galpão precisa de internet?\" Puxe pro quadro as que forem boas.", italic=True, color=COR_CINZA)
    add_para(doc, "Onde dilatar: pergunte quem tem casa de madeira, quem tem sobrado, quem tem galpão ou casa no interior. Cada caso diferente muda o projeto. Deixe os alunos contarem a casa deles.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)

    # BLOCO 4
    add_heading(doc, "BLOCO 4 — ONDE FICA O ROTEADOR · 15 MIN")
    add_heading(doc, "A lâmpada — 7 min", level=2)
    add_para(doc, "Pegue o roteador da bancada na mão.", italic=True, color=COR_CINZA)
    add_para(doc, "\"Imaginem que isso aqui é uma lâmpada, e vocês só podem ter UMA pra iluminar a casa inteira. Onde vocês põem? No canto da sala, perto da porta? Dentro do armário? No chão, atrás do sofá?\"", italic=True)
    add_para(doc, "Deixe a turma responder. Fica claro sozinho: no meio da casa, no alto, sem nada na frente. O Wi-Fi se espalha igual luz: pra todos os lados, e perde força em cada parede.")
    add_table(doc, [["Regra", "Por quê"]] + [list(r) for r in REGRAS_ROTEADOR], col_widths=[4.5, 12])

    add_heading(doc, "Cabo × Wi-Fi — 5 min", level=2)
    add_para(doc, "\"O que fica parado vai no cabo. O que anda vai no Wi-Fi.\"", bold=True, color=COR_PRIMARIA)
    add_bullet(doc, "Parado: TV, computador de mesa, videogame, câmera, impressora. Cabo é mais firme e não divide o sinal com ninguém")
    add_bullet(doc, "Anda: celular, notebook, tablet. Wi-Fi")
    add_bullet(doc, "Cada aparelho que sai do Wi-Fi pro cabo deixa o Wi-Fi mais folgado pros outros")

    add_heading(doc, "E se a fibra entra no canto? — 3 min", level=2)
    add_para(doc, "\"O técnico do provedor deixa a ONU onde a fibra entra. A gente não mexe na fibra: a gente leva um CABO DE REDE da ONU até o meio da casa, e lá põe o roteador. Esse cabo é o serviço de vocês.\" Mostre o pedaço de cabo.", italic=True)

    add_hr(doc)

    # BLOCO 5
    add_heading(doc, "BLOCO 5 — PRÁTICA · CLÍNICA DE PLANTAS · 25 MIN")
    add_para(doc, "No formato das clínicas de cabos (Noite 5) e de links (Noite 7). Equipes de 3. Cada equipe recebe as 8 casas do ANEXO 1 e preenche o laudo na página 2 da folha: o problema, a solução e o que precisa comprar.")
    add_heading(doc, "Regras", level=2)
    add_bullet(doc, "Leiam a queixa do cliente E olhem a planta. O ■ vermelho é onde o roteador está hoje")
    add_bullet(doc, "Pra cada casa, desenhem a lápis, no próprio cartão, onde o roteador deveria ficar")
    add_bullet(doc, "Atenção: pode ter casa em que a solução é não mexer em nada")
    add_bullet(doc, "~15 min pra equipe, depois correção")
    add_heading(doc, "Gabarito (só com você)", level=2)
    add_table(doc, [["#", "Casa", "O problema", "A solução"]] +
              [[str(i), p["titulo"], p["problema"], p["solucao"]] for i, p in enumerate(PLANTAS, 1)],
              col_widths=[0.7, 2.8, 6.2, 6.8])
    add_heading(doc, "Correção — 10 min", level=2)
    add_para(doc, "Casa por casa, uma equipe diferente dá o laudo. Pare em três:")
    add_bullet(doc, "#2 (rack): a solução é de graça. Técnico bom também resolve sem vender nada")
    add_bullet(doc, "#4 (galpão): quem pôs repetidor no meio do pátio? Repetidor precisa de tomada e de sinal bom pra repetir. O certo é cabo")
    add_bullet(doc, "#8 (nada a mexer): quem mandou trocar o roteador? A lentidão em horário fixo é do plano ou do provedor. Mexer na instalação aqui é vender o que o cliente não precisa")
    add_para(doc, "Onde dilatar: volte à galeria da parede e peça que cada equipe escolha uma planta da turma (com licença do dono) e dê o laudo dela, igual às 8 casas. O dono da planta diz se faz sentido.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)
    add_heading(doc, "INTERVALO · 15 MIN", color=COR_CINZA)
    add_hr(doc)

    # BLOCO 6
    add_heading(doc, "BLOCO 6 — PRÁTICA · O CABO NÃO ANDA RETO · 20 MIN")
    add_para(doc, "Cada equipe recebe um dos cinco pares de pontos que você marcou e mede três vezes, anotando na página 2 da folha:")
    add_bullet(doc, "1. No olho: quantos metros vocês acham que dá de A até B? (antes de medir)")
    add_bullet(doc, "2. Em linha reta, com a trena")
    add_bullet(doc, "3. Pelo caminho que o cabo faria de verdade: colado no rodapé, contornando porta e móvel, subindo e descendo na parede. Nunca no meio do chão")
    add_para(doc, "~10 min pra medir. Se der tempo, a equipe troca de par e mede outro.", italic=True, color=COR_CINZA)
    add_heading(doc, "A conta da folga — 10 min", level=2)
    add_para(doc, "Cada equipe diz no quadro os três números. Quase sempre o caminho real dá uma vez e meia a duas vezes a linha reta, e o \"no olho\" erra pra menos. Depois, a regra da folha:")
    add_para(doc, "Cabo pra comprar = caminho real + 1 m em cada ponta (subir até o aparelho) + 10% de sobra", bold=True, color=COR_PRIMARIA, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "\"Cabo curto não tem conserto: não se emenda cabo de rede. Faltou 30 cm, joga o cabo fora e passa outro. Por isso a sobra é obrigação, não desperdício.\"", italic=True, color=COR_ALERTA)

    add_hr(doc)

    # BLOCO 7
    add_heading(doc, "BLOCO 7 — PRÁTICA · O PROJETO DA DUPLA · 30 MIN")
    add_para(doc, "Cada dupla escolhe UMA das duas plantas da galeria (a que tiver mais problema pra resolver) e passa a limpo na malha quadriculada da página 3 da folha.")
    add_bullet(doc, "Escala: cada quadrado é 1 metro. O dono da planta estima o tamanho dos cômodos (\"o quarto tem uns 3 por 3\")")
    add_bullet(doc, "Marquem com letras: E = onde a internet entra, R = onde o roteador VAI ficar, C = cada aparelho que vai no cabo")
    add_bullet(doc, "Com a outra cor, desenhem o caminho de cada cabo, pelas paredes")
    add_bullet(doc, "Contem os quadrados de cada caminho e preencham a tabela de trechos, já com a regra da folga")
    add_para(doc, "Circule. Erros comuns: cabo atravessando o meio do cômodo, roteador ainda no canto \"porque é ali que a fibra chega\", esquecer o cabo da ONU até o roteador. Pergunte em vez de corrigir: \"por onde esse cabo passa?\"", italic=True, color=COR_CINZA)
    add_para(doc, "Onde dilatar: duplas que terminam cedo desenham uma segunda opção (por exemplo, com repetidor em vez de cabo) e comparam qual é melhor pro cliente.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)

    # BLOCO 8
    add_heading(doc, "BLOCO 8 — LISTA DE MATERIAIS E ORÇAMENTO · 25 MIN")
    add_heading(doc, "O que entra no orçamento — 5 min", level=2)
    add_bullet(doc, "Cabo: a metragem total da tabela de trechos, já com a folga")
    add_bullet(doc, "Conectores: 2 por cabo, mais uns 4 de reserva (sempre se perde algum na crimpagem, lembram da Noite 4?)")
    add_bullet(doc, "Canaleta: os metros do cabo que passam por fora da parede")
    add_bullet(doc, "Aparelhos: só se precisar (roteador novo, ponto de acesso, mesh)")
    add_bullet(doc, "Mão de obra: visita, cada ponto de cabo e a configuração do roteador")
    add_para(doc, "Na Noite 3 a turma pesquisou preço em loja de verdade. Hoje usa a tabela de preços de referência da página 4, pra todo mundo chegar a números comparáveis e a correção ser rápida.", italic=True, color=COR_CINZA)
    add_heading(doc, "O orçamento — 20 min", level=2)
    add_para(doc, "A dupla preenche o orçamento do próprio projeto na página 4. O celular serve de calculadora.")
    add_para(doc, "Onde dilatar: peça o orçamento em duas versões, \"econômica\" e \"completa\" (com canaleta e tomada de rede na parede), e que a dupla diga qual recomendaria e por quê.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)

    # BLOCO 9
    add_heading(doc, "BLOCO 9 — PRÁTICA · BANCA DE PROJETOS · 20 MIN")
    add_para(doc, "Cada dupla entrega o projeto (páginas 3 e 4) pra dupla do lado, que confere pelo checklist da página 4. Igual à conferência de cabo no testador: passa ou não passa.")
    add_bullet(doc, "10 min: a dupla conferente marca o checklist, dá o veredito (aprovado, aprovado com ajuste, refazer) e escreve UMA sugestão")
    add_bullet(doc, "Devolve. A dupla dona tem 3 min pra corrigir o que puder")
    add_bullet(doc, "7 min: duas ou três duplas mostram o projeto pra turma, contando o problema da casa e a solução. Escolha as plantas mais diferentes (sobrado, galpão, madeira)")
    add_para(doc, "\"Tudo que vocês conferiram hoje é o que o cliente vai perguntar: por que aí, por que tanto cabo, por que esse preço. Quem sabe responder fecha o serviço.\"", italic=True, color=COR_PRIMARIA)

    add_hr(doc)

    # BLOCO 10
    add_heading(doc, "BLOCO 10 — FECHAMENTO E TAREFA · 10 MIN")
    add_heading(doc, "Quiz de saída", level=2)
    add_bullet(doc, "Onde o roteador deve ficar? Três palavras")
    add_bullet(doc, "TV e celular: quem vai no cabo e quem vai no Wi-Fi? Por quê?")
    add_bullet(doc, "O caminho do cabo deu 12 m. Quanto cabo você compra?")
    add_bullet(doc, "Por que não dá pra comprar o cabo \"certinho\"?")

    add_heading(doc, "A tarefa", level=2)
    add_para(doc, "\"Na próxima noite vocês configuram um roteador do zero. A tarefa: olhem a etiqueta embaixo do roteador de casa e anotem na folha a marca e o modelo. Depois vejam no celular o nome do Wi-Fi da casa: é igual ao que está na etiqueta? Se for, ninguém nunca mexeu nele. E perguntem em casa se alguém sabe a senha de administrador. Não tirem foto da etiqueta nem mandem a senha pra ninguém.\"", italic=True)

    add_heading(doc, "Anuncie a próxima", level=2)
    add_para(doc, "\"Na próxima noite: configurando o roteador. Cada equipe pega um roteador de fábrica e deixa ele pronto pra entregar ao cliente: nome, senha, segurança. A segunda parte do projeto.\"", italic=True)

    add_hr(doc)

    add_heading(doc, "SE PERGUNTAREM")
    add_para(doc, "\"Posso mudar a fibra de lugar?\"", bold=True)
    add_para(doc, "> \"A fibra e a ONU são do provedor: quem mexe é o técnico deles, e geralmente cobram. O que vocês fazem é levar cabo de rede da ONU até onde o roteador vai ficar. É mais barato e é serviço de vocês.\"", italic=True)
    add_para(doc, "\"Repetidor não resolve tudo?\"", bold=True)
    add_para(doc, "> \"Repetidor repete o sinal que recebe: precisa ficar onde o sinal ainda é bom, e divide a velocidade. Ajuda em casa pequena com um cômodo ruim. Pra galpão, sobrado ou casa grande, cabo ou mesh. Na Noite 10 a gente mede isso.\"", italic=True)
    add_para(doc, "\"Até quantos metros vai o cabo de rede?\"", bold=True)
    add_para(doc, "> \"Até 100 metros por trecho. Pra casa, sobra. Pra galpão longe, ainda vai. Passou disso, precisa de outro equipamento no meio.\"", italic=True)
    add_para(doc, "\"Pode passar cabo de rede junto com fio de luz?\"", bold=True)
    add_para(doc, "> \"Não no mesmo eletroduto: a rede elétrica atrapalha o sinal e é norma separar. Canaleta ou eletroduto próprio pra rede.\"", italic=True)
    add_para(doc, "\"Quanto se cobra de verdade por ponto?\"", bold=True)
    add_para(doc, "> \"Varia muito com a cidade e a dificuldade. A tabela de hoje é de referência. Pra cobrar de verdade, perguntem a quem já faz e vejam quanto o provedor cobra pelo mesmo serviço.\"", italic=True)

    add_hr(doc)

    add_heading(doc, "DEPOIS DA AULA")
    add_bullet(doc, "Recolha os projetos das duplas (páginas 3 e 4) ou peça foto: eles voltam na Noite 10 (cobertura) e na Noite 15 (projeto final)")
    add_bullet(doc, "Tire as etiquetas de fita crepe dos pontos de medição")
    add_bullet(doc, "Anote quantas equipes acertaram a casa #8 (nada a mexer) sem ajuda. É o indicador da noite")

    add_hr(doc)

    add_heading(doc, "DESCRITIVO PARA O DIÁRIO")
    add_para(doc, "Desenvolvi atividade de planejamento de rede local doméstica. A partir das plantas residenciais desenhadas pelos estudantes, discutimos a posição do roteador instalado pelo provedor e seus efeitos na cobertura. Apresentei o roteiro de visita técnica e os critérios de posicionamento do roteador, bem como a escolha entre conexão cabeada e sem fio conforme o uso dos aparelhos. Os estudantes preencheram a ficha de visita técnica da própria residência, analisaram em equipe oito situações de instalação com laudo de problema e solução e mediram, nas dependências da escola, o percurso real de cabeamento entre pontos definidos, comparando estimativa, linha reta e trajeto real e aplicando regra de folga. Em dupla, elaboraram o projeto da rede de uma residência em malha quadriculada com escala, calcularam a metragem de cabo e montaram lista de materiais e orçamento com tabela de preços de referência. Ao final, as duplas conferiram os projetos umas das outras por checklist e apresentaram soluções à turma.")

    add_hr(doc)

    add_heading(doc, "CAIXA DE FERRAMENTAS DIDÁTICAS", color=COR_DESTAQUE)
    add_para(doc, "Se sobrar tempo, ou pra usar em outra noite:", italic=True, color=COR_CINZA)
    add_para(doc, "1. O cabo do tamanho certo", bold=True)
    add_para(doc, "A equipe que mediu o par de pontos corta e crimpa um cabo com a metragem que calculou (com a folga) e passa ele de verdade pelo caminho, de A até B, com fita crepe no rodapé. Chegou? Sobrou quanto? Depois testa no testador. Junta Noites 4, 5 e 8. 30 min.")
    add_para(doc, "2. A planta da escola", bold=True)
    add_para(doc, "A turma desenha na malha a planta do corredor e das salas do CEJA e decide onde ficaria um roteador pra cobrir tudo. Guarde: na Noite 10 a turma mede o sinal de verdade e vê se acertou. 25 min.")
    add_para(doc, "3. Preço de verdade", bold=True)
    add_para(doc, "Cada dupla pesquisa no celular o preço real de três itens do seu orçamento (cabo, roteador, canaleta) em duas lojas, como na Noite 3, e compara com a tabela de referência. Ficou mais caro ou mais barato? 15 min.")
    add_para(doc, "4. Cartaz das regras do roteador", bold=True)
    add_para(doc, "Em equipe, um cartaz pra família: \"Onde colocar o roteador\", com as cinco regras desenhadas (sem palavra técnica). Apresentação de 1 minuto por equipe. 20 min.")
    add_para(doc, "5. A casa de dois andares", bold=True)
    add_para(doc, "Desafio pra quem terminou: a dupla projeta o sobrado da casa #3 na malha, os dois andares lado a lado, com o caminho do cabo subindo pela escada, e orça. 20 min.")

    # ANEXO 1 — clínica de plantas
    for k in range(0, len(PLANTAS), 2):
        doc.add_page_break()
        add_heading(doc, f"ANEXO 1 — CLÍNICA DE PLANTAS · CASAS {k + 1} E {k + 2} (1 jogo por equipe)")
        add_para(doc, "■ vermelho = onde o roteador está hoje. Desenhe a lápis onde ele deveria ficar.", italic=True, size=10, color=COR_CINZA)
        for i in (k, k + 1):
            p = PLANTAS[i]
            add_para(doc, f"CASA {i + 1} — {p['titulo'].upper()}", bold=True, size=13, color=COR_PRIMARIA)
            add_para(doc, f"O cliente diz: {p['queixa']}", italic=True, size=11)
            planta(doc, p, altura=2.2, largura=5.6, size=10)
            doc.add_paragraph()

    path = os.path.join(OUTDIR, "Roteiro - Redes 8 - Planejando a rede de casa.docx")
    doc.save(path)
    return path


# ============================================================
# SLIDES
# ============================================================
def build_slides():
    prs = new_presentation()
    P, D, K, R = PPalette.P, PPalette.D, PPalette.K, PPalette.R

    def lista(kicker, titulo, linhas, subtitulo=None, size=20, nota=None, nota_cor=None):
        s = add_slide(prs)
        add_title_block(s, kicker, titulo, subtitulo)
        add_text(s, linhas, 0.7, 2.8, 11.9, 3.8, size=size, color=K)
        if nota:
            add_text(s, nota, 0.7, 6.4, 11.9, 0.5, size=16, italic=True, color=nota_cor or D, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    def frase(kicker, texto, cor=None, size=36):
        s = add_slide(prs)
        add_text(s, kicker, 0.7, 1.6, 11.9, 0.5, size=16, bold=True, color=D, align=PP_ALIGN.CENTER)
        add_text(s, texto.split("\n"), 0.9, 2.4, 11.5, 3.0, size=size, bold=True, color=cor or P, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    def tabela(kicker, titulo, rows, col_widths=None, subtitulo=None, size=16, nota=None, height=3.6):
        s = add_slide(prs)
        add_title_block(s, kicker, titulo, subtitulo)
        add_table_slide(s, rows, top=2.6, height=height, col_widths=col_widths, size=size)
        if nota:
            add_text(s, nota, 0.7, 6.5, 11.9, 0.5, size=15, italic=True, color=D, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    add_slide_cover(prs, "NOITE 8", "A REDE", "DA CASA",
                    "Parte 1: planejando a instalação",
                    "Redes de Computadores · Curso FIC · CEJA Itapiranga – SC")

    lista("RETOMADA", "Da Noite 7 pra hoje", [
        "•  O DNS faz o quê?",
        "•  Quem é o dono de https://caixa.gov.br.consulta-fgts.com?",
        "•  Cadeado quer dizer que o site é bom?",
        "•  Alguém mandou a cartilha pra família?",
    ])

    lista("A GALERIA", "Cole a planta da sua casa na parede", [
        "Passe pela galeria e olhe todas. Depois:",
        "",
        "•  Em quantas casas o roteador está no CANTO?",
        "•  Por que ele está ali?",
        "•  Os cômodos com Wi-Fi fraco ficam perto ou longe do X?",
    ])

    frase("O QUE A GALERIA MOSTRA", "O técnico do provedor deixa o roteador\nonde é fácil pra ELE.\nNão onde é bom pra casa.", size=30)

    tabela("A VISITA TÉCNICA", "Antes de falar preço, pergunte e olhe",
           [["O quê", "O que perguntar ou olhar"]] + [list(r) for r in VISITA],
           col_widths=[2.9, 9.0], size=13, height=4.0)

    lista("PRÁTICA · 20 MIN", "A ficha da minha casa", [
        "•  Página 1 da folha: a visita técnica da SUA casa",
        "•  12 min sozinho",
        "•  8 min em dupla: leia a ficha do outro e faça duas perguntas",
        "   que ela não respondeu",
    ], size=22)

    frase("UMA PERGUNTA", "Se você só pudesse ter UMA lâmpada\npra iluminar a casa inteira,\nonde você colocaria?", size=32)

    frase("A FRASE DA NOITE", "Roteador é lâmpada:\nno centro, no alto e à vista.")

    tabela("ONDE FICA O ROTEADOR", "Cinco regras", [["Regra", "Por quê"]] + [list(r) for r in REGRAS_ROTEADOR],
           col_widths=[3.4, 8.5], size=15, height=3.6)

    frase("CABO × WI-FI", "O que fica parado vai no cabo.\nO que anda vai no Wi-Fi.")

    lista("CABO × WI-FI", "Quem vai onde", [
        "•  Parado: TV, computador, videogame, câmera, impressora  →  CABO",
        "•  Anda: celular, notebook, tablet  →  WI-FI",
        "",
        "Cada aparelho que vai pro cabo deixa o Wi-Fi mais folgado pros outros.",
    ], size=21)

    frase("E SE A FIBRA ENTRA NO CANTO?", "A gente não mexe na fibra.\nLeva um cabo de rede da ONU\naté o meio da casa.\nEsse cabo é o serviço de vocês.", size=30)

    lista("PRÁTICA · 25 MIN", "Clínica de plantas", [
        "•  Equipes de 3. Oito casas",
        "•  Leia a queixa do cliente e olhe a planta (■ = roteador hoje)",
        "•  Desenhe a lápis onde ele deveria ficar",
        "•  Laudo na página 2: problema, solução, o que comprar",
    ], size=21, nota="Atenção: pode ter casa em que a solução é não mexer em nada.", nota_cor=R)

    frase("", "INTERVALO\n15 minutos", cor=K)

    lista("PRÁTICA · 20 MIN", "O cabo não anda reto", [
        "Sua equipe recebe um par de pontos: A e B. Meça três vezes:",
        "",
        "1.  No olho: quantos metros você acha que dá?",
        "2.  Em linha reta, com a trena",
        "3.  Pelo caminho do cabo: rodapé, contornando porta e móvel,",
        "     subindo e descendo. Nunca no meio do chão",
    ], size=20)

    frase("A CONTA DA FOLGA", "Cabo pra comprar =\ncaminho real\n+ 1 m em cada ponta\n+ 10% de sobra", size=32)

    frase("CUIDADO", "Cabo de rede não se emenda.\nFaltou 30 cm: joga fora e passa outro.\nSobra é obrigação, não desperdício.", cor=R, size=30)

    lista("PRÁTICA · 30 MIN", "O projeto da dupla", [
        "•  Escolham UMA das duas plantas e passem pra malha (página 3)",
        "•  Cada quadrado é 1 metro",
        "•  E = entrada da internet   ·   R = roteador   ·   C = aparelho no cabo",
        "•  Com outra cor: o caminho de cada cabo, pelas paredes",
        "•  Contem os quadrados e preencham a tabela de trechos",
    ], size=20)

    lista("O ORÇAMENTO", "O que não pode faltar", [
        "•  Cabo: a metragem com a folga",
        "•  Conectores: 2 por cabo + 4 de reserva",
        "•  Canaleta: o cabo que passa por fora da parede",
        "•  Aparelhos: só se precisar",
        "•  Mão de obra: visita + cada ponto + configuração do roteador",
    ], size=21, nota="Preços: a tabela de referência da página 4.")

    tabela("TABELA DE REFERÊNCIA", "Preços pra atividade (não são de loja)",
           [["Item", "Unid.", "Preço"]] + [[n, u, reais(v)] for n, u, v in PRECOS[:8]],
           col_widths=[7.4, 2.2, 2.3], size=13, height=3.9)

    tabela("TABELA DE REFERÊNCIA", "Aparelhos e mão de obra",
           [["Item", "Unid.", "Preço"]] + [[n, u, reais(v)] for n, u, v in PRECOS[8:]],
           col_widths=[7.4, 2.2, 2.3], size=14, height=3.6)

    lista("PRÁTICA · 20 MIN", "Banca de projetos", [
        "•  Entregue o projeto pra dupla do lado",
        "•  Ela confere pelo checklist da página 4",
        "•  Veredito: aprovado · aprovado com ajuste · refazer",
        "•  E escreve UMA sugestão",
        "•  3 min pra corrigir. Depois, algumas duplas apresentam",
    ], size=21)

    lista("O QUE FICA", "Três frases pra guardar", [
        "\"Roteador é lâmpada: no centro, no alto e à vista.\"",
        "",
        "\"O que fica parado vai no cabo. O que anda vai no Wi-Fi.\"",
        "",
        "\"Cabo não anda reto. E sobra é obrigação.\"",
    ], size=23)

    lista("PARA A PRÓXIMA NOITE", "Configurando o roteador", [
        "A tarefa:",
        "",
        "•  Anote a marca e o modelo do roteador de casa (etiqueta embaixo)",
        "•  O nome do Wi-Fi da casa é igual ao da etiqueta?",
        "•  Alguém em casa sabe a senha de administrador?",
    ], nota="Não tire foto da etiqueta nem mande a senha pra ninguém.", nota_cor=R)

    path = os.path.join(OUTDIR, "Aula - Redes 8 - Planejando a rede de casa.pptx")
    prs.save(path)
    return path


# ============================================================
# FOLHA DO ALUNO
# ============================================================
def linhas(doc, n):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("_" * 100)
        r.font.size = Pt(11)
        r.font.color.rgb = COR_CINZA


def pergunta(doc, texto, size=11):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.font.size = Pt(size)
    r.bold = True


def campo(doc, texto, size=11, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(texto)
    r.font.size = Pt(size)
    return p


def build_folha():
    doc = Document()
    setup_page(doc, margem=Cm(1.8))

    # PÁGINA 1 — Ficha de visita técnica da minha casa
    cabecalho_documento(
        doc,
        "VISITA TÉCNICA: A MINHA CASA",
        f"Redes de Computadores · {NOITE} · Curso FIC · CEJA Itapiranga · 1 de 4",
    )
    campo(doc, "Nome: _________________________________________________     Dupla com: ________________", space_after=8)

    add_heading(doc, "1  A casa", level=2)
    campo(doc, "Quantas pessoas moram: ____    Alguém trabalha ou estuda de casa?  ☐ Sim  ☐ Não")
    campo(doc, "A casa é:  ☐ térrea  ☐ sobrado     ☐ alvenaria  ☐ madeira  ☐ mista     Tem laje?  ☐ Sim  ☐ Não")
    campo(doc, "Tem galpão, garagem ou área externa que precisa de internet?  ☐ Sim  ☐ Não   Qual? ______________")
    campo(doc, "Tem forro ou rodapé por onde um cabo pode passar?  ☐ Sim  ☐ Não  ☐ Não sei")

    add_heading(doc, "2  A internet hoje", level=2)
    campo(doc, "Por onde a internet entra na casa (em que cômodo, em que parede)? ______________________________")
    campo(doc, "Onde está o roteador hoje? ________________________   Fica:  ☐ no alto  ☐ no chão  ☐ dentro de móvel")
    campo(doc, "Cômodos onde o Wi-Fi é fraco ou cai: ______________________________________________________")
    campo(doc, "Fica lento em algum horário? ☐ Não  ☐ Sim, às ______ horas")

    add_heading(doc, "3  Os aparelhos", level=2)
    add_para(doc, "Parado = fica sempre no mesmo lugar. Anda = vai de um cômodo pro outro.", italic=True, size=10, color=COR_CINZA)
    tbl = add_table(doc, [
        ["Aparelho", "Quantos", "Em que cômodo", "Parado ou anda?", "Cabo ou Wi-Fi?"],
        ["Celular", "", "", "", ""],
        ["TV", "", "", "", ""],
        ["Computador / notebook", "", "", "", ""],
        ["Videogame", "", "", "", ""],
        ["Câmera / impressora / outro", "", "", "", ""],
    ], col_widths=[4.6, 1.8, 4.0, 3.3, 3.3])
    for row in tbl.rows[1:]:
        row.height = Cm(0.8)

    add_heading(doc, "4  As perguntas da dupla", level=2)
    add_para(doc, "Sua dupla leu sua ficha e fez duas perguntas que ela não respondia. Anote as perguntas e as respostas.", italic=True, size=10, color=COR_CINZA)
    for i in (1, 2):
        campo(doc, f"Pergunta {i}: ___________________________________________________________________________")
        campo(doc, "Resposta: ____________________________________________________________________________", space_after=8)

    add_heading(doc, "5  Onde fica o roteador", level=2)
    campo(doc, "Roteador é lâmpada: no ______________, no ______________ e à ______________.")
    campo(doc, "O que fica ______________ vai no cabo. O que ______________ vai no Wi-Fi.")

    # PÁGINA 2 — Clínica de plantas + medição
    doc.add_page_break()
    cabecalho_documento(
        doc,
        "CLÍNICA DE PLANTAS E O CABO QUE NÃO ANDA RETO",
        f"Laudo da equipe · {NOITE} · 2 de 4",
    )
    campo(doc, "Equipe: ______________________________________________________________")

    add_heading(doc, "1  Clínica de plantas", level=2)
    add_para(doc, "Pra cada casa do cartão: qual o problema, onde o roteador deveria ficar e o que precisa comprar.", italic=True, size=10, color=COR_CINZA)
    tbl = add_table(doc, [["#", "O problema", "A solução", "Comprar"]] +
                    [[f"{i}", "", "", "☐ nada  ☐ cabo\n☐ aparelho"] for i in range(1, len(PLANTAS) + 1)],
                    col_widths=[0.7, 6.4, 6.4, 3.0])
    for row in tbl.rows[1:]:
        row.height = Cm(1.0)
    campo(doc, "Qual casa enganou a equipe? ______   Por quê? ________________________________________________")

    add_heading(doc, "2  O cabo não anda reto", level=2)
    campo(doc, "Par de pontos que a equipe mediu: ☐ 1  ☐ 2  ☐ 3  ☐ 4  ☐ 5")
    tbl = add_table(doc, [
        ["Medida", "Metros"],
        ["1. No olho (antes de medir)", ""],
        ["2. Em linha reta, com a trena", ""],
        ["3. Pelo caminho do cabo (rodapé, porta, subida)", ""],
        ["Quantas vezes o caminho real é maior que a linha reta? (3 ÷ 2)", ""],
    ], col_widths=[13, 4.5])
    for row in tbl.rows[1:]:
        row.height = Cm(0.8)

    add_para(doc, "Cabo pra comprar = caminho real + 1 m em cada ponta + 10% de sobra", bold=True, size=11, color=COR_PRIMARIA, align=WD_ALIGN_PARAGRAPH.CENTER)
    campo(doc, "Faça a conta do seu par:   ________ + 2 m = ________   +10% = ________ m de cabo")
    pergunta(doc, "Por que não dá pra comprar o cabo \"certinho\", sem sobra?")
    linhas(doc, 2)

    # PÁGINA 3 — O projeto na malha
    doc.add_page_break()
    cabecalho_documento(
        doc,
        "O PROJETO DA DUPLA",
        f"A rede de uma casa · {NOITE} · 3 de 4",
    )
    campo(doc, "Dupla: ______________________________________   Casa de: __________________________")
    add_para(doc, "Cada quadrado = 1 metro.  E = entrada da internet  ·  R = roteador  ·  C = aparelho no cabo. Com outra cor, o caminho de cada cabo, pelas paredes.", italic=True, size=10, color=COR_CINZA)

    malha = doc.add_table(rows=14, cols=17)
    malha.style = 'Table Grid'
    for row in malha.rows:
        row.height = Cm(1.0)
        for cell in row.cells:
            cell.width = Cm(1.0)

    add_heading(doc, "Os trechos de cabo", level=2)
    tbl = add_table(doc, [
        ["Trecho (de onde até onde)", "Quadrados", "+ 2 m pontas", "+ 10%", "Total (m)"],
        ["ONU → roteador", "", "", "", ""],
        ["Roteador → ", "", "", "", ""],
        ["Roteador → ", "", "", "", ""],
        ["Roteador → ", "", "", "", ""],
        ["TOTAL DE CABO", "", "", "", ""],
    ], col_widths=[6.5, 2.6, 2.8, 2.4, 3.2])
    for row in tbl.rows[1:]:
        row.height = Cm(0.75)

    # PÁGINA 4 — Orçamento + banca + tarefa
    doc.add_page_break()
    cabecalho_documento(
        doc,
        "ORÇAMENTO E BANCA",
        f"Lista de materiais · {NOITE} · 4 de 4",
    )

    add_heading(doc, "1  Tabela de preços de referência", level=2)
    add_para(doc, "Valores pra atividade em sala, não de loja. No serviço de verdade, pesquise (como na Noite 3).", italic=True, size=9, color=COR_CINZA)
    metade = (len(PRECOS) + 1) // 2
    esq, dir_ = PRECOS[:metade], PRECOS[metade:]
    rows = [["Item", "Preço", "Item", "Preço"]]
    for i in range(metade):
        a = esq[i]
        b = dir_[i] if i < len(dir_) else ("", "", None)
        rows.append([f"{a[0]} ({a[1]})", reais(a[2]),
                     f"{b[0]} ({b[1]})" if b[0] else "", reais(b[2]) if b[2] is not None else ""])
    tbl = add_table(doc, rows, col_widths=[6.4, 2.2, 6.7, 2.2])
    for row in tbl.rows:
        for cell in row.cells:
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(8.5)

    add_heading(doc, "2  O orçamento do projeto", level=2)
    tbl = add_table(doc, [
        ["Item", "Qtd", "Preço unit.", "Total"],
        ["Cabo de rede", "", "", ""],
        ["Conectores RJ45 (2 por cabo + 4)", "", "", ""],
        ["Canaleta", "", "", ""],
        ["", "", "", ""],
        ["Mão de obra: visita", "1", "", ""],
        ["Mão de obra: pontos de cabo", "", "", ""],
        ["Mão de obra: configuração do roteador", "", "", ""],
        ["TOTAL", "", "", ""],
    ], col_widths=[8.5, 2, 3.4, 3.6])
    for row in tbl.rows[1:]:
        row.height = Cm(0.62)

    add_heading(doc, "3  Banca: conferência da dupla vizinha", level=2)
    campo(doc, "Projeto da dupla: ____________________________   Conferido por: ____________________________", size=10)
    meio = (len(CHECKLIST) + 1) // 2
    tbl = add_table(doc, [[CHECKLIST[i], "☐", CHECKLIST[i + meio] if i + meio < len(CHECKLIST) else "", "☐"]
                          for i in range(meio)], header=False, col_widths=[8.0, 0.7, 8.0, 0.7])
    for row in tbl.rows:
        for cell in row.cells:
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(8.5)
    campo(doc, "Veredito:  ☐ aprovado   ☐ aprovado com ajuste   ☐ refazer", size=10)
    campo(doc, "Uma sugestão: _________________________________________________________________________", size=10)

    add_heading(doc, "TAREFA DA SEMANA", level=2, color=COR_DESTAQUE)
    campo(doc, "Roteador de casa — marca e modelo (etiqueta embaixo): _____________________________________", size=10)
    campo(doc, "O nome do Wi-Fi da casa é igual ao da etiqueta?  ☐ Sim  ☐ Não     Alguém sabe a senha de administrador?  ☐ Sim  ☐ Não", size=10)
    add_para(doc, "Não tire foto da etiqueta nem mande a senha pra ninguém. Traga esta folha na Noite 9 · CEJA Itapiranga – SC", italic=True, size=9, color=COR_CINZA, align=WD_ALIGN_PARAGRAPH.CENTER)

    path = os.path.join(OUTDIR, "Folha do Aluno - Redes 8 - Planejando a rede de casa.docx")
    doc.save(path)
    return path


# ============================================================
# RESUMO DO ALUNO
# ============================================================
def build_resumo():
    doc = Document()
    setup_page(doc, margem=Cm(1.8))

    cabecalho_documento(
        doc,
        "A REDE DA CASA: PLANEJANDO A INSTALAÇÃO",
        f"Resumo do aluno · {NOITE} · Redes de Computadores · Curso FIC · CEJA Itapiranga",
    )

    add_heading(doc, "A ideia principal", level=2)
    add_para(doc, "Rede boa começa antes do alicate: visita, planta, posição do roteador, caminho do cabo, lista de materiais e orçamento. O técnico do provedor deixa o roteador onde é fácil pra ele; o seu serviço é deixar onde é bom pra casa.")

    add_heading(doc, "A visita técnica: o que perguntar e o que olhar", level=2)
    add_table(doc, [["O quê", "O que perguntar ou olhar"]] + [list(r) for r in VISITA], col_widths=[3.8, 11.7])

    add_heading(doc, "Onde fica o roteador", level=2)
    add_table(doc, [["Regra", "Por quê"]] + [list(r) for r in REGRAS_ROTEADOR], col_widths=[4.2, 11.3])
    add_bullet(doc, "A fibra entra no canto? Não mexa nela: leve um cabo de rede da ONU até o meio da casa e ponha o roteador lá")

    add_heading(doc, "Cabo ou Wi-Fi", level=2)
    add_bullet(doc, "Fica parado (TV, computador, videogame, câmera, impressora): cabo")
    add_bullet(doc, "Anda (celular, notebook, tablet): Wi-Fi")
    add_bullet(doc, "Galpão ou cômodo longe: cabo até lá (até 100 m) e um aparelho na ponta. Repetidor no meio do pátio é gambiarra")

    add_heading(doc, "Quanto cabo comprar", level=2)
    add_bullet(doc, "Meça o caminho de verdade: pelo rodapé, forro ou canaleta, contornando portas. Nunca no meio do chão")
    add_bullet(doc, "O caminho real costuma dar uma vez e meia a duas vezes a linha reta")
    add_para(doc, "Cabo pra comprar = caminho real + 1 m em cada ponta + 10% de sobra", bold=True, color=COR_PRIMARIA)
    add_bullet(doc, "Cabo de rede não se emenda. Faltou, passa outro: por isso a sobra")
    add_bullet(doc, "Conectores: 2 por cabo, mais uns 4 de reserva")

    add_heading(doc, "O que entra no orçamento", level=2)
    add_bullet(doc, "Cabo, conectores, canaleta ou eletroduto")
    add_bullet(doc, "Aparelhos, só se precisar (roteador, ponto de acesso, mesh)")
    add_bullet(doc, "Mão de obra: visita, cada ponto de cabo e a configuração do roteador")
    add_bullet(doc, "Se a instalação já está certa, não venda nada: lentidão em horário fixo é do plano ou do provedor")

    add_heading(doc, "Três frases pra guardar", level=2)
    add_para(doc, "\"Roteador é lâmpada: no centro, no alto e à vista.\"", bold=True, color=COR_PRIMARIA)
    add_para(doc, "\"O que fica parado vai no cabo. O que anda vai no Wi-Fi.\"", bold=True, color=COR_PRIMARIA)
    add_para(doc, "\"Cabo não anda reto. E sobra é obrigação.\"", bold=True, color=COR_PRIMARIA)

    add_heading(doc, "Pra se testar em casa", level=2)
    add_bullet(doc, "Onde está o roteador da sua casa? Ele obedece às cinco regras? O que você mudaria?")
    add_bullet(doc, "O caminho do cabo até o quarto dá 9 m. Quanto cabo você compra?")
    add_bullet(doc, "A internet de um cliente fica lenta todo dia às 20 h, e o roteador está bem posicionado. O que você faz?")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Este material é seu. Guarde e volte quando precisar. · CEJA Itapiranga – SC")
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = COR_CINZA

    path = os.path.join(OUTDIR, "Resumo do Aluno - Redes 8 - Planejando a rede de casa.docx")
    doc.save(path)
    return path


# ============================================================
# MAIN
# ============================================================
if __name__ == "__main__":
    print("Gerando Roteiro...")
    print("  →", build_roteiro())
    print("Gerando Slides...")
    print("  →", build_slides())
    print("Gerando Folha do Aluno...")
    print("  →", build_folha())
    print("Gerando Resumo do Aluno...")
    print("  →", build_resumo())
    print("\nPronto.")
