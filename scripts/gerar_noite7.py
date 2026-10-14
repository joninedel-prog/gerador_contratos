"""Gera os 4 arquivos da Noite 7 do curso de Redes.

Noite 7 — Nomes na rede: DNS, URL e o que o HTTPS resolve
Padrão prática-pesada, sem role-play e sem duelo. Parte da tarefa da Noite 6
(três endereços de sites anotados com https://). A turma encena o DNS ("agenda viva"),
aprende a achar o dono de um link lendo o endereço, faz uma clínica de links
(verdadeiro × golpe), passa por estações (ping por nome × por número, DNS sabotado,
cadeado no celular), fabrica links falsos pra outra equipe desmascarar e fecha
com uma cartilha anti-golpe pra família.
Precisa de um pouco de internet: hotspot do celular do professor no computador
e o 4G dos alunos na estação do cadeado.
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from common import (
    COR_PRIMARIA, COR_ALERTA, COR_DESTAQUE, COR_CINZA,
    add_heading, add_para, add_bullet, add_table, add_hr,
    setup_page, cabecalho_documento,
    PPalette, new_presentation, add_slide, add_text, add_title_block,
    add_slide_cover, add_footer, add_table_slide,
)
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PInches

OUTDIR = "/home/nedel/curso"
NOITE = "Noite 7"
FOOTER = f"Redes de Computadores · {NOITE}"


# A analogia da agenda do celular
AGENDA = [
    ("O nome do contato (\"Mãe\")", "O nome do site (google.com.br)"),
    ("O número de telefone", "O endereço IP do site (Noite 6)"),
    ("A agenda que transforma nome em número", "O DNS"),
    ("Contato salvo errado", "DNS adulterado: você chama o nome certo e cai no lugar errado"),
    ("Agenda apagada", "DNS fora do ar: você sabe o nome, mas não acha o número"),
    ("Últimas ligações (o celular lembra)", "O cache: o aparelho guarda os números que já procurou"),
]

# Os sites da agenda viva — números inventados pra brincadeira
SITES_VIVOS = [
    ("google.com.br", "142.250.0.10", "Página do Google: uma caixa de busca"),
    ("youtube.com", "142.250.0.30", "Página do YouTube: um vídeo de gatinho"),
    ("bb.com.br", "170.66.0.20", "Página do banco: \"Digite sua senha\" + SELO: certificado de bb.com.br"),
    ("itapiranga.sc.gov.br", "200.19.0.40", "Página da Prefeitura: notícias da cidade"),
    ("mercadolivre.com.br", "54.94.0.50", "Página da loja: ofertas"),
]
GOLPISTA = ("(o golpista)", "66.66.0.66", "Página IGUALZINHA à do banco: \"Digite sua senha\" — SEM selo")

# Clínica de links: (link, dono, veredito, explicação)
LINKS = [
    ("https://www.bb.com.br/pbb/pagina-inicial", "bb.com.br", "Verdadeiro",
     "Antes da primeira barra termina em bb.com.br."),
    ("https://bb.com.br.atualizacao-cadastral.com/login", "atualizacao-cadastral.com", "Golpe",
     "O \"bb.com.br\" no começo é enfeite. Quem manda é o final: atualizacao-cadastral.com."),
    ("https://www.caixa.gov.br/beneficios", "caixa.gov.br", "Verdadeiro",
     ".gov.br só órgão público consegue registrar."),
    ("https://caixa-gov-br.app/fgts-liberado", "caixa-gov-br.app", "Golpe",
     "Traço no lugar do ponto. Termina em .app, não em .gov.br."),
    ("http://www.itapiranga.sc.gov.br/noticias", "itapiranga.sc.gov.br", "Verdadeiro, sem cadeado",
     "O dono é a Prefeitura, mas é http (sem S): o navegador mostra \"Não seguro\". Pode ler; não digite senha."),
    ("https://www.mercadolivre.com.br/ofertas", "mercadolivre.com.br", "Verdadeiro",
     "Termina em mercadolivre.com.br."),
    ("https://mercadolivre.ofertas-relampago.shop/iphone", "ofertas-relampago.shop", "Golpe",
     "\"mercadolivre\" está no começo, enfeitando. O dono é ofertas-relampago.shop."),
    ("https://www.gov.br/receitafederal/restituicao", "gov.br", "Verdadeiro",
     "O dono é gov.br. Tudo depois da barra é só o caminho dentro do site."),
    ("https://gov-br.restituicao-consulta.online", "restituicao-consulta.online", "Golpe",
     "\"gov-br\" com traço não é gov.br. O final .online entrega."),
    ("https://www.correios.com.br/rastreamento", "correios.com.br", "Verdadeiro",
     "Termina em correios.com.br."),
    ("https://correios.taxa-liberacao.info/pagar", "taxa-liberacao.info", "Golpe",
     "O golpe da \"taxa pra liberar a encomenda\". O dono é taxa-liberacao.info."),
    ("https://www.arnazon.com.br", "arnazon.com.br", "Golpe",
     "Letra trocada: \"rn\" de longe parece \"m\". Leia devagar."),
    ("https://wa.me/5549999990000", "wa.me", "Verdadeiro",
     "Parece estranho, mas wa.me é o link oficial do WhatsApp. Estranho não é sinônimo de golpe."),
    ("https://bit.ly/3xYz9Kq", "bit.ly (encurtador)", "Não dá pra saber",
     "Encurtador esconde o destino. Não abra sem saber quem mandou e pra onde vai."),
    ("https://nubank.com.br@login-seguro.net/app", "login-seguro.net", "Golpe",
     "Pegadinha do @: tudo antes do @ é ignorado. O dono é o que vem depois."),
    ("https://drive.google.com/file/d/1AbC/view", "google.com", "Verdadeiro (com cuidado)",
     "O dono é o Google, mas qualquer pessoa guarda arquivo no Drive. Site certo não garante arquivo bom."),
]

# Telas da estação A
TELAS = [
    ("Ping pelo número",
     ["C:\\> ping 8.8.8.8",
      "Resposta de 8.8.8.8: bytes=32 tempo=18ms TTL=117",
      "Resposta de 8.8.8.8: bytes=32 tempo=17ms TTL=117"]),
    ("Ping pelo nome — DNS funcionando",
     ["C:\\> ping google.com",
      "Disparando google.com [142.250.79.46] com 32 bytes de dados:",
      "Resposta de 142.250.79.46: bytes=32 tempo=19ms TTL=117"]),
    ("Ping pelo nome — DNS quebrado",
     ["C:\\> ping google.com",
      "A solicitação ping não pôde encontrar o host google.com.",
      "Verifique o nome e tente novamente."]),
]


# ============================================================
# ROTEIRO
# ============================================================
def build_roteiro():
    doc = Document()
    setup_page(doc)

    cabecalho_documento(
        doc,
        "NOMES NA REDE: DNS, URL E O QUE O HTTPS RESOLVE",
        f"Roteiro do professor · Redes de Computadores · {NOITE} · Curso FIC · CEJA Itapiranga – SC",
    )

    # A IDEIA
    add_heading(doc, "A IDEIA DESTA NOITE")
    add_para(doc, "Na Noite 6 a turma descobriu que todo aparelho tem um número (o IP). Mas ninguém digita número pra entrar num site — digita nome. Hoje a turma descobre quem transforma o nome em número (o DNS), aprende a ler um endereço de site pra saber quem é o dono de verdade, e entende o que o cadeado do navegador protege — e o que ele NÃO protege.")
    add_para(doc, "É a noite mais útil pra vida deles fora do curso: quase todo golpe por WhatsApp e SMS é um link. Quem sabe ler o endereço não cai, e ainda protege a família.")
    add_para(doc, "A frase da noite: \"O dono do link é o que vem antes da primeira barra, lido de trás pra frente.\"", bold=True, color=COR_PRIMARIA)

    add_heading(doc, "O que se resolve nesta noite", level=2)
    add_para(doc, "Cada aluno sai sabendo explicar o DNS com a analogia da agenda do celular, achar o dono de um link, separar link verdadeiro de golpe, conferir o cadeado e o certificado no celular, e reconhecer o defeito de DNS no campo: \"o número responde, o nome não\".")

    add_heading(doc, "Esta é a sétima de quinze noites", level=2)
    add_para(doc, "Fecha o bloco do \"invisível\" (Noites 6 e 7). A Noite 8 começa o projeto da rede de casa (planejamento). O DNS volta na Noite 9, quando a turma configura o roteador, e a segurança volta inteira na Noite 11 — hoje é só a parte do link e do cadeado.")
    add_para(doc, "Não entre em criptografia, chaves, tipos de registro de DNS nem servidores raiz. Pra técnico de campo e pra família, basta: agenda, dono do link, cadeado.", italic=True, color=COR_ALERTA)

    add_hr(doc)

    # ANTES DA AULA
    add_heading(doc, "ANTES DA AULA")

    add_heading(doc, "Material da mesa", level=2)
    add_bullet(doc, "1 ou 2 computadores/notebooks (os da sala) e o seu celular")
    add_bullet(doc, "Internet pro computador: o hotspot do seu celular (Wi-Fi do computador LIGADO hoje, ao contrário das noites 5 e 6). Se a escola tiver Wi-Fi com internet, melhor ainda")
    add_bullet(doc, "Pra agenda viva (Bloco 3): o ANEXO 1 impresso — a Agenda, as placas dos sites e os cartões de pedido. Fita crepe")
    add_bullet(doc, "Pra clínica de links (Bloco 5): o ANEXO 2 impresso e recortado — um jogo de 16 cartões por equipe")
    add_bullet(doc, "Pra cartilha (Bloco 9): uma folha A4 ou meia cartolina por equipe e canetões")

    add_heading(doc, "Prepare as estações (Bloco 7)", level=2)
    add_bullet(doc, "ESTAÇÃO A — computador com internet funcionando normal. Teste: ping 8.8.8.8, ping google.com e nslookup google.com")
    add_bullet(doc, "ESTAÇÃO B — o DNS sabotado. No segundo computador: Painel de Controle → Central de Rede → Alterar configurações do adaptador → (o adaptador que está com internet) → Propriedades → Protocolo IP Versão 4 → \"Usar os seguintes endereços de servidor DNS\" → 192.168.0.250 (um número que ninguém tem). OK. Depois, no Prompt: ipconfig /flushdns")
    add_bullet(doc, "Teste a Estação B: ping 8.8.8.8 tem que RESPONDER e ping google.com tem que dar \"não pôde encontrar o host\". O navegador não abre nenhum site")
    add_bullet(doc, "Só tem um computador? Faça a Estação B no seu celular Android: Configurações → Rede e internet → DNS privado → \"Nome do host do provedor\" → digite dns.naoexiste.com.br. O celular passa a mostrar \"Conectado, sem internet\" ou \"Não é possível acessar o servidor DNS privado\"")
    add_bullet(doc, "ESTAÇÃO C — não precisa preparar: é o celular dos alunos, com 4G. Gasta poucos dados")
    add_para(doc, "Combine com a turma: os links da clínica são SÓ NO PAPEL. Ninguém digita nenhum deles no celular. Alguns podem até existir de verdade.", bold=True, color=COR_ALERTA)

    add_heading(doc, "Imprima", level=2)
    add_bullet(doc, "Folha do aluno (3 páginas), uma por pessoa")
    add_bullet(doc, "Resumo do aluno, uma por pessoa (pra levar pra casa)")
    add_bullet(doc, "ANEXO 1 (agenda viva), 1 jogo. ANEXO 2 (cartões de links), 1 jogo por equipe")

    add_hr(doc)

    # QUADRO DE TEMPO
    add_heading(doc, "QUADRO DE TEMPO — 210 MINUTOS")
    add_table(doc, [
        ["Bloco", "Tempo", "O que acontece"],
        ["1. Retomada + censo dos endereços", "15 min", "Os sites da tarefa no quadro. O que eles têm em comum?"],
        ["2. Nome e número: o DNS", "15 min", "Analogia da agenda do celular"],
        ["3. PRÁTICA — A agenda viva", "25 min", "A turma encena o DNS: pedido, memória, nome errado, agenda fora do ar, agenda adulterada"],
        ["4. Como se lê um endereço", "10 min", "Protocolo, dono, caminho. A regra da primeira barra"],
        ["5. PRÁTICA — Clínica de links", "35 min", "Equipes analisam 16 links: dono, verdadeiro ou golpe, e a pista"],
        ["INTERVALO", "15 min", ""],
        ["6. O cadeado: o que ele resolve", "10 min", "Cartão-postal × carta lacrada. O que o cadeado NÃO garante"],
        ["7. PRÁTICA — Estações", "30 min", "A: ping por nome e por número. B: o DNS sabotado. C: o cadeado no celular"],
        ["8. PRÁTICA — Fábrica de golpes", "20 min", "Cada equipe cria links falsos pra outra desmascarar"],
        ["9. PRÁTICA — Cartilha pra família", "25 min", "\"Como não cair em link falso\" em linguagem de vó. Apresentação"],
        ["10. Fechamento + tarefa", "10 min", "Quiz de saída e anúncio da Noite 8"],
    ], col_widths=[6.5, 2, 8])
    add_para(doc, "Blocos 3, 5, 7, 8 e 9: 135 min de atividade do aluno (64% do total), mais o censo do Bloco 1. A exposição está em três blocos curtos (35 min somados), cada um seguido de prática.", italic=True, color=COR_CINZA)

    add_hr(doc)

    # BLOCO 1
    add_heading(doc, "BLOCO 1 — RETOMADA + CENSO DOS ENDEREÇOS · 15 MIN")
    add_heading(doc, "Recap relâmpago — 5 min", level=2)
    add_bullet(doc, "MAC é o quê? E IP? (o CPF e o número do apartamento)")
    add_bullet(doc, "Quem dá o IP pro celular? (o roteador — o porteiro, o DHCP)")
    add_bullet(doc, "O que quer dizer 169.254? (ninguém deu endereço)")
    add_bullet(doc, "Qual era o número da portaria? (192.168.0.1)")

    add_heading(doc, "O censo — 10 min", level=2)
    add_para(doc, "Cada aluno escreve no quadro um dos três endereços que anotou na tarefa. Quem não fez, abre um site agora no celular e copia da barra. Depois pergunte, sem explicar:")
    add_bullet(doc, "\"O que todos têm em comum?\" (começam com https://, quase todos têm .com ou .br)")
    add_bullet(doc, "\"Alguém anotou um número em vez de nome?\" (ninguém — e na Noite 6 a gente só usou número)")
    add_bullet(doc, "\"Quantos tinham cadeado?\" (quase todos)")
    add_para(doc, "\"Semana passada, pra entrar no roteador, a gente digitou um número: 192.168.0.1. Aqui no quadro ninguém digitou número. Mas o celular só sabe ir até número. Então alguém está traduzindo. Quem?\"", italic=True, color=COR_PRIMARIA)
    add_para(doc, "Guarde o quadro: no Bloco 4 esses endereços viram material de aula.", italic=True, color=COR_CINZA)

    add_hr(doc)

    # BLOCO 2
    add_heading(doc, "BLOCO 2 — NOME E NÚMERO: O DNS · 15 MIN")
    add_heading(doc, "A agenda do celular — 8 min", level=2)
    add_para(doc, "\"Quando vocês ligam pra mãe, digitam o número dela? Não. Tocam em 'Mãe'. A agenda acha o número e o celular disca. Na internet é igual: vocês digitam o nome, e um serviço chamado DNS acha o número.\"", italic=True)
    add_table(doc, [["Na agenda do celular…", "Na internet…"]] + [list(r) for r in AGENDA], col_widths=[6.5, 10])
    add_heading(doc, "Quem é a agenda — 4 min", level=2)
    add_bullet(doc, "Em casa, quem responde é o roteador — mas ele mesmo pergunta pra agenda do provedor. Por isso, quando a agenda do provedor cai, \"a internet cai\" pra todo mundo, mesmo com tudo ligado")
    add_bullet(doc, "Existem agendas públicas, gratuitas: 8.8.8.8 (Google) e 1.1.1.1 (Cloudflare). Lembram do teste do balanço? A gente pingava 8.8.8.8")
    add_bullet(doc, "Na Noite 9, quando configurar o roteador, a turma vai ver onde se troca a agenda")
    add_heading(doc, "O sintoma de campo — 3 min", level=2)
    add_para(doc, "\"O cliente diz: 'tá conectado, o WhatsApp até manda mensagem, mas nenhum site abre'. Muitas vezes é a agenda. O que já estava conectado continua, porque já sabe o número. O que precisa procurar nome novo, falha. Teste do técnico: o número responde, o nome não.\"", italic=True, color=COR_PRIMARIA)

    add_hr(doc)

    # BLOCO 3
    add_heading(doc, "BLOCO 3 — PRÁTICA · A AGENDA VIVA · 25 MIN")
    add_para(doc, "A turma encena o DNS, no mesmo jeito do porteiro vivo da Noite 6. Use o ANEXO 1.")
    add_heading(doc, "Montagem — 3 min", level=2)
    add_bullet(doc, "Um aluno é a AGENDA (o DNS). Fica numa mesa com a folha da Agenda")
    add_bullet(doc, "Cinco alunos são os SITES. Cada um prende no peito a placa com o número dele e segura a \"página\". Um sexto é o GOLPISTA (66.66.0.66) — só você e ele sabem")
    add_bullet(doc, "O resto é NAVEGADOR: cada um recebe um cartão de pedido (\"Abrir google.com.br\"…) e um papelzinho, que é a memória dele (o cache)")
    add_bullet(doc, "Turma pequena? Os sites podem ser placas em cima de cadeiras, com a página embaixo. Não precisa de gente")
    add_heading(doc, "As rodadas", level=2)
    add_para(doc, "Troque a Agenda a cada rodada. Depois de cada uma, pergunte: \"isso acontece de verdade quando?\"", italic=True, color=COR_CINZA)
    add_table(doc, [
        ["Rodada", "O que acontece", "A lição"],
        ["1. O pedido", "O navegador vai à Agenda: \"Qual o número de google.com.br?\" A Agenda olha e responde. O navegador anota no papelzinho, vai até o site com aquele número e recebe a página", "O DNS transforma o nome no número; o resto é com o número"],
        ["2. A memória", "Todos pedem o MESMO site de novo. Quem anotou no papelzinho vai direto, sem passar na Agenda", "O cache: a segunda vez é mais rápida"],
        ["3. Nome errado", "Entregue pedidos com \"gogle.com.br\" e \"youtub.com\". A Agenda responde: \"não conheço\"", "\"Não é possível acessar esse site\": às vezes é só digitação"],
        ["4. Agenda fora do ar", "A Agenda sai da mesa. Quem tem número anotado no papelzinho ainda chega. Quem pede site novo, não", "\"O WhatsApp funciona, mas os sites não abrem\""],
        ["5. Agenda adulterada", "Sem ninguém ver, risque na Agenda o número do bb.com.br e escreva 66.66.0.66. Mande um navegador abrir bb.com.br. Ele vai até o golpista e recebe a página igualzinha pedindo senha", "Digitou o nome certo e caiu no golpe. Acontece quando invadem o roteador — lembra da senha de fábrica?"],
        ["6. O selo", "Repita a rodada 5, mas agora o navegador confere o SELO antes de digitar a senha. O banco de verdade tem; o golpista não tem", "É o que o cadeado do HTTPS faz — depois do intervalo a gente vê"],
    ], col_widths=[2.6, 8.6, 5.3])
    add_para(doc, "Onde dilatar: depois da rodada 6, pergunte \"e se o golpista criasse um site com nome PARECIDO, tipo bb-com-br.app, e fizesse um selo pra esse nome?\" Deixe a turma pensar. A resposta abre o Bloco 4: o selo confere o nome, mas quem confere o nome é você.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)

    # BLOCO 4
    add_heading(doc, "BLOCO 4 — COMO SE LÊ UM ENDEREÇO · 10 MIN")
    add_para(doc, "Escreva no quadro, bem grande, e marque as três partes com cores diferentes:")
    add_para(doc, "https://  www.caixa.gov.br  /beneficios/bolsa", bold=True, size=14, color=COR_PRIMARIA, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_table(doc, [
        ["Parte", "Exemplo", "O que é"],
        ["Protocolo", "https://", "Como a conversa vai acontecer. Com S = conversa lacrada (Bloco 6)"],
        ["Nome (domínio)", "www.caixa.gov.br", "O DONO. É isso que a agenda transforma em número"],
        ["Caminho", "/beneficios/bolsa", "Qual página, dentro do site. Qualquer um escreve o que quiser aqui"],
    ], col_widths=[3.5, 4.5, 8.5])
    add_heading(doc, "A regra da primeira barra", level=2)
    add_bullet(doc, "Ache a primeira barra depois do https://. Tudo que vem depois dela é caminho — não importa")
    add_bullet(doc, "Antes dela, leia de TRÁS PRA FRENTE: .br (Brasil) → .gov (governo) → caixa. O dono é a palavra que vem logo antes do final (.com.br, .gov.br, .com, .net…)")
    add_bullet(doc, "Tudo que estiver antes do dono é enfeite: o dono do site cria o que quiser ali. Quem é dono de golpe.com cria bb.com.br.golpe.com")
    add_bullet(doc, ".gov.br só órgão público registra. .com.br qualquer pessoa registra, com CPF ou CNPJ")
    add_para(doc, "Volte ao quadro do censo e ache o dono de cada endereço que a turma escreveu. É rápido e fixa a regra.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)

    # BLOCO 5
    add_heading(doc, "BLOCO 5 — PRÁTICA · CLÍNICA DE LINKS · 35 MIN")
    add_para(doc, "No formato da clínica de cabos da Noite 5. Equipes de 3. Cada equipe recebe o jogo de 16 cartões do ANEXO 2 e preenche o laudo na página 2 da folha: quem é o dono, o veredito (verdadeiro, golpe ou não dá pra saber) e a pista.")
    add_heading(doc, "Regras", level=2)
    add_bullet(doc, "Os links ficam NO PAPEL. Ninguém digita no celular")
    add_bullet(doc, "Primeiro achem o dono (a regra da primeira barra). Só depois decidam")
    add_bullet(doc, "Separem os cartões em três montes em cima da mesa: VERDADEIRO, GOLPE, NÃO DÁ PRA SABER")
    add_bullet(doc, "~20 min pra equipe, depois correção")
    add_heading(doc, "Gabarito (só com você)", level=2)
    add_table(doc, [["#", "Dono", "Veredito", "Por quê"]] +
              [[str(i), dono, ver, expl] for i, (_, dono, ver, expl) in enumerate(LINKS, 1)],
              col_widths=[0.8, 4.2, 3.3, 8.2])
    add_heading(doc, "Correção — 15 min", level=2)
    add_para(doc, "Link por link, uma equipe diferente diz o veredito e a pista. Pare em quatro:")
    add_bullet(doc, "#2 e #7: o nome do banco ou da loja no COMEÇO. É o truque mais comum que existe")
    add_bullet(doc, "#13 (wa.me): quem condenou? Estranho não é golpe — tem que achar o dono")
    add_bullet(doc, "#15 (o @): quase ninguém acerta. Diga que é raro, mas existe")
    add_bullet(doc, "#16 (Drive): o dono é bom, o arquivo pode não ser")
    add_para(doc, "Onde dilatar: peça que cada aluno abra o WhatsApp e procure, nas conversas e nos grupos da família, um link que recebeu. A equipe aplica a regra e diz o dono — sem abrir. Quase sempre aparece um golpe de verdade, e a aula fica real.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)
    add_heading(doc, "INTERVALO · 15 MIN", color=COR_CINZA)
    add_hr(doc)

    # BLOCO 6
    add_heading(doc, "BLOCO 6 — O CADEADO: O QUE ELE RESOLVE · 10 MIN")
    add_heading(doc, "Cartão-postal × carta lacrada — 5 min", level=2)
    add_bullet(doc, "http (sem S) é cartão-postal: todo mundo no caminho pode ler — o vizinho no mesmo Wi-Fi, o Wi-Fi da rodoviária. O navegador avisa: \"Não seguro\"")
    add_bullet(doc, "https (com S) é carta lacrada: só o site lê. E o site tem que mostrar um documento (o certificado) provando que é dono daquele nome. É o SELO da rodada 6")
    add_heading(doc, "O que o cadeado NÃO garante — 5 min", level=2)
    add_para(doc, "\"O golpista também tem cadeado. Ele compra o nome dele — atualizacao-cadastral.com — e tira o selo DESSE nome, de graça. A conversa com ele vai lacradinha… direto pro golpista. O cadeado diz que ninguém está ouvindo. Não diz se quem está do outro lado é do bem.\"", italic=True, color=COR_ALERTA)
    add_para(doc, "Por isso as duas coisas juntas: cadeado E dono certo. Uma sem a outra não vale.", bold=True, color=COR_PRIMARIA)

    add_hr(doc)

    # BLOCO 7
    add_heading(doc, "BLOCO 7 — PRÁTICA · ESTAÇÕES · 30 MIN")
    add_para(doc, "Três estações, ~10 min cada, as equipes rodam. Página 3 da folha.")
    add_heading(doc, "Estação A — O número e o nome (computador normal)", level=2)
    add_bullet(doc, "ping 8.8.8.8 → responde (o número)")
    add_bullet(doc, "ping google.com → a primeira linha mostra o número que a agenda achou. Anote")
    add_bullet(doc, "nslookup google.com → \"Servidor\" é quem respondeu (o roteador ou o celular do professor); \"Endereços\" é o número do Google")
    add_bullet(doc, "nslookup de um site da tarefa de cada um")
    add_heading(doc, "Estação B — O computador doente (DNS sabotado)", level=2)
    add_bullet(doc, "Abra o navegador: nenhum site abre")
    add_bullet(doc, "ping 8.8.8.8 → responde! O caminho até a internet está bom")
    add_bullet(doc, "ping google.com → \"não pôde encontrar o host\"")
    add_bullet(doc, "A equipe escreve o diagnóstico: o número responde, o nome não → é a agenda (DNS)")
    add_bullet(doc, "Pra equipe que acertou: mostre onde estava a sabotagem (Propriedades do IPv4). Uma equipe por vez pode consertar (\"Obter o endereço dos servidores DNS automaticamente\") — e você sabota de novo pra próxima")
    add_heading(doc, "Estação C — O cadeado no celular (4G)", level=2)
    add_bullet(doc, "Abra um site da sua tarefa. Toque no cadeado (ou no ícone ao lado do endereço) → \"A conexão é segura\" → \"Certificado é válido\". Anote \"Emitido para\" — bate com o dono?")
    add_bullet(doc, "Abra http://neverssl.com — um site feito de propósito sem cadeado. O que o navegador mostra? (\"Não seguro\")")
    add_bullet(doc, "No iPhone, o Safari mostra menos detalhes: vale só ver o cadeado e o \"Não seguro\"")
    add_para(doc, "\"O que vocês viram na Estação B é um dos chamados mais comuns: 'tem internet mas não abre nada'. Sem saber disso, o técnico troca cabo, reinicia roteador, liga pro provedor. Com dois pings, em um minuto ele sabe que é a agenda.\"", italic=True, color=COR_PRIMARIA)

    add_hr(doc)

    # BLOCO 8
    add_heading(doc, "BLOCO 8 — PRÁTICA · FÁBRICA DE GOLPES · 20 MIN")
    add_para(doc, "Agora as equipes trocam de lado: cada uma fabrica links pra outra desmascarar. Quem faz golpe entende como ele é feito.")
    add_bullet(doc, "Sorteie pra cada equipe um site de verdade: o banco da cidade, uma loja, o gov.br, os Correios, a Celesc, o Instagram…")
    add_bullet(doc, "Em 8 min, a equipe escreve na folha 4 links: 3 golpes, usando truques diferentes (nome no começo, traço no lugar do ponto, letra trocada, final estranho, o @), e 1 verdadeiro. Em ordem misturada")
    add_bullet(doc, "Troca a folha com a equipe ao lado. A outra equipe tem 5 min pra achar o dono de cada link e apontar o verdadeiro")
    add_bullet(doc, "Desvira e confere. Ponto pra quem desmascarou — e ponto pra quem conseguiu enganar, se o golpe era bem feito")
    add_para(doc, "Onde dilatar: cada equipe escreve também a MENSAGEM que acompanharia o link (\"Seu CPF foi bloqueado, regularize em 24h…\"). A turma lista no quadro os truques da mensagem: pressa, medo, prêmio, autoridade.", italic=True, color=COR_DESTAQUE)

    add_hr(doc)

    # BLOCO 9
    add_heading(doc, "BLOCO 9 — PRÁTICA · CARTILHA PRA FAMÍLIA · 25 MIN")
    add_para(doc, "Cada equipe faz, numa folha A4 ou meia cartolina, uma cartilha: \"Como não cair em link falso\". O público é a avó, o tio, o vizinho — gente que não fez o curso.")
    add_heading(doc, "O que a cartilha precisa ter — 15 min pra fazer", level=2)
    add_bullet(doc, "No máximo 3 passos, em frases curtas, sem palavra técnica (nada de DNS, domínio, protocolo)")
    add_bullet(doc, "Um exemplo de link verdadeiro e um de golpe, com o dono marcado")
    add_bullet(doc, "Uma frase sobre o cadeado (\"o cadeado sozinho não basta\")")
    add_bullet(doc, "O que fazer na dúvida (\"não clica; abre o app do banco ou liga pro número do cartão\")")
    add_heading(doc, "Apresentação — ~2 min por equipe", level=2)
    add_bullet(doc, "Cada equipe lê a cartilha como se estivesse explicando pra avó")
    add_bullet(doc, "A turma diz uma palavra que a avó não entenderia. A equipe troca na hora")
    add_para(doc, "No fim, cada um tira foto das cartilhas. Sugira mandar no grupo da família — é o primeiro \"atendimento\" deles fora da sala.", italic=True, color=COR_PRIMARIA)

    add_hr(doc)

    # BLOCO 10
    add_heading(doc, "BLOCO 10 — FECHAMENTO E TAREFA · 10 MIN")
    add_heading(doc, "Quiz de saída", level=2)
    add_bullet(doc, "O que o DNS faz, numa frase?")
    add_bullet(doc, "Quem é o dono de https://caixa.gov.br.consulta-fgts.com/app?")
    add_bullet(doc, "O site tem cadeado. Isso quer dizer que é seguro digitar a senha?")
    add_bullet(doc, "O cliente diz que o WhatsApp funciona e nenhum site abre. Quais os dois pings que você faz?")

    add_heading(doc, "A tarefa", level=2)
    add_para(doc, "\"Na próxima noite a gente começa o projeto de verdade: a rede de uma casa. A tarefa: desenhem a planta da casa de vocês, à mão, sem régua — os cômodos vistos de cima. Marquem com um X onde fica o modem ou roteador, com uma seta por onde o cabo da internet chega, e pintem os cômodos onde o Wi-Fi é fraco.\"", italic=True)

    add_heading(doc, "Anuncie a próxima", level=2)
    add_para(doc, "\"Na próxima noite: planejando a rede de uma casa. Onde colocar o roteador, por onde passar cabo, quanto material comprar. É a primeira parte do projeto final.\"", italic=True)

    add_hr(doc)

    # SE PERGUNTAREM
    add_heading(doc, "SE PERGUNTAREM")
    add_para(doc, "\"Se eu clicar num link de golpe, já era?\"", bold=True)
    add_para(doc, "> \"Na maioria das vezes, só abrir não faz o estrago — o estrago vem de digitar senha, código, cartão, ou baixar e instalar alguma coisa. Clicou sem querer? Fecha, não digita nada. Digitou? Troca a senha na hora e liga pro banco.\"", italic=True)
    add_para(doc, "\"E o QR Code do Pix ou do boleto?\"", bold=True)
    add_para(doc, "> \"QR Code é um link escondido. A câmera do celular mostra o endereço antes de abrir: leia o dono. No Pix, confira o nome de quem vai receber antes de confirmar.\"", italic=True)
    add_para(doc, "\"Por que no 4G às vezes abre e no Wi-Fi de casa não?\"", bold=True)
    add_para(doc, "> \"No 4G a agenda é a da operadora; no Wi-Fi, é a do roteador e do provedor. Se só o Wi-Fi falha, suspeite da agenda do roteador — o teste dos dois pings resolve.\"", italic=True)
    add_para(doc, "\"Navegação anônima protege?\"", bold=True)
    add_para(doc, "> \"Ela só não guarda histórico no SEU celular. O site, o provedor e o golpista veem tudo igual. Não protege de golpe.\"", italic=True)
    add_para(doc, "\"Quem decide quem é dono de um nome .br?\"", bold=True)
    add_para(doc, "> \"O Registro.br. Qualquer pessoa pode consultar lá quem registrou um .com.br — é uma das ferramentas da caixa.\"", italic=True)

    add_hr(doc)

    # DEPOIS DA AULA
    add_heading(doc, "DEPOIS DA AULA")
    add_bullet(doc, "DESFAÇA a sabotagem da Estação B: Propriedades do IPv4 → \"Obter o endereço dos servidores DNS automaticamente\". Se usou o celular, volte o DNS privado pra \"Automático\"")
    add_bullet(doc, "Guarde os cartões de links e as cartilhas. As cartilhas voltam na Noite 11 (segurança)")
    add_bullet(doc, "Anote quantas equipes acertaram o #2 e o #15 sem ajuda. É o indicador da noite")

    add_hr(doc)

    # DESCRITIVO
    add_heading(doc, "DESCRITIVO PARA O DIÁRIO")
    add_para(doc, "Desenvolvi atividade sobre resolução de nomes e navegação segura. Apresentei o serviço DNS por analogia com a agenda telefônica, relacionando nome de domínio e endereço IP, cache e servidores DNS públicos, e a estrutura de uma URL (protocolo, domínio e caminho), com ênfase na identificação do domínio responsável pelo endereço. Os estudantes encenaram o funcionamento do DNS, incluindo falha do servidor e adulteração de registro, e analisaram em equipe um conjunto de endereços verdadeiros e fraudulentos, registrando laudo. Em estações práticas, compararam testes de conectividade por endereço IP e por nome (ping e nslookup), diagnosticaram uma falha de DNS preparada e verificaram o certificado HTTPS no navegador do celular, discutindo o que a conexão segura garante e o que não garante. Ao final, criaram endereços fraudulentos para análise mútua e produziram uma cartilha de prevenção a golpes destinada às famílias.")

    add_hr(doc)

    # CAIXA DE FERRAMENTAS DIDÁTICAS
    add_heading(doc, "CAIXA DE FERRAMENTAS DIDÁTICAS", color=COR_DESTAQUE)
    add_para(doc, "Se sobrar tempo, ou pra usar em outra noite:", italic=True, color=COR_CINZA)
    add_para(doc, "1. Quem é o dono? (Registro.br)", bold=True)
    add_para(doc, "No celular, entrem em registro.br e consultem um .com.br de uma loja da cidade: aparece quem registrou e desde quando. Depois tentem um nome que não existe: \"disponível\" — qualquer um pode comprar. Mostra como o golpista consegue o nome dele. 15 min.")
    add_para(doc, "2. A memória do computador", bold=True)
    add_para(doc, "Na Estação A: ipconfig /displaydns mostra todos os nomes que o computador já guardou no cache (o papelzinho da rodada 2). Depois ipconfig /flushdns apaga. Liga com o \"às vezes resolve reiniciando\". 10 min.")
    add_para(doc, "3. Leia o QR Code antes de abrir", bold=True)
    add_para(doc, "Gere no seu celular (ou imprima) 3 QR Codes: um pro site da Prefeitura, um pra wa.me e um pra um endereço com cara de golpe que você inventar. A turma aponta a câmera e lê o endereço SEM abrir. Diz o dono de cada um. 10 min.")
    add_para(doc, "4. O link encurtado", bold=True)
    add_para(doc, "Mostre que um link bit.ly pode ser visto antes de abrir: em muitos encurtadores, acrescentar um + no final mostra o destino. A turma discute: por que golpista adora encurtador? 10 min.")
    add_para(doc, "5. Caça ao golpe real", bold=True)
    add_para(doc, "Cada um procura, no próprio celular (SMS, e-mail, WhatsApp), uma mensagem de golpe que já recebeu. Em equipe, montam a \"ficha\" do golpe: o link, o dono, o truque do link e o truque da mensagem. 20 min.")

    # ANEXO 1 — agenda viva
    doc.add_page_break()
    add_heading(doc, "ANEXO 1 — AGENDA VIVA · A AGENDA (imprimir 1)")
    add_para(doc, "Números inventados pra brincadeira. Na rodada 5 você risca o número do bb.com.br e escreve 66.66.0.66.", italic=True, size=10, color=COR_CINZA)
    tbl = add_table(doc, [["Nome", "Número (IP)"]] + [[n, ip] for n, ip, _ in SITES_VIVOS],
                    col_widths=[9, 7.5])
    for row in tbl.rows:
        row.height = Cm(1.3)
        for cell in row.cells:
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(18)

    doc.add_page_break()
    add_heading(doc, "ANEXO 1 — PLACAS DOS SITES E PÁGINAS (imprimir e recortar)")
    add_para(doc, "Cada site prende a placa no peito e segura a página. A do golpista fica escondida até a rodada 5.", italic=True, size=10, color=COR_CINZA)
    tbl = add_table(doc, [["Placa (no peito)", "Página (na mão)"]] +
                    [[f"{ip}\n{n}", pag] for n, ip, pag in SITES_VIVOS + [GOLPISTA]],
                    col_widths=[7, 9.5])
    for row in tbl.rows[1:]:
        row.height = Cm(2.6)
        for run in row.cells[0].paragraphs[0].runs:
            run.font.size = Pt(18)
            run.bold = True

    add_heading(doc, "Cartões de pedido (recortar)", level=2)
    pedidos = ["google.com.br", "youtube.com", "bb.com.br", "itapiranga.sc.gov.br", "mercadolivre.com.br",
               "google.com.br", "youtube.com", "bb.com.br", "mercadolivre.com.br", "itapiranga.sc.gov.br",
               "gogle.com.br", "youtub.com"]
    tbl = add_table(doc, [[f"Abrir\n{pedidos[i]}" for i in range(r, r + 4)] for r in range(0, 12, 4)],
                    header=False, col_widths=[4.1] * 4)
    for row in tbl.rows:
        row.height = Cm(1.8)

    # ANEXO 2 — cartões de links
    doc.add_page_break()
    add_heading(doc, "ANEXO 2 — CLÍNICA DE LINKS (1 jogo por equipe, recortar)")
    add_para(doc, "Só no papel. Ninguém digita estes links.", bold=True, size=10, color=COR_ALERTA)
    tbl = add_table(doc, [[f"#{i + 1}\n{LINKS[i][0]}", f"#{i + 2}\n{LINKS[i + 1][0]}"] for i in range(0, 16, 2)],
                    header=False, col_widths=[8.3, 8.3])
    for row in tbl.rows:
        row.height = Cm(2.3)
        for cell in row.cells:
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(11)
                run.font.name = "Consolas"

    path = os.path.join(OUTDIR, "Roteiro - Redes 7 - Nomes na rede.docx")
    doc.save(path)
    return path


# ============================================================
# SLIDES
# ============================================================
def build_slides():
    prs = new_presentation()
    P, D, K, R = PPalette.P, PPalette.D, PPalette.K, PPalette.R

    def lista(kicker, titulo, linhas, subtitulo=None, size=20, nota=None, nota_cor=None):
        s = add_slide(prs)
        add_title_block(s, kicker, titulo, subtitulo)
        add_text(s, linhas, 0.7, 2.8, 11.9, 3.8, size=size, color=K)
        if nota:
            add_text(s, nota, 0.7, 6.4, 11.9, 0.5, size=16, italic=True, color=nota_cor or D, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    def frase(kicker, texto, cor=None, size=36):
        s = add_slide(prs)
        add_text(s, kicker, 0.7, 1.6, 11.9, 0.5, size=16, bold=True, color=D, align=PP_ALIGN.CENTER)
        add_text(s, texto.split("\n"), 0.9, 2.4, 11.5, 3.0, size=size, bold=True, color=cor or P, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    def tabela(kicker, titulo, rows, col_widths=None, subtitulo=None, size=16, nota=None, height=3.6):
        s = add_slide(prs)
        add_title_block(s, kicker, titulo, subtitulo)
        add_table_slide(s, rows, top=2.6, height=height, col_widths=col_widths, size=size)
        if nota:
            add_text(s, nota, 0.7, 6.5, 11.9, 0.5, size=15, italic=True, color=D, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    def tela(kicker, titulo, linhas_tela, nota=None, nota_cor=None):
        s = add_slide(prs)
        add_title_block(s, kicker, titulo)
        fundo = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, PInches(0.7), PInches(2.5), PInches(11.9), PInches(3.0))
        fundo.fill.solid()
        fundo.fill.fore_color.rgb = K
        fundo.line.fill.background()
        tb = add_text(s, linhas_tela, 1.0, 2.75, 11.3, 2.6, size=20, color=PPalette.W)
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = "Consolas"
        if nota:
            add_text(s, nota, 0.7, 5.8, 11.9, 0.9, size=18, italic=True, color=nota_cor or D, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    def link(kicker, url, partes, nota=None, nota_cor=None):
        """Slide de endereço grande + explicação das partes."""
        s = add_slide(prs)
        add_text(s, kicker, 0.7, 0.6, 11.9, 0.5, size=16, bold=True, color=D)
        fundo = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, PInches(0.7), PInches(1.3), PInches(11.9), PInches(1.2))
        fundo.fill.solid()
        fundo.fill.fore_color.rgb = PPalette.W
        fundo.line.color.rgb = P
        tb = add_text(s, url, 0.9, 1.55, 11.5, 0.8, size=26, bold=True, color=K, align=PP_ALIGN.CENTER)
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = "Consolas"
        add_text(s, partes, 0.7, 3.0, 11.9, 3.0, size=22, color=K)
        if nota:
            add_text(s, nota, 0.7, 6.3, 11.9, 0.6, size=18, italic=True, bold=True, color=nota_cor or D, align=PP_ALIGN.CENTER)
        add_footer(s, FOOTER)
        return s

    # Capa
    add_slide_cover(prs, "NOITE 7", "NOMES", "NA REDE",
                    "DNS, endereço de site e o que o cadeado resolve",
                    "Redes de Computadores · Curso FIC · CEJA Itapiranga – SC")

    lista("RETOMADA", "Da Noite 6 pra hoje", [
        "•  MAC é o quê? E IP?",
        "•  Quem dá o IP pro celular?",
        "•  O que quer dizer 169.254?",
        "•  Qual era o número da portaria?",
    ])

    lista("O CENSO", "Escreva no quadro um site da sua tarefa", [
        "Agora olhe o quadro:",
        "",
        "•  O que todos têm em comum?",
        "•  Alguém escreveu um NÚMERO em vez de nome?",
        "•  Quantos tinham cadeado?",
    ])

    frase("A PERGUNTA DA NOITE", "A gente digita nome.\nO celular só sabe ir até número.\nQuem traduz?")

    frase("A RESPOSTA", "Você não digita o número da sua mãe.\nVocê toca em \"Mãe\".\nA agenda acha o número.", size=32)

    tabela("O DNS É A AGENDA DA INTERNET", "Quem é quem", [["Na agenda do celular…", "Na internet…"]] +
           [list(r) for r in AGENDA], col_widths=[4.9, 7.0], size=15, height=4.0)

    lista("QUEM É A AGENDA", "Onde mora o DNS", [
        "•  Em casa, o roteador responde — e pergunta pra agenda do provedor",
        "•  Agenda do provedor caiu? \"A internet caiu\" pra todo mundo",
        "•  Agendas públicas e gratuitas: 8.8.8.8 (Google) e 1.1.1.1",
        "•  Na Noite 9 a gente vê onde se troca a agenda do roteador",
    ])

    frase("O SINTOMA DE CAMPO", "\"Tá conectado, o WhatsApp manda mensagem,\nmas nenhum site abre.\"\n\nO número responde. O nome não.", size=30)

    lista("PRÁTICA · 25 MIN", "A agenda viva", [
        "•  Um é a AGENDA, com a folha de nomes e números",
        "•  Cinco são SITES: placa com o número no peito, página na mão",
        "•  O resto é NAVEGADOR: um pedido e um papelzinho (a memória)",
        "",
        "Navegador:  \"Qual o número de google.com.br?\"",
    ], size=20, nota="Seis rodadas. Cada uma é um problema de verdade.")

    tabela("AS RODADAS", "O que cada uma mostra", [
        ["Rodada", "Na vida real"],
        ["1. O pedido", "O DNS transforma o nome no número"],
        ["2. A memória", "O cache: a segunda vez é mais rápida"],
        ["3. Nome errado", "\"Não é possível acessar esse site\""],
        ["4. Agenda fora do ar", "O WhatsApp funciona, os sites não abrem"],
        ["5. Agenda adulterada", "Nome certo, lugar errado: golpe"],
        ["6. O selo", "O que o cadeado faz (depois do intervalo)"],
    ], col_widths=[3.9, 8.0], size=16, height=3.8)

    link("COMO SE LÊ UM ENDEREÇO", "https://  www.caixa.gov.br  /beneficios/bolsa", [
        "https://  →  como a conversa acontece (com S = lacrada)",
        "www.caixa.gov.br  →  o DONO. É isso que a agenda traduz",
        "/beneficios/bolsa  →  a página lá dentro. Qualquer um escreve o que quiser",
    ])

    frase("A REGRA DA PRIMEIRA BARRA", "Ache a primeira barra.\nAntes dela, leia de trás pra frente.\nO dono é o que vem antes do final.", size=32)

    link("ONDE O GOLPE SE ESCONDE", "https://bb.com.br.atualizacao-cadastral.com/login", [
        "1.  Primeira barra: depois de .com",
        "2.  De trás pra frente: .com → atualizacao-cadastral",
        "3.  O dono: atualizacao-cadastral.com",
        "",
        "O \"bb.com.br\" no começo é enfeite. Quem é dono escreve o que quiser ali.",
    ], nota="Nome do banco no COMEÇO do link: o truque mais comum que existe.", nota_cor=R)

    lista(".gov.br × .com.br", "Quem pode ter cada final", [
        "•  .gov.br — só órgão público registra",
        "•  .com.br — qualquer pessoa, com CPF ou CNPJ",
        "•  .com, .net, .app, .shop, .online, .info — qualquer pessoa no mundo",
        "",
        "Volte ao quadro do censo: quem é o dono de cada site da turma?",
    ])

    lista("PRÁTICA · 35 MIN", "Clínica de links", [
        "•  Equipes de 3. Dezesseis cartões",
        "•  Primeiro achem o DONO. Depois decidam",
        "•  Três montes: VERDADEIRO · GOLPE · NÃO DÁ PRA SABER",
        "•  Laudo na página 2 da folha",
    ], size=22, nota="Os links ficam NO PAPEL. Ninguém digita no celular.", nota_cor=R)

    frase("", "INTERVALO\n15 minutos", cor=K)

    tabela("O CADEADO", "Cartão-postal × carta lacrada", [
        ["", "http (sem S)", "https (com S)"],
        ["Quem lê no caminho", "Qualquer um: vizinho, Wi-Fi público", "Ninguém"],
        ["O site prova quem é?", "Não", "Sim: mostra o certificado (o selo)"],
        ["O navegador mostra", "\"Não seguro\"", "O cadeado"],
    ], col_widths=[3.6, 4.2, 4.1], size=17)

    frase("CUIDADO", "O golpista também tem cadeado.\nO cadeado diz que ninguém está ouvindo.\nNão diz se quem está do outro lado é do bem.", cor=R, size=30)

    frase("AS DUAS COISAS JUNTAS", "Cadeado  +  dono certo.\nUma sem a outra não vale.")

    lista("PRÁTICA · 30 MIN", "Três estações", [
        "A.  O número e o nome — ping e nslookup",
        "B.  O computador doente — o que está errado?",
        "C.  O cadeado no celular — quem emitiu? E o \"Não seguro\"?",
        "",
        "~10 minutos em cada. Anote na página 3.",
    ], size=22)

    for titulo, linhas_tela in TELAS:
        tela("ESTAÇÕES A E B", titulo, linhas_tela)

    frase("O DIAGNÓSTICO EM 1 MINUTO", "ping 8.8.8.8 responde\nping google.com não\n=\na agenda (DNS)", size=32)

    lista("ESTAÇÃO C", "O cadeado no celular", [
        "1.  Abra um site da sua tarefa",
        "2.  Toque no cadeado → \"A conexão é segura\" → \"Certificado é válido\"",
        "3.  \"Emitido para\" bate com o dono?",
        "4.  Abra http://neverssl.com — o que aparece?",
    ], size=21)

    lista("PRÁTICA · 20 MIN", "Fábrica de golpes", [
        "•  Sua equipe recebe um site de verdade",
        "•  Escreva 4 links: 3 golpes com truques diferentes + 1 verdadeiro",
        "•  Troque com a equipe ao lado. Ela tem 5 min pra achar o verdadeiro",
        "",
        "Truques: nome no começo · traço no lugar do ponto · letra trocada · final estranho · o @",
    ], size=19)

    lista("PRÁTICA · 25 MIN", "Cartilha pra família", [
        "\"Como não cair em link falso\" — pra avó entender",
        "",
        "•  No máximo 3 passos, sem palavra técnica",
        "•  Um link verdadeiro e um golpe, com o dono marcado",
        "•  Uma frase sobre o cadeado",
        "•  O que fazer na dúvida",
    ], size=20, nota="Depois: cada equipe lê como se estivesse explicando pra avó.")

    lista("O QUE FICA", "Três frases pra guardar", [
        "\"O DNS é a agenda da internet.\"",
        "",
        "\"O dono do link vem antes da primeira barra, lido de trás pra frente.\"",
        "",
        "\"Cadeado mais dono certo. Um sem o outro não vale.\"",
    ], size=23)

    lista("PARA A PRÓXIMA NOITE", "Planejando a rede de uma casa", [
        "A tarefa:",
        "",
        "•  Desenhe a planta da sua casa à mão, vista de cima",
        "•  X onde fica o modem ou roteador",
        "•  Uma seta por onde o cabo da internet chega",
        "•  Pinte os cômodos onde o Wi-Fi é fraco",
    ])

    path = os.path.join(OUTDIR, "Aula - Redes 7 - Nomes na rede.pptx")
    prs.save(path)
    return path


# ============================================================
# FOLHA DO ALUNO
# ============================================================
def linhas(doc, n):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("_" * 100)
        r.font.size = Pt(11)
        r.font.color.rgb = COR_CINZA


def pergunta(doc, texto, size=11):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.font.size = Pt(size)
    r.bold = True


def build_folha():
    doc = Document()
    setup_page(doc, margem=Cm(1.8))

    # PÁGINA 1 — Agenda + como se lê + tarefa
    cabecalho_documento(
        doc,
        "NOME, NÚMERO E DONO",
        f"Redes de Computadores · {NOITE} · Curso FIC · CEJA Itapiranga · 1 de 3",
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Nome: ")
    r.font.size = Pt(11)
    r.bold = True
    r = p.add_run("_________________________________________________     Equipe: ________")
    r.font.size = Pt(11)

    add_heading(doc, "1  O DNS é a agenda da internet", level=2)
    add_para(doc, "Complete a coluna da direita durante a explicação.", italic=True, size=10, color=COR_CINZA)
    tbl = add_table(doc, [["Na agenda do celular…", "Na internet…"]] + [[a, ""] for a, _ in AGENDA],
                    col_widths=[6.5, 11])
    for row in tbl.rows[1:]:
        row.height = Cm(0.8)

    add_heading(doc, "2  As três partes de um endereço", level=2)
    add_para(doc, "https://  www.caixa.gov.br  /beneficios/bolsa", bold=True, size=13, color=COR_PRIMARIA, align=WD_ALIGN_PARAGRAPH.CENTER)
    tbl = add_table(doc, [
        ["Parte", "Neste exemplo", "Pra que serve"],
        ["Protocolo", "", ""],
        ["Nome (o dono)", "", ""],
        ["Caminho", "", ""],
    ], col_widths=[3.5, 5.5, 8.5])
    for row in tbl.rows[1:]:
        row.height = Cm(0.8)

    add_para(doc, "A regra: ache a primeira barra depois do https://. Antes dela, leia de trás pra frente. O dono é a palavra que vem logo antes do final (.com.br, .gov.br, .com…).", italic=True, size=10, color=COR_PRIMARIA)

    add_heading(doc, "3  Os sites da sua tarefa", level=2)
    tbl = add_table(doc, [
        ["Endereço que você anotou", "Quem é o dono?", "Cadeado?"],
        ["", "", "☐ Sim  ☐ Não"],
        ["", "", "☐ Sim  ☐ Não"],
        ["", "", "☐ Sim  ☐ Não"],
    ], col_widths=[9, 5, 3.5])
    for row in tbl.rows[1:]:
        row.height = Cm(0.9)

    pergunta(doc, "Na agenda viva, como o navegador caiu no golpista mesmo digitando o nome certo?")
    linhas(doc, 2)

    # PÁGINA 2 — Clínica de links
    doc.add_page_break()
    cabecalho_documento(
        doc,
        "CLÍNICA DE LINKS · LAUDO DA EQUIPE",
        f"Verdadeiro ou golpe · {NOITE} · 2 de 3",
    )
    add_para(doc, "Primeiro achem o DONO (regra da primeira barra). Depois decidam. Os links ficam no papel: ninguém digita no celular.", italic=True, size=10, color=COR_CINZA)
    p = doc.add_paragraph()
    r = p.add_run("Equipe: ______________________________________________________________")
    r.font.size = Pt(11)

    tbl = add_table(doc, [["#", "Quem é o dono?", "Veredito", "A pista"]] +
                    [[str(i), "", "☐ V   ☐ G   ☐ ?", ""] for i in range(1, len(LINKS) + 1)],
                    col_widths=[0.8, 5.2, 3.6, 7.9])
    for row in tbl.rows[1:]:
        row.height = Cm(0.95)
    add_para(doc, "V = verdadeiro   ·   G = golpe   ·   ? = não dá pra saber", italic=True, size=9, color=COR_CINZA)

    p = doc.add_paragraph()
    r = p.add_run("Quantos a equipe acertou?  ______ de 16        Qual enganou vocês? ______        Por quê?")
    r.font.size = Pt(11)
    linhas(doc, 1)

    # PÁGINA 3 — Estações + fábrica + tarefa
    doc.add_page_break()
    cabecalho_documento(
        doc,
        "ESTAÇÕES E FÁBRICA DE GOLPES",
        f"Diagnóstico · {NOITE} · 3 de 3",
    )

    add_heading(doc, "1  Estação A — O número e o nome", level=2)
    tbl = add_table(doc, [
        ["Comando", "Respondeu?", "O número que apareceu"],
        ["ping 8.8.8.8", "☐ Sim  ☐ Não", ""],
        ["ping google.com", "☐ Sim  ☐ Não", ""],
        ["nslookup ____________________", "☐ Sim  ☐ Não", ""],
    ], col_widths=[6.5, 3.5, 7.5])
    for row in tbl.rows[1:]:
        row.height = Cm(0.8)
    p = doc.add_paragraph()
    r = p.add_run("No nslookup, quem foi o \"Servidor\" (a agenda) que respondeu? ________________________")
    r.font.size = Pt(11)

    add_heading(doc, "2  Estação B — O computador doente", level=2)
    p = doc.add_paragraph()
    r = p.add_run("ping 8.8.8.8:  ☐ respondeu  ☐ não        ping google.com:  ☐ respondeu  ☐ não")
    r.font.size = Pt(11)
    pergunta(doc, "Diagnóstico da equipe: o que está doente nesse computador? Como vocês sabem?")
    linhas(doc, 2)

    add_heading(doc, "3  Estação C — O cadeado no celular", level=2)
    p = doc.add_paragraph()
    r = p.add_run("Site: ______________________   Certificado emitido para: ______________________   Bate com o dono?  ☐ Sim  ☐ Não")
    r.font.size = Pt(10)
    p = doc.add_paragraph()
    r = p.add_run("O que o navegador mostrou no http://neverssl.com? ____________________________________________")
    r.font.size = Pt(10)

    add_heading(doc, "4  Fábrica de golpes", level=2)
    p = doc.add_paragraph()
    r = p.add_run("Site que a equipe recebeu: ______________________   Escreva 3 golpes e 1 verdadeiro, misturados:")
    r.font.size = Pt(11)
    for i in "ABCD":
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{i})  https://__________________________________________________________________")
        r.font.size = Pt(11)
    p = doc.add_paragraph()
    r = p.add_run("A outra equipe achou o verdadeiro?  ☐ Sim  ☐ Não        Qual golpe enganou? ______")
    r.font.size = Pt(11)

    add_heading(doc, "TAREFA DA SEMANA", level=2, color=COR_DESTAQUE)
    add_bullet(doc, "Numa folha à parte, desenhe a planta da sua casa à mão, vista de cima, com os cômodos")
    add_bullet(doc, "Marque com X onde fica o modem/roteador e com uma seta por onde chega o cabo da internet")
    add_bullet(doc, "Pinte os cômodos onde o Wi-Fi é fraco. Traga na Noite 8")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Traga esta folha na próxima aula · CEJA Itapiranga – SC")
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = COR_CINZA

    path = os.path.join(OUTDIR, "Folha do Aluno - Redes 7 - Nomes na rede.docx")
    doc.save(path)
    return path


# ============================================================
# RESUMO DO ALUNO
# ============================================================
def build_resumo():
    doc = Document()
    setup_page(doc, margem=Cm(1.8))

    cabecalho_documento(
        doc,
        "NOMES NA REDE: DNS, ENDEREÇO DE SITE E CADEADO",
        f"Resumo do aluno · {NOITE} · Redes de Computadores · Curso FIC · CEJA Itapiranga",
    )

    add_heading(doc, "A ideia principal", level=2)
    add_para(doc, "A gente digita nomes; os aparelhos só sabem ir até números (IP). Quem traduz é o DNS, a agenda da internet. E pra não cair em golpe, o que importa é saber quem é o DONO do endereço — o cadeado sozinho não basta.")

    add_heading(doc, "O DNS é a agenda da internet", level=2)
    add_table(doc, [["Na agenda do celular…", "Na internet…"]] + [list(r) for r in AGENDA], col_widths=[6, 9.5])

    add_heading(doc, "Como achar o dono de um link", level=2)
    add_bullet(doc, "Ache a primeira barra depois do https://. O que vem depois dela é só o caminho")
    add_bullet(doc, "Antes da barra, leia de trás pra frente. O dono é a palavra logo antes do final (.com.br, .gov.br, .com…)")
    add_bullet(doc, "Tudo antes do dono é enfeite: bb.com.br.atualizacao-cadastral.com é de atualizacao-cadastral.com")
    add_bullet(doc, "Pegadinha do @: tudo antes do @ é ignorado. O dono vem depois")
    add_bullet(doc, ".gov.br só órgão público registra. .com.br, .com, .app, .shop… qualquer pessoa")
    add_bullet(doc, "Link encurtado (bit.ly e parecidos) esconde o dono. Na dúvida, não abra")

    add_heading(doc, "Os truques mais comuns", level=2)
    add_table(doc, [
        ["Truque", "Exemplo", "O dono de verdade"],
        ["Nome no começo", "correios.taxa-liberacao.info", "taxa-liberacao.info"],
        ["Traço no lugar do ponto", "caixa-gov-br.app", "caixa-gov-br.app"],
        ["Letra trocada", "arnazon.com.br (rn parece m)", "arnazon.com.br"],
        ["O @", "nubank.com.br@login-seguro.net", "login-seguro.net"],
    ], col_widths=[4.3, 6.2, 5])

    add_heading(doc, "O cadeado (https)", level=2)
    add_bullet(doc, "http (sem S) é cartão-postal: quem está no caminho pode ler. O navegador mostra \"Não seguro\"")
    add_bullet(doc, "https (com S) é carta lacrada, e o site mostra um certificado provando que é dono daquele nome")
    add_bullet(doc, "O golpista TAMBÉM tem cadeado — pro nome dele. Cadeado mais dono certo: um sem o outro não vale")
    add_bullet(doc, "Na dúvida: não clique. Abra o app do banco ou ligue pro número que está no seu cartão")

    add_heading(doc, "No atendimento: \"tem internet, mas nenhum site abre\"", level=2)
    add_para(doc, "No Prompt: ping 8.8.8.8 e depois ping google.com. Se o número responde e o nome não (\"não pôde encontrar o host\"), o problema é a agenda (DNS): do roteador, do provedor ou configurada errada no aparelho.")

    add_heading(doc, "Três frases pra guardar", level=2)
    add_para(doc, "\"O DNS é a agenda da internet.\"", bold=True, color=COR_PRIMARIA)
    add_para(doc, "\"O dono do link vem antes da primeira barra, lido de trás pra frente.\"", bold=True, color=COR_PRIMARIA)
    add_para(doc, "\"Cadeado mais dono certo. Um sem o outro não vale.\"", bold=True, color=COR_PRIMARIA)

    add_heading(doc, "Pra se testar em casa", level=2)
    add_bullet(doc, "Quem é o dono de https://gov.br.consulta-cpf.site/regularizar? É golpe?")
    add_bullet(doc, "Por que \"o site tem cadeado\" não quer dizer que é seguro digitar a senha?")
    add_bullet(doc, "Procure um link que você recebeu no WhatsApp e diga quem é o dono — sem abrir")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Este material é seu. Guarde e volte quando precisar. · CEJA Itapiranga – SC")
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = COR_CINZA

    path = os.path.join(OUTDIR, "Resumo do Aluno - Redes 7 - Nomes na rede.docx")
    doc.save(path)
    return path


# ============================================================
# MAIN
# ============================================================
if __name__ == "__main__":
    print("Gerando Roteiro...")
    print("  →", build_roteiro())
    print("Gerando Slides...")
    print("  →", build_slides())
    print("Gerando Folha do Aluno...")
    print("  →", build_folha())
    print("Gerando Resumo do Aluno...")
    print("  →", build_resumo())
    print("\nPronto.")
